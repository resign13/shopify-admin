import unittest
from io import BytesIO
from unittest.mock import Mock
from openpyxl import load_workbook
from openpyxl.drawing.image import Image
from PIL import Image as PILImage
from order_matrix_export import build, group_items, size_label, SIZES


class OrderMatrixExportTest(unittest.TestCase):
    def orders(self):
        items = [{'productId': 1, 'sku': 'CS5402-Brown', 'image': 'fixture', 'sizeCode': size, 'quantity': qty}
                 for size, qty in [('S', 10), ('28', 2), ('30/M', 8), ('L', 3), ('XL', 4), ('XXL', 5), ('Tall XL', 6)]]
        items += [{'productId': 2, 'sku': 'CS5402-Black', 'image': 'fixture', 'sizeCode': 'S', 'quantity': 7}]
        return [{'orderNo': 'ORDER/1', 'userName': '业务员甲', 'items': items, 'note': '=1+1'},
                {'orderNo': 'ORDER/1', 'userName': '业务员乙', 'items': items[:1]}]

    def export(self, orders=None, split=False, include_images=False):
        source = BytesIO(); PILImage.new('RGB', (100, 100), 'brown').save(source, format='PNG')
        self.fetch = Mock(return_value=source.getvalue())
        return build(orders if orders is not None else self.orders(), split=split, include_images=include_images,
                     fetch_image=self.fetch, make_image=lambda data, **kw: Image(BytesIO(data)),
                     attachments=lambda order: ([], ['https://example.test/label.pdf']))

    def test_aliases_and_special_sizes(self):
        for values, expected in zip([['28', 'S', 'S/28'], ['30', 'M'], ['32', 'L'], ['34', 'XL'], ['36', '2XL', 'XXL', '36/2XL']], SIZES):
            for value in values:
                self.assertEqual(size_label(value), expected)
        self.assertEqual(size_label('Tall XL'), 'Tall XL')
        self.assertEqual(size_label('28/M'), '28/M')
        grouped = group_items(self.orders()[0]['items'])
        self.assertEqual(len(grouped), 2)
        self.assertEqual(grouped[0]['sizes']['28/S'], 12)
        self.assertEqual(sum(grouped[0]['sizes'].values()), 38)

    def test_both_modes_preserve_order_color_and_totals(self):
        for split in [False, True]:
            book = load_workbook(self.export(split=split))
            self.assertEqual(len(book.worksheets), 2 if split else 1)
            sheet = book.worksheets[0]
            self.assertEqual(sheet.freeze_panes, 'A6')
            self.assertFalse(sheet.sheet_view.pane.xSplit)
            self.assertEqual([sheet.cell(5, c).value for c in range(3, 8)], SIZES)
            self.assertEqual(sheet['H5'].value, 'Tall XL')
            self.assertEqual(sheet['I5'].value, '合计 pcs')
            self.assertEqual(sheet['C6'].value, 12)
            self.assertEqual(sheet['I6'].value, 38)
            self.assertEqual(sheet['I7'].value, 7)
            self.assertEqual(sheet['I8'].value, 45)
            all_values = [cell.value for ws in book for row in ws for cell in row]
            self.assertEqual(all_values.count('CS5402-Brown'), 2)
            self.assertEqual(all_values.count('CS5402-Black'), 1)
            self.assertTrue(any('业务员乙' in str(v) for v in all_values))
            self.assertEqual(sheet['A10'].data_type, 's')

    def test_images_once_per_group_and_disabled(self):
        book = load_workbook(self.export(include_images=True))
        self.assertEqual(len(book.active._images), 3)
        self.assertEqual(self.fetch.call_count, 1)
        self.export(include_images=False)
        self.fetch.assert_not_called()

    def test_empty_export(self):
        for split in [False, True]:
            book = load_workbook(self.export([], split=split))
            self.assertEqual(book.active['A1'].value, '暂无订单')


if __name__ == '__main__':
    unittest.main()
