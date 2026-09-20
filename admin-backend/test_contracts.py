import io
import unittest
import uuid
from concurrent.futures import ThreadPoolExecutor
from unittest.mock import patch
from PIL import Image
from openpyxl import load_workbook
from test_workbench import seed, application, db
import contracts

class ContractsTest(unittest.TestCase):
    def setUp(self):
        self.tokens,_=seed()
        application.app.config['TESTING']=True
    def call(self,path='contracts',method='GET',body=None,role='admin'):
        with application.app.test_client() as c:
            return c.open('/api/admin/'+path,method=method,json=body,headers={'Authorization':'Bearer '+self.tokens[role]})
    def payload(self):
        return {'requestId':str(uuid.uuid4()),'partyA':'甲方测试','partyB':'万芳测试','contractNo':'TEST-001','styleImage':'/uploads/fixture.jpg','sizeChartImage':'/uploads/fixture.jpg','factoryName':'万芳测试','deliveryDate':'2026-10-10','note':'测试','items':[{'productId':1,'quantities':{'M':5,'Tall XL':7}}]}
    def inventory(self):
        return db._fetch_all('SELECT size_code,stock,contract_pending,pending_inbound FROM product_size_prices WHERE product_id=1 ORDER BY id')
    def test_atomic_snapshot(self):
        before=self.inventory();body=self.payload();r=self.call(method='POST',body=body)
        self.assertEqual(r.status_code,200,r.json);item=r.json['item'];self.assertEqual(item['quantity'],12)
        for old,new in zip(before,self.inventory()):
            self.assertEqual(old['stock'],new['stock']);self.assertEqual(old['pending_inbound'],new['pending_inbound'])
            self.assertEqual(new['contract_pending'],old['contract_pending']+body['items'][0]['quantities'].get(old['size_code'],0))
        logs=db._fetch_all("SELECT entity_table,batch_id FROM admin_audit_logs WHERE module='contracts'")
        self.assertIn('purchase_contracts',[r['entity_table'] for r in logs]);self.assertIn('product_size_prices',[r['entity_table'] for r in logs]);self.assertEqual(len({r['batch_id'] for r in logs}),1)
        db._fetch_one("UPDATE products SET sku='CHANGED' WHERE id=1 RETURNING id")
        self.assertEqual(self.call('contracts/1').json['item']['items'][0]['sku'],item['items'][0]['sku'])
        self.assertEqual(self.call('contracts?keyword=万芳').json['total'],1)
        self.assertEqual(self.call('contracts?keyword=no-match').json['total'],0)
    def test_concurrent_retry(self):
        body=self.payload()
        with ThreadPoolExecutor(max_workers=2) as pool: responses=list(pool.map(lambda _:self.call(method='POST',body=body),range(2)))
        self.assertEqual([r.status_code for r in responses],[200,200]);self.assertEqual(self.inventory()[0]['contract_pending'],17);self.assertEqual(self.call().json['total'],1)
        body['factoryName']='different';self.assertEqual(self.call(method='POST',body=body).status_code,400)
    def test_rollback(self):
        before=self.inventory()
        for q in [-1,1.5,None,True,'',2147483648]:
            body=self.payload();body['items'].append({'productId':2,'quantities':{'M':q}})
            self.assertEqual(self.call(method='POST',body=body).status_code,400);self.assertEqual(self.inventory(),before)
        for row in [{'productId':2,'quantities':{'WRONG':2}},{'productId':1,'quantities':{'M':2}},{'productId':9999,'quantities':{'M':2}}]:
            body=self.payload();body['items'].append(row);self.assertEqual(self.call(method='POST',body=body).status_code,400)
        body=self.payload();body['sizeChartImage']=''
        self.assertEqual(self.call(method='POST',body=body).status_code,400)
        self.assertEqual(self.call().json['total'],0);self.assertEqual(db._fetch_one("SELECT count(*) AS n FROM admin_audit_logs WHERE module='contracts'")['n'],0)
    def test_permissions(self):
        for role in ['customer','warehouse']:
            self.assertEqual(self.call(role=role).status_code,403);self.assertEqual(self.call(method='POST',body=self.payload(),role=role).status_code,403)
        db._fetch_one("UPDATE admin_users SET permissions='[]' WHERE role='sales' RETURNING id");self.assertEqual(self.call(role='sales').status_code,403)
        db._fetch_one("UPDATE admin_users SET permissions='[\"contracts\"]' WHERE role='sales' RETURNING id")
        self.assertEqual(self.call(method='POST',body=self.payload(),role='sales').status_code,200);self.assertEqual(self.call('products/1',role='sales').status_code,200)
    def test_export(self):
        body=self.payload();body['factoryName']='=formula';item=self.call(method='POST',body=body).json['item']
        raw=io.BytesIO();Image.new('RGB',(1800,1200),'navy').save(raw,format='PNG')
        ws=load_workbook(contracts.export(item,lambda url:raw.getvalue(),application.build_excel_image)).active
        self.assertEqual([ws.cell(4,c).value for c in range(3,8)],['S','M','L','XL','XXL']);self.assertEqual([ws.cell(5,c).value for c in [4,5,8,9]],[5,0,7,12]);self.assertEqual(len(ws._images),2);self.assertIsNone(ws.freeze_panes);self.assertEqual(ws['A2'].data_type,'s')
        with patch.object(application,'fetch_image_bytes',return_value=raw.getvalue()):self.assertEqual(self.call('contracts/1/export').status_code,200)
        with self.assertRaises(ValueError):contracts.export(item,lambda url:None,application.build_excel_image)

    def test_modify_cancel_delta_and_conflict(self):
        before=self.inventory()
        original=self.call(method='POST',body=self.payload()).json['item']
        body=self.payload();body['revision']=original['revision']
        body['items'][0]['quantities']={'M':8,'Tall XL':3}
        r=self.call('contracts/1','PUT',body)
        self.assertEqual(r.status_code,200,r.json)
        self.assertEqual(self.inventory()[0]['contract_pending'],20)
        self.assertEqual(self.call('contracts/1','PUT',body).status_code,409)
        revision=r.json['item']['revision']
        r=self.call('contracts/1/cancel','POST',{'revision':revision})
        self.assertEqual(r.status_code,200,r.json)
        self.assertEqual(self.inventory(),before)
        self.assertEqual(self.call('contracts/1/cancel','POST',{'revision':revision}).status_code,200)
        self.assertEqual(self.inventory(),before)

    def test_cancel_insufficient_pending_rolls_back(self):
        self.call(method='POST',body=self.payload())
        db._fetch_one("UPDATE product_size_prices SET contract_pending=4 WHERE product_id=1 AND size_code='M' RETURNING id")
        before=self.inventory()
        r=self.call('contracts/1/cancel','POST',{'revision':1})
        self.assertEqual(r.status_code,400,r.json)
        self.assertEqual(self.inventory(),before)
        self.assertFalse(self.call('contracts/1').json['item']['cancelled'])

    def test_standard_mapping_and_creator(self):
        body=self.payload();body['creatorName']='forged'
        body['items'][0].update(quantities={'S':0,'M':3,'L':0,'XL':4,'XXL':0},sizeMapping={'M':'M','XL':'Tall XL'})
        r=self.call(method='POST',body=body)
        self.assertEqual(r.status_code,200,r.json)
        item=r.json['item'];self.assertEqual(item['items'][0]['quantities'],{'M':3,'Tall XL':4})
        self.assertNotEqual(item['creatorName'],'forged')
        self.assertEqual(self.call().json['items'][0]['contractNo'],'TEST-001')
        body=self.payload();body['items'][0].update(quantities={'S':1},sizeMapping={})
        self.assertEqual(self.call(method='POST',body=body).status_code,400)
