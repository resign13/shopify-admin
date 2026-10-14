"""Current cutover scope: real PostgreSQL/cookies, deferred features stay off.

All fixture writes inherit test_workbench's explicit loopback *_test guard.
"""
import io
import os
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import Mock, patch

from flask import Flask
from flask.testing import FlaskClient
from PIL import Image

from test_workbench import application, db, seed
from security_core import Security, validate_password


def photo():
    raw = io.BytesIO()
    Image.new('RGB', (4, 4), 'white').save(raw, format='PNG')
    return raw.getvalue()


class PhaseOneSecurityTest(unittest.TestCase):
    def setUp(self):
        self.env = patch.dict(os.environ, {'SECURITY_ATTACHMENT_MODE': 'legacy', 'SECURITY_LOGIN_PROTECTION_ENABLED': '0'})
        self.env.start(); self.addCleanup(self.env.stop)
        self.media = patch.object(application.media, 'mode', 'legacy')
        self.media.start(); self.addCleanup(self.media.stop)
        self.login_gate = patch.object(application.security, 'login_protection', False)
        self.login_gate.start(); self.addCleanup(self.login_gate.stop)
        self.tokens, self.products = seed()
        self.client = FlaskClient(application.app, application.app.response_class)
        csrf = self.client.get('/api/auth/csrf').json['csrfToken']
        self.preheaders = {'Origin': 'http://localhost', 'X-CSRF-Token': csrf}

    def login(self):
        response = self.client.post('/api/auth/admin/login', json={'email': 'admin@gingtto.test', 'password': 'Workbench-Test-2026'}, headers=self.preheaders)
        self.assertEqual(response.status_code, 200, response.json)
        self.headers = {'Origin': 'http://localhost', 'X-CSRF-Token': response.json['csrfToken'], 'X-Auth-Session': response.json['sessionId']}
        return response

    def test_production_cookie_boundary_has_no_private_media_prerequisite(self):
        with patch.dict(os.environ, {'APP_ENV': 'production', 'SECURITY_SECRET': 'phase-one-production-probe-0123456789abcdef', 'SECURITY_ALLOWED_HOSTS': 'admin.fixture.test', 'SECURITY_LEGACY_SOURCES_CLOSED': '0'}):
            unused_db = Mock()
            unused_db._connect.side_effect = AssertionError('Legacy mode must not require a private-media marker')
            boundary = Security(Flask('phase-one-prod'), unused_db, 'admin')
            self.assertTrue(boundary.secure)
            self.assertEqual(boundary.cookie, '__Host-gingtto-admin-session')
            self.assertEqual(boundary.attachment_mode, 'legacy')
            self.assertFalse(boundary.login_protection)
            unused_db._connect.assert_not_called()

    def test_private_profile_still_requires_honest_source_closure(self):
        with patch.dict(os.environ, {'APP_ENV': 'production', 'SECURITY_SECRET': 'phase-one-production-probe-0123456789abcdef', 'SECURITY_ATTACHMENT_MODE': 'private', 'SECURITY_LEGACY_SOURCES_CLOSED': '0'}):
            with self.assertRaisesRegex(RuntimeError, 'closure'):
                Security(Flask('future-private-profile'), Mock(), 'admin')

    def test_deferred_login_policies_do_not_create_login_counters(self):
        for _ in range(101):
            r = self.client.post('/api/auth/admin/login', json={'email': 'sales@gingtto.test', 'password': 'wrong'}, headers=self.preheaders)
            self.assertEqual(r.status_code, 401)
            self.assertEqual(r.json['message'], 'Invalid email or password')
            self.assertNotIn('Retry-After', r.headers)
        self.login()
        keys = [application.security.digest('quota:admin:login-ip:127.0.0.1'), application.security.digest('login-account:admin:sales@gingtto.test')]
        self.assertIsNone(db._fetch_one('SELECT 1 FROM security_rate_buckets WHERE bucket_key=ANY(%s)', (keys,)))

    def test_deferred_password_rule_preserves_account_creation_and_change(self):
        self.login()
        validate_password('short')
        # Synthetic accounts only. No synced or production account is changed.
        r = self.client.post('/api/admin/admin-users', json={'name': 'Phase one fixture', 'email': 'phase-one@fixture.test', 'password': 'short', 'role': 'sales'}, headers=self.headers)
        self.assertEqual(r.status_code, 201, r.json)
        uid = r.json['user']['id']
        r = self.client.put('/api/admin/admin-users/' + str(uid), json={'password': '123456'}, headers=self.headers)
        self.assertEqual(r.status_code, 200, r.json)
        other = self.client.post('/api/admin/store-users', json={'name': 'Buyer fixture', 'email': 'phase-buyer@fixture.test', 'password': 'short'}, headers=self.headers)
        self.assertEqual(other.status_code, 201, other.json)
        r = self.client.put('/api/admin/store-users/' + str(other.json['user']['id']), json={'password': '123456'}, headers=self.headers)
        self.assertEqual(r.status_code, 200, r.json)

    def test_cookie_csrf_and_old_endpoint_controls_remain_active(self):
        self.assertEqual(self.client.get('/api/admin/orders').status_code, 401)
        self.assertEqual(self.client.get('/api/public/products').status_code, 410)
        self.assertEqual(self.client.get('/api/internal/products').status_code, 404)
        signed = self.login()
        self.assertNotIn('token', signed.json)
        self.assertTrue(self.client.get_cookie(application.security.cookie).http_only)
        self.assertEqual(self.client.post('/api/admin/orders', json={}).status_code, 403)
        self.assertEqual(self.client.get('/api/auth/me').headers['Cache-Control'], 'no-store')

    def test_only_explicit_auth_routes_receive_login_exemption(self):
        probe = Flask('exact-login-routes')
        Security(probe, db, 'admin')
        probe.add_url_rule('/api/admin/reports/login', 'not_a_login', lambda: {'secret': 'fixture'})
        guest = FlaskClient(probe, probe.response_class)
        self.assertEqual(guest.get('/api/admin/reports/login').status_code, 401)

    def test_three_legacy_exports_keep_images_and_attachment_links(self):
        from openpyxl import load_workbook
        from zipfile import ZipFile
        self.login()
        with TemporaryDirectory() as directory, patch.dict(application.media.uploads, {'admin': Path(directory)}):
            (Path(directory) / 'fixture.png').write_bytes(photo())
            pdf = 'https://old.fixture.test/packing.pdf'
            order = {'orderNo': 'LM-000001', 'contactName': 'Fixture Buyer', 'shippingFee': 10, 'note': 'Fixture packing note', 'labelImageUrls': ['/uploads/fixture.png', pdf], 'labelPdfUrl': pdf, 'items': [{'sku': 'ONE', 'productName': 'Fixture', 'sizeCode': 'M', 'quantity': 2, 'unitPrice': 20, 'image': '/uploads/fixture.png'}]}
            for builder in (application.build_orders_export, application.build_orders_sheet_export, application.build_order_invoice_export):
                with self.subTest(export=builder.__name__), application.app.test_request_context('/api/admin/orders/export'):
                    stream = builder(order if builder == application.build_order_invoice_export else [order])
                    book = load_workbook(stream)
                    with ZipFile(stream) as archive:
                        self.assertGreaterEqual(len([n for n in archive.namelist() if n.startswith('xl/media/')]), 2)
                    # Existing files list URLs in their remarks cells; this
                    # phase preserves that format instead of adding hyperlinks.
                    values = '\n'.join(str(c.value) for ws in book for row in ws for c in row if c.value is not None)
                    self.assertIn(pdf, values)
                    self.assertIn('LM-000001', values)
                    self.assertNotIn('/api/attachments/', values)

    def test_legacy_upload_url_preview_and_reference_are_preserved(self):
        self.login()
        with TemporaryDirectory() as directory, patch.dict(application.media.uploads, {'admin': Path(directory)}), patch.object(application, 'cloudinary_enabled', return_value=False):
            uploaded = self.client.post('/api/admin/uploads', data={'files': (io.BytesIO(photo()), 'fixture.png'), 'purpose': 'order'}, headers=self.headers)
            self.assertEqual(uploaded.status_code, 200, uploaded.json)
            url = uploaded.json['urls'][0]
            self.assertIn('/uploads/', url)
            self.assertNotIn('/api/attachments/', url)
            guest = FlaskClient(application.app, application.app.response_class)
            response = guest.get(url)
            self.assertEqual(response.status_code, 200)
            response.close()
            original = {'labelImageUrls': [url], 'labelPdfUrl': 'https://old.fixture.test/document.pdf'}
            self.assertIs(application.media.rewrite(original), original)
            application.media.bind('order', 9999, [url])
            self.assertIsNone(db._fetch_one('SELECT 1 FROM security_file_refs WHERE entity_type=%s AND entity_id=%s', ('order', 9999)))
            self.assertEqual(application.media.canonical(url), url)
            with application.app.test_request_context('/api/admin/orders/export'):
                self.assertEqual(application.fetch_image_bytes(url, attachment=True), photo())

    def test_three_exports_show_missing_images_without_altering_amount_formulas(self):
        from openpyxl import load_workbook
        order = {'orderNo': 'LM-000001', 'shippingFee': 10, 'labelImageUrls': ['https://outside.fixture.test/attachment.png'], 'items': [{'sku': 'ONE', 'productName': 'Fixture', 'sizeCode': 'M', 'quantity': 2, 'unitPrice': 20, 'image': 'https://outside.fixture.test/product.png'}]}
        for builder in (application.build_orders_export, application.build_orders_sheet_export, application.build_order_invoice_export):
            with self.subTest(export=builder.__name__), patch.object(application, 'fetch_image_bytes', return_value=None):
                stream = builder(order if builder == application.build_order_invoice_export else [order])
                book = load_workbook(stream)
                values = '\n'.join(str(c.value) for ws in book for row in ws for c in row if c.value is not None)
                self.assertIn('Image unavailable', values)
                self.assertIn('Attachment images unavailable: 1', values)
                if builder == application.build_order_invoice_export:
                    self.assertEqual(book['PI']['J23'].value, '=J21+J22')
                    self.assertEqual(book['PI']['I13'].value, 20)

    def test_legacy_exports_read_owned_files_without_network_and_reject_traversal(self):
        with TemporaryDirectory() as directory, patch.dict(application.media.uploads, {'admin': Path(directory)}):
            (Path(directory) / 'fixture.png').write_bytes(photo())
            with application.app.test_request_context('/api/admin/orders/export'):
                with patch('safe_images.fetch_image', side_effect=AssertionError('Owned local upload must not use HTTP')):
                    self.assertEqual(application.fetch_image_bytes('/uploads/fixture.png', attachment=True), photo())
                for value in ['/uploads/../secret.png', '/uploads/a/../../secret.png', '/uploads/%2e%2e/secret.png']:
                    with self.assertRaises(ValueError):
                        application.media.legacy_export_source(value)
                self.assertIsNone(application.media.legacy_export_source('https://outside.invalid/uploads/fixture.png'))

    def test_legacy_mode_still_validates_new_upload_content(self):
        self.login()
        with patch.object(application, 'save_file_locally') as writer:
            for raw, name in [(b'MZ invalid image', 'bad.png'), (photo(), 'mismatch.jpg'), (b'<script>bad</script>', 'bad.html')]:
                r = self.client.post('/api/admin/uploads', data={'files': (io.BytesIO(raw), name)}, headers=self.headers)
                self.assertEqual(r.status_code, 400, r.json)
            writer.assert_not_called()

    def test_fixture_upload_and_thumbnail_cache_never_use_manual_directories(self):
        import hashlib
        import image_delivery
        original = application.media.uploads['admin'] / 'fixture.jpg'
        before = hashlib.sha256(original.read_bytes()).hexdigest()
        self.assertNotIn(application.BASE_DIR, application.media.cache.parents)
        with patch.object(image_delivery, 'r2_client', side_effect=AssertionError('Local fixture must not fetch external objects')):
            response = self.client.get('/uploads/fixture.jpg?w=160')
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.content_type, 'image/webp')
            response.close()
        self.assertTrue(list(application.media.cache.glob('*.webp')))
        self.assertEqual(hashlib.sha256(original.read_bytes()).hexdigest(), before)

    def test_fixture_cannot_reuse_real_or_cached_storage_clients(self):
        import image_delivery
        import private_storage
        from werkzeug.datastructures import FileStorage
        with patch.object(image_delivery, '_r2_client') as public, patch.object(private_storage, '_client') as private, \
                patch.object(application, 'CLOUDINARY_CLOUD_NAME', 'fixture'), patch.object(application, 'CLOUDINARY_API_KEY', 'fixture'), patch.object(application, 'CLOUDINARY_API_SECRET', 'fixture'):
            self.assertFalse(application.cloudinary_enabled())
            for get_client in (image_delivery.r2_client, private_storage.client):
                with self.assertRaisesRegex(RuntimeError, 'isolated fixtures'): get_client()
            with self.assertRaisesRegex(RuntimeError, 'isolated fixtures'):
                application.upload_file_to_cloudinary(FileStorage(stream=io.BytesIO(photo()), filename='fixture.png'))
            public.assert_not_called(); private.assert_not_called()
            # Outside test mode the same configured providers remain available;
            # these explicit doubles make no network call.
            with patch.dict(os.environ, {'APP_ENV': 'local'}):
                self.assertIs(image_delivery.r2_client(), public.return_value)
                self.assertIs(private_storage.client(), private.return_value)
                self.assertTrue(application.cloudinary_enabled())


if __name__ == '__main__':
    unittest.main()
