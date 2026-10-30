"""XLSX and filter parity checks, on the guarded isolated *_test database only."""
import io
import unittest
from datetime import datetime
from unittest.mock import patch
from urllib.parse import urlencode

from openpyxl import load_workbook
import test_dashboard as fixture
from test_workbench import application, db
import dashboard_export


class DashboardExportTest(unittest.TestCase):
    setUp = fixture.DashboardTest.setUp
    product = fixture.DashboardTest.product
    order = fixture.DashboardTest.order
    line = fixture.DashboardTest.line
    call = fixture.DashboardTest.call
    data = fixture.DashboardTest.data

    def export(self, role='admin', **params):
        query = {'dateFrom': '2026-09-18', 'dateTo': '2026-09-18', **params}
        with application.app.test_client() as client:
            headers = {'Authorization': 'Bearer ' + self.tokens[role]} if role else {}
            return client.get('/api/admin/dashboard/style-performance/export?' + urlencode(query), headers=headers)

    def workbook(self, **params):
        response = self.export(**params)
        self.assertEqual(response.status_code, 200, response.json if response.is_json else response.status)
        self.assertIn('spreadsheetml.sheet', response.content_type)
        self.assertEqual(response.headers['Cache-Control'], 'no-store')
        self.assertIn('.xlsx', response.headers['Content-Disposition'])
        self.addCleanup(response.close)
        book = load_workbook(io.BytesIO(response.data))
        self.addCleanup(book.close)
        return book.active

    def records(self, sheet):
        columns = [cell.value for cell in sheet[dashboard_export.HEADER_ROW]]
        return [dict(zip(columns, row)) for row in sheet.iter_rows(
            min_row=dashboard_export.HEADER_ROW + 1, max_row=sheet.max_row - 1, values_only=True)]

    def test_full_filtered_result_not_page_with_formats_totals_and_beijing_time(self):
        for index in range(12):
            self.product(f'EXTRA-{index:02}', stock=index)
        query = dict(sort='styleCode', direction='asc')
        page = self.data(page=2, pageSize=10, **query)
        self.assertEqual(len(page['items']), 3)
        sheet = self.workbook(page=2, pageSize=10, **query)
        rows = self.records(sheet)
        self.assertEqual(len(rows), 13)
        self.assertEqual([r['款号'] for r in rows], sorted(r['款号'] for r in rows))
        total = dict(zip([c.value for c in sheet[1]], [c.value for c in sheet[sheet.max_row]]))
        self.assertEqual(total['销量（件）'], page['summary']['units'])
        self.assertAlmostEqual(total['商品金额（USD，不含运费）'], page['summary']['amount'])
        self.assertEqual(total['现货余额'], page['summary']['stock'])
        self.assertEqual(sheet['A1'].value, '款号')
        self.assertFalse(sheet.merged_cells.ranges)
        self.assertEqual(sheet.freeze_panes, 'A2')
        self.assertEqual(sheet.sheet_view.pane.xSplit, None)
        self.assertEqual(sheet.sheet_view.pane.ySplit, 1)
        self.assertTrue(sheet.auto_filter.ref.startswith('A1:'))
        self.assertTrue(sheet.auto_filter.ref.endswith(str(sheet.max_row - 1)))
        self.assertEqual(sheet['A2'].data_type, 's')
        self.assertEqual(sheet['D2'].number_format, '#,##0.00')
        sold = next(row for row in rows if row['款号'] == 'GT-2026')
        self.assertEqual(sold['最近销售时间（北京时间）'], datetime(2026, 9, 18, 17))
        self.assertEqual(sold['库存状态'], '库存健康')
        self.assertEqual(sold['预计可售（天）'], 70)

    def test_filter_and_sort_parity_including_current_stock(self):
        coat = self.product('COAT', stock=5)
        self.order(coat, 4, price=100)
        self.product('UNSOLD', stock=0)
        db._fetch_one("UPDATE product_size_prices SET stock=-5 WHERE product_id=%s AND size_code='M' RETURNING id", (self.products[0]['id'],))
        for query in [
            {'country': 'Germany', 'category': 'denim', 'keyword': '水洗灰'},
            {'style': 'GT-2026-1'}, {'risk': 'backordered'}, {'risk': 'out_of_stock'},
            {'sort': 'amount', 'direction': 'asc'}, {'sort': 'stock', 'direction': 'desc'},
            {'velocityWindow': 30, 'sort': 'estimatedDays', 'direction': 'asc'},
        ]:
            with self.subTest(query=query):
                expected = self.data(**query)
                sheet = self.workbook(**query)
                rows = self.records(sheet)
                self.assertEqual([r['款号'] for r in rows], [r['styleCode'] for r in expected['items']])
                for row, item in zip(rows, expected['items']):
                    for label, key in [('销量（件）', 'units'), ('商品金额（USD，不含运费）', 'amount'),
                                       ('现货余额', 'stock'), ('可用现货', 'availableStock'), ('欠货件数', 'shortageUnits'),
                                       ('合同未送', 'contractPending'), ('待验货', 'pendingInspection'),
                                       ('合格', 'pendingInbound'), ('待打回次品', 'defectivePending'),
                                       ('临时待入库', 'temporaryInbound'), ('日均销量（件）', 'averageDailyUnits')]:
                        self.assertEqual(row[label], item[key], label)

    def test_personal_scope_does_not_export_company_stock_or_other_owner_sales(self):
        own = db._fetch_one("SELECT id FROM admin_users WHERE role='sales'")['id']
        db._fetch_one("UPDATE orders SET owner_admin_id=%s WHERE status='paid' RETURNING id", (own,))
        for params in [{'role': 'sales', 'salespersonId': 'all', '_ownerId': 1},
                       {'salespersonId': own}, {'salespersonId': 'unassigned'}]:
            with self.subTest(params=params):
                expected = self.data(**params)
                sheet = self.workbook(**params)
                headers = [c.value for c in sheet[1]]
                self.assertNotIn('现货余额', headers)
                self.assertNotIn('库存状态', headers)
                self.assertNotIn('合同未送', headers)
                self.assertEqual(sum(r['销量（件）'] for r in self.records(sheet)), expected['summary']['units'])
                self.assertAlmostEqual(sum(r['商品金额（USD，不含运费）'] for r in self.records(sheet)), expected['summary']['amount'])
        all_rows = self.records(self.workbook(salespersonId='all'))
        self.assertIn('现货余额', all_rows[0])
        self.assertEqual(sum(r['销量（件）'] for r in all_rows), 6)

    def test_permissions_and_revocation_match_dashboard(self):
        self.assertEqual(self.export(role=None).status_code, 401)
        for role in ('customer', 'warehouse'):
            self.assertEqual(self.export(role=role).status_code, 403)
        db._fetch_one("UPDATE admin_users SET permissions='[]'::jsonb WHERE role='sales' RETURNING id")
        self.assertEqual(self.export(role='sales').status_code, 403)
        db._fetch_one("UPDATE admin_users SET permissions='[\"dashboard\"]'::jsonb WHERE role='warehouse' RETURNING id")
        self.assertEqual(self.export(role='warehouse').status_code, self.call(role='warehouse').status_code)

    def test_invalid_filters_empty_results_and_export_quota(self):
        for params in [{'dateFrom': 'bad'}, {'dateFrom': '2020-01-01'}, {'dateFrom': '2026-09-20'},
                       {'risk': 'bad'}, {'velocityWindow': 14}, {'salespersonId': '9999999'}]:
            self.assertEqual(self.export(**params).status_code, 400, params)
        sheet = self.workbook(keyword='does-not-exist')
        self.assertEqual(self.records(sheet), [])
        self.assertEqual(sheet['C2'].value, 0)
        with patch.dict('os.environ', {'SECURITY_LIMIT_EXPORT': '1'}):
            response = self.export()
            self.assertEqual(response.status_code, 429)
            self.assertIn('Retry-After', response.headers)

    def test_untrusted_product_text_remains_literal_and_export_does_not_write_business(self):
        db._fetch_one("UPDATE products SET style_code='0000123' WHERE id=1 RETURNING id")
        db._fetch_one("UPDATE product_translations SET name='=1+1' WHERE product_id=1 RETURNING product_id")
        before = db._fetch_all('SELECT id,stock FROM product_size_prices ORDER BY id')
        before_orders = db._fetch_all('SELECT id,total_amount,status,updated_at FROM orders ORDER BY id')
        sheet = self.workbook(style='0000123')
        self.assertEqual(sheet['A2'].value, '0000123')
        self.assertEqual(sheet['A2'].number_format, '@')
        self.assertEqual(sheet['B2'].value, '=1+1')
        self.assertEqual(sheet['B2'].data_type, 's')
        self.assertEqual(before, db._fetch_all('SELECT id,stock FROM product_size_prices ORDER BY id'))
        self.assertEqual(before_orders, db._fetch_all('SELECT id,total_amount,status,updated_at FROM orders ORDER BY id'))


if __name__ == '__main__':
    unittest.main()
