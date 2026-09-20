"""Synthetic invoice export regression tests; no real order writes."""
import unittest
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile
from unittest.mock import patch
from openpyxl import load_workbook
from PIL import Image
from test_workbench import application


class InvoiceExportTest(unittest.TestCase):
    def test_all_exports_include_note_and_all_nine_attachments(self):
        photo = BytesIO(); Image.new('RGB', (40, 60), '#335577').save(photo, format='PNG')
        order = {'orderNo': 'NINE-ATTACHMENTS', 'shippingFee': 10, 'note': 'Nine photo packing instructions',
                 'labelImageUrls': [f'/uploads/attachment-{i}.png' for i in range(9)],
                 'items': [{'productId': 1, 'sku': 'BLUE', 'sizeCode': 'M', 'quantity': 1, 'unitPrice': 20}]}
        for builder in [application.build_orders_export, application.build_orders_sheet_export, application.build_order_invoice_export]:
            with self.subTest(builder=builder.__name__), patch.object(application, 'fetch_image_bytes', side_effect=lambda url, **kwargs: photo.getvalue() if url else None):
                stream = builder(order if builder == application.build_order_invoice_export else [order])
                book = load_workbook(stream)
                baseline_order = {**order, 'labelImageUrls': []}
                baseline = load_workbook(builder(baseline_order if builder == application.build_order_invoice_export else [baseline_order]))
                self.assertEqual(sum(len(ws._images) for ws in book) - sum(len(ws._images) for ws in baseline), 9)
                self.assertTrue(any(order['note'] in str(c.value) for ws in book for row in ws for c in row))
                self.assertFalse(any('/uploads/attachment-' in str(c.value) for ws in book for row in ws for c in row))
                attachment_pictures = [pic for ws in book for pic in ws._images if pic.width == 40 and pic.height == 60]
                self.assertEqual(len(attachment_pictures), 9)
                self.assertEqual(len({pic.anchor._from.row for pic in attachment_pictures}), 1)
                positions = [(pic.anchor._from.col, pic.anchor._from.colOff) for pic in attachment_pictures]
                self.assertEqual(positions, sorted(set(positions)))


    def test_both_exports_embed_clear_images_and_reuse_downloads(self):
        with TemporaryDirectory() as directory:
            Image.new('RGB', (2400, 3200), '#335577').save(Path(directory) / 'fixture.jpg')
            order = {'orderNo': 'IMAGE-FIXTURE', 'items': [
                {'sku': 'SKU-1', 'productName': 'Fixture', 'image': 'https://img.smawell.shop/uploads/fixture.jpg',
                 'sizeCode': size, 'quantity': 1, 'unitPrice': 20}
                for size in ['S', 'M', 'L']
            ]}
            for builder in (application.build_orders_export, application.build_orders_sheet_export):
                with self.subTest(export=builder.__name__), patch.object(application, 'UPLOAD_DIR', Path(directory)), \
                        application.app.test_request_context('/api/admin/orders/export'):
                    from image_delivery import deliver_image
                    with patch('image_delivery.deliver_image', wraps=deliver_image) as read:
                        stream = builder([order])
                        self.assertEqual(read.call_count, 1)
                    with ZipFile(stream) as archive:
                        media = [name for name in archive.namelist() if name.startswith('xl/media/')]
                        self.assertEqual(len(media), 1)
                        for name in media:
                            content = archive.read(name)
                            self.assertLess(len(content), 50000)
                            with Image.open(BytesIO(content)) as image:
                                self.assertEqual(image.width, 480)
                                self.assertEqual(image.height, 640)

    def test_attachment_quality_is_independent_of_product_thumbnail(self):
        with TemporaryDirectory() as directory:
            Image.new('RGB', (2400, 3200), '#335577').save(Path(directory) / 'large.jpg')
            url = '/uploads/large.jpg'
            order = {'orderNo': 'CLEAR-ATTACHMENT', 'shippingFee': 10, 'labelImageUrls': [url, url],
                     'items': [{'sku': 'ONE', 'image': url, 'quantity': 1, 'unitPrice': 20}]}
            for builder in (application.build_orders_export, application.build_orders_sheet_export, application.build_order_invoice_export):
                with self.subTest(export=builder.__name__), patch.object(application, 'UPLOAD_DIR', Path(directory)), application.app.test_request_context('/'):
                    stream = builder(order if builder == application.build_order_invoice_export else [order])
                    book = load_workbook(stream)
                    with ZipFile(stream) as archive:
                        dimensions = []
                        for name in archive.namelist():
                            if name.startswith('xl/media/'):
                                with Image.open(BytesIO(archive.read(name))) as picture:
                                    dimensions.append(picture.size)
                        self.assertIn((1200, 1600), dimensions)
                    pictures = [pic for ws in book for pic in ws._images if pic.width == 1200]
                    self.assertEqual(len(pictures), 2)
                    self.assertEqual(pictures[0].anchor._from.row, pictures[1].anchor._from.row)
                    self.assertGreater(pictures[1].anchor._from.col, pictures[0].anchor._from.col)
                    for ws in book:
                        for picture in ws._images:
                            if picture.width == 1200:
                                self.assertGreaterEqual(ws.row_dimensions[picture.anchor._from.row + 1].height, 300)

    def test_invoice_product_picture_keeps_640_pixels(self):
        photo = BytesIO(); Image.new('RGB', (1800, 1800), '#335577').save(photo, format='PNG')
        order = {'orderNo': 'CLEAR-PI', 'shippingFee': 10, 'items': [
            {'sku': 'ONE', 'image': 'https://example.test/product.png', 'quantity': 1, 'unitPrice': 20}]}
        with patch.object(application, 'fetch_image_bytes', return_value=photo.getvalue()):
            book = load_workbook(application.build_order_invoice_export(order))
        self.assertTrue(any(pic.width == 640 and pic.height == 640 for pic in book.active._images))

    def test_parallel_prefetch_and_disk_cache_reuse_originals(self):
        from threading import Barrier
        from unittest.mock import MagicMock
        barrier = Barrier(2)
        photo = BytesIO(); Image.new('RGB', (1800, 1800), '#335577').save(photo, format='PNG')
        urls = ['https://example.test/one.png', 'https://example.test/two.png']
        order = {'items': [{'image': urls[0]}], 'labelImageUrls': urls}
        def download(*args, **kwargs):
            barrier.wait(timeout=10)
            response = MagicMock()
            response.__enter__.return_value = response
            response.headers = {'Content-Type': 'image/png'}
            response.read.return_value = photo.getvalue()
            return response
        with TemporaryDirectory() as directory, patch.object(application, 'BASE_DIR', Path(directory)), \
                patch.object(application.urllib_request, 'urlopen', side_effect=download) as network:
            for repeat in range(2):
                with application.app.test_request_context('/'):
                    application.prefetch_export_images([order])
                    for url in urls:
                        self.assertEqual(application.fetch_image_bytes(url, attachment=True), photo.getvalue())
                    self.assertEqual(application.fetch_image_bytes(urls[0]), photo.getvalue())
            self.assertEqual(network.call_count, 2)

    def test_invoice_with_few_and_many_product_rows(self):
        for count in (1, 8, 20, 60):
            with self.subTest(products=count), patch.object(application, 'fetch_image_bytes', return_value=None):
                order = {
                    'orderNo': 'SYNTHETIC-INVOICE', 'shippingFee': 15,
                    'contactName': 'Fixture', 'items': [
                        {'productName': f'Product {index}', 'sku': f'SKU-{index}',
                         'sizeCode': 'M', 'quantity': index + 1, 'unitPrice': 20}
                        for index in range(count)
                    ],
                }
                output = application.build_order_invoice_export(order)
                book = load_workbook(output)
                sheet = book['PI']
                self.assertEqual(sheet.max_column, 10)
                self.assertTrue(all((d.max or d.min or 0) <= 10 for d in sheet.column_dimensions.values()))
                last_item = 12 + max(count, 8)
                self.assertEqual(sheet['B5'].value, 'SYNTHETIC-INVOICE')
                self.assertEqual(sheet.cell(last_item + 1, 10).value, f'=SUM(J13:J{last_item})')
                self.assertEqual(sheet.cell(last_item + 2, 10).value, 15)
                for index in range(count):
                    row = 13 + index
                    self.assertEqual(sheet.cell(row, 8).value, f'=SUM(C{row}:G{row})')
                    self.assertEqual(sheet.cell(row, 10).value, f'=H{row}*I{row}')
                    self.assertIn(f'SKU-{index}', str(sheet.cell(row, 1).value))
                self.assertTrue(any('BANK' in str(cell.value).upper() for row in sheet for cell in row))
                book.close()


if __name__ == '__main__':
    unittest.main()
