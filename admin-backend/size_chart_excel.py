"""Full-value maintenance workbooks; English customer workbooks never contain hidden data."""
from copy import deepcopy
from decimal import Decimal
from io import BytesIO
import hmac
import json
import re
from zipfile import ZipFile

from flask import current_app, g, request, send_file
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.utils import get_column_letter

import size_charts as charts
import workbench

VERSION = 'gingtto-size-charts-v2'
LEGACY_VERSION = 'gingtto-size-charts-v1'
META = '_尺码表模板'
MAX_BYTES, MAX_EXPANDED = 10 * 1024 * 1024, 50 * 1024 * 1024


def packed(value): return json.dumps(value, ensure_ascii=False, separators=(',', ':'), sort_keys=True)
def sign(value): return current_app.extensions['security'].digest('size-charts:' + packed(value))
def structure(doc):
    return {k: doc[k] for k in ('styleCode', 'version', 'catalogVersion', 'sizes', 'dimensions')}


def write(sheet, row, col, value):
    cell = sheet.cell(row, col, value)
    if isinstance(value, str): cell.data_type = 's'
    cell.alignment = Alignment(vertical='center', wrap_text=True)
    cell.font = Font(name='Arial', size=11, color='243850')
    if isinstance(value, Decimal): cell.number_format = '0.00'
    return cell


def export(payload):
    if not isinstance(payload, dict): raise ValueError('导出参数无效')
    codes, mode = payload.get('styleCodes'), payload.get('mode', 'maintenance')
    target_unit = payload.get('unit')
    if target_unit is not None and (not isinstance(target_unit, str) or target_unit not in {'cm', 'in'}):
        raise ValueError('导出单位须为 cm 或 in')
    if mode not in {'maintenance', 'english'}: raise ValueError('导出版本无效')
    if not isinstance(codes, list) or not 1 <= len(codes) <= charts.MAX_STYLES or any(not isinstance(c, str) for c in codes) or len(set(codes)) != len(codes):
        raise ValueError('请选择 1–100 个不同款号')
    cats = charts.catalog(codes); docs = [charts.detail(c, cats) for c in codes]
    book = Workbook(); book.remove(book.active)
    metadata = []
    for index, saved_doc in enumerate(docs, 1):
        doc = charts.unified_document(saved_doc, target_unit) if target_unit else saved_doc
        if not doc['sizes'] or len(doc['sizes']) > charts.MAX_SIZES: raise ValueError(doc['styleCode'] + '：需配置 1–128 个真实尺码')
        dims = charts.english(doc) if mode == 'english' else doc['dimensions']
        title = re.sub(r'[\\/*?:\[\]]', '_', doc['styleCode']).strip("'")[:23] or 'Style'
        sheet = book.create_sheet(f'{index:03}-{title}')
        english = mode == 'english'
        headers = ['Measurement', 'Unit', *doc['sizes']] if english else ['测量维度', 'English', '单位', *doc['sizes']]
        start = 3 if english else 4
        last_col = max(len(headers), 5)
        sheet.column_dimensions['A'].width = 25
        sheet.column_dimensions['B'].width = 12 if english else 27
        if not english: sheet.column_dimensions['C'].width = 10
        for col in range(start, last_col+1): sheet.column_dimensions[get_column_letter(col)].width = 12
        sheet.merge_cells('A1:A5')
        for row in range(1,6):
            sheet.row_dimensions[row].height = 24 if row != 1 else 30
            for col in range(1,last_col+1):
                sheet.cell(row,col).fill = PatternFill('solid',fgColor='F3F6FA')
        # Compact information block beside the photo; measurement grid uses full width below.
        title_cell = write(sheet, 1, 2, doc['titleEn'] if english else doc['title'])
        title_cell.font = Font(name='Arial',size=16,bold=True,color='243850')
        sheet.merge_cells(start_row=1,start_column=2,end_row=1,end_column=last_col)
        for row,label,value in [(2,'Style No.' if english else '款号',doc['styleCode']),
                                (3,'Material' if english else '材质',(doc['materialEn'] if english else doc['material']) or '—'),
                                (4,'Updated' if english else '更新日期（北京时间）',doc['updatedDate'] or '—'),
                                (5,'Unit' if english else '整款单位',doc['unit'])]:
            label_cell=write(sheet,row,2,label);label_cell.font=Font(name='Arial',size=10,color='64748B')
            write(sheet,row,3,value);sheet.merge_cells(start_row=row,start_column=3,end_row=row,end_column=last_col)
        instructions = '仅填写测量值；空白将清空原值，0 表示零。单位在页面统一换算，请勿修改表格结构。'
        if doc['unit'] != saved_doc['unit']:
            instructions = f"仅填写测量值；空白清空，0 表示零。本表为 {doc['unit']}，导入自动换算为保存单位 {saved_doc['unit']}；未改数值保留原精度。"
        write(sheet,6,1,'SIZE CHART · 1 in = 2.54 cm' if english else instructions)
        sheet.merge_cells(start_row=6,start_column=1,end_row=6,end_column=last_col)
        sheet.row_dimensions[6].height=30;sheet.row_dimensions[7].height=8
        border=Border(left=Side(style='thin',color='D8E2EF'),right=Side(style='thin',color='D8E2EF'),top=Side(style='thin',color='D8E2EF'),bottom=Side(style='thin',color='D8E2EF'))
        for col, label in enumerate(headers, 1):
            cell = write(sheet, 8, col, label); cell.fill = PatternFill('solid', fgColor='DFE8F5'); cell.font = Font(name='Arial',size=11,bold=True,color='243850')
            cell.border=border;cell.alignment=Alignment(horizontal='left' if col<start-1 else 'center',vertical='center',wrap_text=True)
        sheet.row_dimensions[8].height=28
        for row, d in enumerate(dims, 9):
            labels = [d['nameEn'], d['unit']] if english else [d['name'], d['nameEn'], d['unit']]
            for col, label in enumerate(labels, 1): write(sheet, row, col, label)
            for col, size in enumerate(doc['sizes'], start):
                val = doc['values'][d['id']].get(size)
                write(sheet, row, col, Decimal(val) if val is not None else None)
            for col in range(1,len(headers)+1):
                cell=sheet.cell(row,col);cell.border=border
                cell.alignment=Alignment(horizontal='left' if col<start-1 else 'center',vertical='center',wrap_text=True)
                if col < start: cell.fill=PatternFill('solid',fgColor='F3F6FA')
                else:
                    cell.number_format='0.00'
                    cell.font=Font(name='Arial',size=11,color='243850' if english else '2457A6')
            sheet.row_dimensions[row].height = 26
        if doc['image']:
            # The URL must still belong to this style; reuse existing bounded, allowlisted image reader.
            if doc['image'] not in doc['images']: raise ValueError(doc['styleCode'] + '：代表图已从商品移除，请重新选择')
            import app
            raw = app.fetch_image_bytes(doc['image'])
            img = app.build_excel_image(raw, width=144, height=144) if raw else None
            if not img: raise ValueError(doc['styleCode'] + '：图片读取失败，请重试导出')
            img.anchor=OneCellAnchor(_from=AnchorMarker(col=0,row=0,colOff=pixels_to_EMU((180-img.width)/2),rowOff=pixels_to_EMU((168-img.height)/2)),ext=XDRPositiveSize2D(pixels_to_EMU(img.width),pixels_to_EMU(img.height)))
            sheet.add_image(img)
        else: write(sheet, 1, 1, 'No image' if english else '未设置代表图')
        sheet.freeze_panes = None
        sheet.sheet_view.showGridLines = True
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = 'landscape'; sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1; sheet.page_setup.fitToHeight = 1
        sheet.print_options.horizontalCentered = True
        sheet.print_area = f'A1:{get_column_letter(last_col)}{8+len(dims)}'
        if not english:
            meta = {'sheet': sheet.title, 'layout': 'compact-v2', **structure(saved_doc)}
            if doc['unit'] != saved_doc['unit']: meta['exportUnit'] = doc['unit']
            if len(packed(meta)) > 30000: raise ValueError('该款尺码表结构过大，请减少维度后导出')
            metadata.append(meta)
    if mode == 'maintenance':
        meta_sheet = book.create_sheet(META); meta_sheet.sheet_state = 'veryHidden'
        write(meta_sheet, 1, 1, VERSION)
        for row, item in enumerate(metadata, 2): write(meta_sheet, row, 1, packed(item)); write(meta_sheet, row, 2, sign(item))
    stream = BytesIO(); book.save(stream)
    if len(stream.getvalue()) > MAX_BYTES: raise ValueError('导出文件超过 10 MB，请减少款号后分批导出；未截断内容')
    stream.seek(0)
    return send_file(stream, as_attachment=True, download_name='size-charts-english.xlsx' if mode == 'english' else '商品尺码表.xlsx', mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


def parse(raw):
    if len(raw) > MAX_BYTES: raise ValueError('文件不能超过 10 MB')
    try:
        with ZipFile(BytesIO(raw)) as archive:
            if len(archive.infolist()) > 2000 or sum(i.file_size for i in archive.infolist()) > MAX_EXPANDED: raise ValueError('Excel 解压内容超过限制')
            if any(n.endswith('vbaProject.bin') for n in archive.namelist()): raise ValueError('不接受含宏文件')
        book = load_workbook(BytesIO(raw), read_only=True, data_only=False, keep_links=False)
    except ValueError: raise
    except Exception as exc: raise ValueError('请上传有效的系统 .xlsx 完整维护表') from exc
    try:
        if META not in book.sheetnames or book[META]['A1'].value not in {VERSION,LEGACY_VERSION}: raise ValueError('仅接受系统导出的完整维护表；英文客户表不支持导入')
        compact=book[META]['A1'].value==VERSION
        meta = book[META]
        if meta.max_row > charts.MAX_STYLES+1 or meta.max_column != 2: raise ValueError('模板元数据无效')
        mapping = {}; codes = set()
        for cells in meta.iter_rows(min_row=2, max_col=2):
            try:
                item = json.loads(cells[0].value)
                if any(c.data_type == 'f' for c in cells) or not hmac.compare_digest(str(cells[1].value), sign(item)): raise ValueError()
                if item['sheet'] in mapping or item['styleCode'] in codes: raise ValueError()
                if (compact and item.get('layout')!='compact-v2') or (not compact and 'layout' in item): raise ValueError()
                mapping[item['sheet']] = item; codes.add(item['styleCode'])
            except (ValueError, TypeError, KeyError): raise ValueError('模板元数据已损坏，请重新导出')
        rows, changes, errors = [], [], []
        cats = charts.catalog(list(codes))
        for sheet in book:
            if sheet.title == META: continue
            if sheet.title not in mapping: raise ValueError('存在未知或重复工作表，请保留原工作表名称')
            item = mapping[sheet.title]; code = item['styleCode']
            try:
                current = charts.detail(code, cats)
                if structure(current) != {k: item[k] for k in structure(current)}: raise workbench.Conflict('尺码表或商品尺码已变化，请重新导出')
                export_unit = item.get('exportUnit', current['unit'])
                if export_unit not in {'cm', 'in'}: raise ValueError('导出单位元数据无效')
                dims, sizes = item['dimensions'], item['sizes']
                first=1 if compact else 4
                max_row=len(dims)+8 if compact else max(len(dims)+9,22)
                max_col=max(3+len(sizes),5) if compact else max(6+len(sizes),8)
                if sheet.max_row > max_row or sheet.max_column > max_col: raise ValueError('存在额外行列')
                # Images are not loaded by the read-only parser; cap cell traversal explicitly.
                matrix = list(sheet.iter_rows(max_row=max_row, max_col=max_col))
                def cell_at(row, col):
                    if row > len(matrix) or col > len(matrix[row-1]): raise ValueError('必要行列已删除，请重新导出完整维护表')
                    return matrix[row-1][col-1]
                for row in matrix:
                    for cell in row:
                        if cell.data_type == 'f': raise ValueError(f'{cell.coordinate}：不接受公式')
                if cell_at(2,3 if compact else 2).value != code: raise ValueError('款号已改变')
                if compact and cell_at(5,3).value != export_unit: raise ValueError('单位已改变，请在页面切换单位后重新导出')
                expected = ['测量维度', 'English', '单位', *sizes]
                if [cell_at(8, i).value for i in range(first,first+3+len(sizes))] != expected: raise ValueError('尺码列或表头已改变')
                if not compact and cell_at(9+len(dims),4).value != 'END': raise ValueError('维度行已删除或移动')
                candidate = deepcopy(current)
                for row_index, d in enumerate(dims, 9):
                    if [cell_at(row_index,i).value or '' for i in range(first,first+3)] != [d['name'],d['nameEn'],export_unit]: raise ValueError(f'第 {row_index} 行：维度或单位已改变')
                    for col, size in enumerate(sizes, first+3):
                        cell = cell_at(row_index,col)
                        try: new = charts.number(cell.value)
                        except ValueError as exc: raise ValueError(f'{cell.coordinate}：{exc}')
                        old = current['values'][d['id']].get(size)
                        # Untouched rounded previews must not alter the stored precision.
                        expected = charts.convert_value(old, current['unit'], export_unit)
                        if new == expected: new = old
                        else:
                            try: new = charts.convert_value(new, export_unit, current['unit'])
                            except ValueError as exc: raise ValueError(f'{cell.coordinate}：{exc}')
                        candidate['values'][d['id']][size] = new
                        if old != new: rows.append({'styleCode': code, 'dimension': d['name'], 'size': size, 'cell': cell.coordinate, 'unit': current['unit'], 'old': old, 'new': new, 'cleared': new is None})
                changes.append(candidate)
            except (ValueError, workbench.Conflict) as exc: errors.append({'styleCode': code, 'sheet': sheet.title, 'message': str(exc)})
        if not changes and not errors: raise ValueError('文件中没有完整款号工作表')
        payload = {'changes': changes, 'actorId': g.current_user['id']}
        return {'rows': rows, 'errors': errors, 'styleCount': len(changes), 'changedCount': len(rows), 'clearCount': sum(r['cleared'] for r in rows),
                'proposal': payload if not errors else None, 'previewToken': sign(payload) if not errors else None}
    finally: book.close()


def preview_upload():
    upload = request.files.get('file')
    if not upload or not upload.filename.lower().endswith('.xlsx'): raise ValueError('请选择系统导出的 .xlsx 完整维护表')
    return parse(upload.stream.read(MAX_BYTES+1))


def confirm(payload):
    if not isinstance(payload, dict) or not isinstance(payload.get('proposal'), dict): raise ValueError('缺少导入预览')
    proposal = payload['proposal']
    if proposal.get('actorId') != g.current_user['id'] or not hmac.compare_digest(str(payload.get('previewToken', '')), sign(proposal)):
        raise ValueError('导入预览无效，请重新上传预览')
    return charts.save_batch({'requestId': payload.get('requestId'), 'changes': proposal.get('changes')})
