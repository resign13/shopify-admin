"""Filtered style performance XLSX; shares the dashboard query and access scope."""
from datetime import datetime
from io import BytesIO

from flask import send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import safe_excel
import sales_ownership
import workbench


RISK_NAMES = {
    'backordered': '欠货', 'out_of_stock': '无库存', 'no_sales': '周转期无销量',
    'critical': '7 天内售罄', 'warning': '7–30 天售罄', 'healthy': '库存健康',
}
BASE_COLUMNS = [
    ('styleCode', '款号', 24, '@'), ('name', '商品名称', 48, '@'),
    ('units', '销量（件）', 16, '#,##0'),
    ('amount', '商品金额（USD，不含运费）', 24, '#,##0.00'),
    ('colorSkuCount', '颜色 SKU 数', 16, '#,##0'),
]
STOCK_COLUMNS = [
    ('stock', '现货余额', 16, '#,##0'), ('availableStock', '可用现货', 16, '#,##0'),
    ('shortageUnits', '欠货件数', 16, '#,##0'), ('shortageSizeCount', '欠货尺码数', 16, '#,##0'),
    ('contractPending', '合同未送', 16, '#,##0'), ('pendingInspection', '待验货', 16, '#,##0'),
    ('pendingInbound', '合格', 16, '#,##0'), ('defectivePending', '待打回次品', 16, '#,##0'),
    ('temporaryInbound', '临时待入库', 16, '#,##0'),
]
HEADER_ROW = 1


def _local_time(value):
    if not value:
        return None
    if isinstance(value, str):
        value = datetime.fromisoformat(value.replace('Z', '+00:00'))
    return value.astimezone(workbench.SHANGHAI).replace(tzinfo=None)


def _days(item):
    if item['shortageUnits'] > 0:
        return '欠货待补'
    if item['availableStock'] <= 0:
        return '无库存'
    if not item['velocityUnits'] or item['estimatedDays'] is None:
        return '暂无销量'
    return item['estimatedDays']


def build_workbook(data):
    personal = data.get('scope', {}).get('personal', False)
    columns = [*BASE_COLUMNS, *([] if personal else STOCK_COLUMNS),
               ('averageDailyUnits', '日均销量（件）', 18, '#,##0.00')]
    if not personal:
        columns += [('estimatedDays', '预计可售（天）', 18, '0.0'), ('risk', '库存状态', 20, '@')]
    columns += [('lastSoldAt', '最近销售时间（北京时间）', 27, 'yyyy-mm-dd hh:mm:ss')]
    book = Workbook()
    sheet = book.active
    sheet.title = '款号商品表现'
    sheet.sheet_view.showGridLines = False
    sheet.append([label for _, label, _, _ in columns])
    for item in data['items']:
        values = {**item, 'lastSoldAt': _local_time(item.get('lastSoldAt'))}
        if not personal:
            values.update(estimatedDays=_days(item), risk=RISK_NAMES[item['risk']])
        sheet.append([values.get(key) for key, _, _, _ in columns])
    last_data_row = sheet.max_row
    # Only additive metrics are totalled. Turnover and per-style SKU counts are not sums.
    summary = data['summary']
    sheet.append(['合计' if key == 'styleCode' else f"{data['total']} 个款号" if key == 'name'
                  else summary.get(key) for key, _, _, _ in columns])
    total_row = sheet.max_row
    for col, (key, _, width, number_format) in enumerate(columns, 1):
        sheet.column_dimensions[get_column_letter(col)].width = width
        for row in range(HEADER_ROW, total_row + 1):
            cell = sheet.cell(row, col)
            heading = row in {HEADER_ROW, total_row}
            cell.font = Font(name='Microsoft YaHei', size=11, bold=heading,
                             color='FFFFFF' if row == HEADER_ROW else '1E293B')
            cell.alignment = Alignment(vertical='center', wrap_text=True,
                                       horizontal='left' if key in {'styleCode', 'name', 'risk'} else 'right')
            if row > HEADER_ROW:
                cell.number_format = number_format
                if key == 'styleCode':
                    cell.data_type = 's'
            if row == HEADER_ROW:
                cell.fill = PatternFill('solid', fgColor='243850')
            elif row == total_row or row % 2 == 0:
                cell.fill = PatternFill('solid', fgColor='EFF6FF' if row == total_row else 'F8FAFC')
    for row in range(HEADER_ROW, total_row + 1):
        sheet.row_dimensions[row].height = 40 if row == HEADER_ROW else 42
    sheet.freeze_panes = f'A{HEADER_ROW + 1}'
    sheet.auto_filter.ref = f'A{HEADER_ROW}:{get_column_letter(len(columns))}{last_data_row}'
    sheet.print_title_rows = f'1:{HEADER_ROW}'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    return safe_excel.prepare(book)


def export(args, user):
    # The route's existing snapshot_admin_call supplies one read-only MVCC snapshot.
    owner_id = sales_ownership.scope(args, user)
    scoped = {**args, '_ownerId': owner_id}
    data = sales_ownership.personalize(workbench.style_performance(scoped, paginate=False), owner_id)
    book = build_workbook(data)
    out = BytesIO()
    book.save(out)
    book.close()
    out.seek(0)
    stamp = datetime.now(workbench.SHANGHAI).strftime('%Y%m%d-%H%M%S')
    return send_file(out, as_attachment=True, download_name=f'款号商品表现-{stamp}.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
