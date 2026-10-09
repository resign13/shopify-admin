"""Cumulative contract tolerance regression. All writes use guarded *_test fixtures."""
import io
import json
from pathlib import Path
import subprocess
import sys
import uuid
import unittest
from concurrent.futures import ThreadPoolExecutor
from unittest.mock import patch
from openpyxl import load_workbook
import test_workbench as fixtures
import inventory_overdelivery as policy
db = fixtures.db


class OverdeliveryTest(unittest.TestCase):
    setUp = fixtures.WorkbenchTest.setUp
    call = fixtures.WorkbenchTest.call
    detail = fixtures.WorkbenchTest.detail
    product_payload = fixtures.WorkbenchTest.product_payload
    upload = fixtures.WorkbenchTest.upload

    def prepare(self, quantities=None):
        quantities = quantities or {'M':100}
        db._fetch_one('UPDATE product_size_prices SET contract_pending=0,pending_inspection=0,pending_inbound=0 WHERE product_id=1 RETURNING id')
        body={'requestId':str(uuid.uuid4()),'partyA':'采购','partyB':'工厂','contractNo':'EXTRA-TEST',
              'styleImage':'/uploads/fixture.jpg','sizeChartImage':'/uploads/fixture.jpg','deliveryDate':'2026-10-20',
              'items':[{'productId':1,'quantities':quantities}]}
        result=self.call('contracts','POST',body)
        self.assertEqual(result.status_code,200,result.json)
        return result.json['item']

    def size(self, code='M'): return next(s for s in self.detail()['sizePrices'] if s['sizeCode']==code)
    def save(self, **fields):
        return self.call('inventory/1','PUT',{'version':self.detail()['version'],'sizeStocks':{},**fields})
    def receive(self, source='normal', body=None):
        path='receive' if source=='normal' else 'temporary/receive' if source=='temporary' else 'defective/return'
        return self.call('inventory/1/'+path,'POST',body or {'version':self.detail()['version'],'requestId':str(uuid.uuid4())})

    def test_100_accepts_115_not_116_then_receives_stock(self):
        self.prepare(); stock=self.size()['stock']
        self.assertEqual(self.save(pendingInspectionBySize={'M':116}).status_code,400)
        r=self.save(pendingInspectionBySize={'M':115});self.assertEqual(r.status_code,200,r.json)
        s=self.size();self.assertEqual((s['contractPending'],s['pendingInspection'],s['overdeliveryUsed'],s['overdeliveryInspection']),(0,115,15,15))
        self.assertGreater(db._fetch_one("SELECT count(*) AS n FROM admin_audit_logs WHERE entity_table='inventory_overdelivery'")['n'],0)
        self.assertEqual(s['stock'],stock)
        self.assertEqual(self.save(pendingInboundBySize={'M':115}).status_code,200)
        self.assertEqual(self.size()['overdeliveryQualified'],15)
        body={'version':self.detail()['version'],'requestId':str(uuid.uuid4())}
        self.assertEqual(self.receive(body=body).status_code,200)
        self.assertEqual(self.size()['stock'],stock+115)
        self.assertEqual(self.size()['overdeliveryUsed'],15)
        self.assertEqual(self.size()['overdeliveryQualified'],0)
        self.assertTrue(self.receive(body=body).json['replayed'])

    def test_split_batches_share_original_allowance_after_receipt(self):
        self.prepare()
        for amount in [50,50]:
            self.assertEqual(self.save(pendingInspectionBySize={'M':amount}).status_code,200)
            self.assertEqual(self.save(pendingInboundBySize={'M':amount}).status_code,200)
            self.receive()
        self.assertEqual(self.size()['contractPending'],0)
        self.assertEqual(self.save(pendingInspectionBySize={'M':10}).status_code,200)
        self.save(pendingInboundBySize={'M':10});self.receive()
        self.assertEqual(self.save(pendingInspectionBySize={'M':6}).status_code,400)
        self.assertEqual(self.save(pendingInspectionBySize={'M':5}).status_code,200)
        self.assertEqual(self.size()['overdeliveryUsed'],15)

    def test_return_extra_first_never_replenishes_allowance(self):
        self.prepare();self.save(pendingInspectionBySize={'M':115})
        self.save(defectivePendingBySize={'M':20})
        before=self.size(); body={'version':self.detail()['version'],'requestId':str(uuid.uuid4())}
        result=self.receive('defective',body)
        self.assertEqual(result.status_code,200,result.json)
        s=self.size();self.assertEqual((s['contractPending'],s['pendingInspection'],s['overdeliveryUsed'],s['overdeliveryInspection']),(5,95,15,0))
        self.assertEqual(result.json['changes'][0]['extraReturned'],15)
        self.assertTrue(self.receive('defective',body).json['replayed'])
        self.assertEqual(s['stock'],before['stock'])
        self.assertEqual(self.save(pendingInspectionBySize={'M':101}).status_code,400)
        self.assertEqual(self.save(pendingInspectionBySize={'M':100}).status_code,200)

    def test_reserve_qualified_reverse_and_manual_correction(self):
        self.prepare();self.save(pendingInspectionBySize={'M':115});self.save(defectivePendingBySize={'M':10})
        self.assertEqual(self.save(pendingInboundBySize={'M':106}).status_code,400)
        self.assertEqual(self.save(pendingInboundBySize={'M':105}).status_code,200)
        self.assertEqual((self.size()['overdeliveryInspection'],self.size()['overdeliveryQualified']),(10,5))
        self.assertEqual(self.save(pendingInboundBySize={'M':0},pendingInspectionBySize={'M':0},defectivePendingBySize={'M':0}).status_code,200)
        self.assertEqual((self.size()['contractPending'],self.size()['overdeliveryUsed']),(100,15))
        self.assertEqual(self.save(pendingInspectionBySize={'M':101}).status_code,400)

    def test_rounding_size_isolation_and_no_contract_no_allowance(self):
        self.prepare({'M':6,'L':7})
        self.assertEqual(self.save(pendingInspectionBySize={'M':7}).status_code,400)
        self.assertEqual(self.save(pendingInspectionBySize={'L':8}).status_code,200)
        self.assertEqual(self.size('L')['overdeliveryUsed'],1)
        self.assertEqual(self.save(contractPendingBySize={'Tall XL':100}).status_code,200)
        self.assertEqual(self.size('Tall XL')['overdeliveryLimit'],0)
        self.assertEqual(self.save(pendingInspectionBySize={'Tall XL':101}).status_code,400)

    def test_multi_size_rollback_and_concurrency(self):
        self.prepare({'M':100,'L':100});before=self.detail()
        self.assertEqual(self.save(pendingInspectionBySize={'M':115,'L':116}).status_code,400)
        self.assertEqual(self.detail(),before)
        body={'version':before['version'],'sizeStocks':{},'pendingInspectionBySize':{'M':115}}
        with ThreadPoolExecutor(max_workers=2) as pool:
            rs=list(pool.map(lambda _:self.call('inventory/1','PUT',body),range(2)))
        self.assertEqual(sorted(r.status_code for r in rs),[200,409])
        self.assertEqual(self.size()['overdeliveryUsed'],15)

    def test_audit_failure_rolls_back_ledger_and_stock(self):
        self.prepare();before=self.detail()
        with patch('workbench.record_inventory_registration',side_effect=ValueError('audit failure')):
            self.assertEqual(self.save(pendingInspectionBySize={'M':115}).status_code,400)
        self.assertEqual(self.detail(),before)

    def test_contract_edit_cannot_shrink_spent_budget_and_increase_grants_only_delta(self):
        contract=self.prepare();self.save(pendingInspectionBySize={'M':115})
        self.save(pendingInspectionBySize={'M':0})
        body={**contract,'requestId':str(uuid.uuid4()),'revision':contract['revision']}
        body['items'][0]['quantities']['M']=90
        self.assertEqual(self.call('contracts/1','PUT',body).status_code,400)
        body['items'][0]['quantities']['M']=200
        r=self.call('contracts/1','PUT',body);self.assertEqual(r.status_code,200,r.json)
        self.assertEqual((self.size()['originalContractQuantity'],self.size()['overdeliveryRemaining']),(200,15))

    def test_product_edits_preserve_ledger_and_customer_omits_fields(self):
        self.prepare();self.save(pendingInspectionBySize={'M':115})
        response=self.call('products/1','PUT',self.product_payload(1))
        self.assertEqual(response.status_code,200,response.json)
        self.assertEqual(self.size()['overdeliveryUsed'],15)
        for route in ['inventory/1','inventory?page=1']:
            public=self.call(route,role='customer')
            self.assertNotIn('overdelivery',str(public.json))
            self.assertNotIn('originalContractQuantity',str(public.json))
        self.save(pendingInspectionBySize={'M':0})
        payload=self.product_payload(1);payload['sizes']=['L','Tall XL'];payload['sizePrices']=[s for s in payload['sizePrices'] if s['sizeCode']!='M']
        self.assertEqual(self.call('products/1','PUT',payload).status_code,400)

    def test_excel_v6_and_legacy_cumulative_limit(self):
        self.prepare()
        raw=self.call('inventory/export').data
        book=load_workbook(io.BytesIO(raw)); sheet=book.active
        self.assertEqual(sheet['A2'].value,'inventory-v6');self.assertEqual(sheet['T2'].value,100)
        sheet['O2']=115;sheet['T2']=999999;sheet['U2']=999999
        out=io.BytesIO();book.save(out)
        r=self.upload('preview',out.getvalue());self.assertEqual(r.status_code,200,r.json)
        self.assertEqual(r.json['errors'],[])
        r=self.upload('confirm',out.getvalue(),r.json['fileHash'])
        self.assertEqual(r.status_code,200,r.json);self.assertEqual(self.size()['overdeliveryUsed'],15)
        book=load_workbook(io.BytesIO(self.call('inventory/export').data));sheet=book.active
        sheet['O2']=116;sheet.delete_cols(20,5)
        for row in range(2,sheet.max_row+1):sheet.cell(row,1,'inventory-v5')
        out=io.BytesIO();book.save(out)
        r=self.upload('preview',out.getvalue());self.assertTrue(r.json['errors'])

    def test_repeated_migration_preserves_ledger(self):
        self.prepare();self.save(pendingInspectionBySize={'M':115})
        before=self.size()
        with db.get_connection() as conn,conn.cursor() as cur:
            policy.migrate(cur);policy.migrate(cur)
        self.assertEqual(self.size(),before)

    def test_manual_pending_balance_does_not_inflate_original_allowance(self):
        self.prepare();self.save(pendingInspectionBySize={'M':100})
        self.save(pendingInboundBySize={'M':100});self.receive()
        self.save(contractPendingBySize={'M':100})
        self.assertEqual(self.save(pendingInspectionBySize={'M':16}).status_code,400)
        self.assertEqual(self.save(pendingInspectionBySize={'M':15}).status_code,200)
        self.assertEqual(self.size()['originalContractQuantity'],100)
        self.assertEqual(self.size()['overdeliveryUsed'],15)

    def test_empty_product_deletion_preserves_cumulative_history(self):
        self.prepare();self.save(pendingInspectionBySize={'M':115})
        self.save(pendingInspectionBySize={'M':0},sizeStocks={'M':0,'L':0,'Tall XL':0})
        self.assertEqual(self.save(contractPendingBySize={'M':0}).status_code,200)
        db._fetch_one('DELETE FROM order_items WHERE product_id=1 RETURNING id')
        result=self.call('products/1','DELETE')
        self.assertEqual(result.status_code,200,result.json)
        self.assertFalse(db._fetch_one('SELECT is_active FROM products WHERE id=1')['is_active'])
        self.assertEqual(db._fetch_one("SELECT used FROM inventory_overdelivery WHERE product_id=1 AND size_code='M'")['used'],15)
        self.assertEqual(self.size()['overdeliveryUsed'],15)

    def test_release_probe_is_read_only_with_populated_ledger(self):
        self.prepare();self.save(pendingInspectionBySize={'M':115})
        before=self.detail()
        audit=db._fetch_one('SELECT count(*) AS n FROM admin_audit_logs')['n']
        probe=Path(__file__).resolve().parents[1]/'scripts/verify-inventory-overdelivery.py'
        result=subprocess.run([sys.executable,str(probe),'admin-backend'],capture_output=True,text=True)
        self.assertEqual(result.returncode,0,result.stderr)
        self.assertTrue(json.loads(result.stdout)['readOnly'])
        self.assertEqual(self.detail(),before)
        self.assertEqual(db._fetch_one('SELECT count(*) AS n FROM admin_audit_logs')['n'],audit)


class ReachablePolicyTest(unittest.TestCase):
    def test_all_small_reachable_states_conserve_normal_and_extra_sources(self):
        # Exhaustive real transitions, not a second copy of the policy algorithm.
        # Includes corrections, qualified reversals, receipts and defect returns.
        names=('contract_pending','pending_inspection','pending_inbound','overdeliveryUsed',
               'overdeliveryInspection','overdeliveryQualified','contractReceived')
        start=(7,0,0,0,0,0,0);pending=[start];seen={start};transitions=0
        def accept(values, previous):
            nonlocal transitions
            p,i,b,u,ei,eb,received=values
            self.assertTrue(all(v>=0 for v in values),values)
            self.assertLessEqual(ei,i);self.assertLessEqual(eb,b)
            self.assertLessEqual(ei+eb,u);self.assertLessEqual(u,1)
            self.assertGreaterEqual(u,previous[3])
            self.assertEqual(p+received,7)
            self.assertLessEqual(i+b-ei-eb,received)
            transitions+=1
            if values not in seen:seen.add(values);pending.append(values)
        while pending:
            previous=pending.pop();row=dict(zip(names,previous))
            row.update(originalContractQuantity=7,overdeliveryLimit=1)
            for ni in range(9):
                for nb in range(9):
                    try:p,i,b,state=policy.transition(row,{'pendingInspection':ni,'pendingInbound':nb})
                    except ValueError:continue
                    accept((p,i,b,state['overdeliveryUsed'],state['overdeliveryInspection'],state['overdeliveryQualified'],state['contractReceived']),previous)
            if row['pending_inbound']:
                values=list(previous);values[2]=0;values[5]=0
                accept(tuple(values),previous)
            for quantity in range(1,row['pending_inspection']+1):
                state=policy.return_preview(row,quantity)
                accept((state['contractAfter'],state['inspectionAfter'],row['pending_inbound'],
                        state['overdeliveryUsed'],state['overdeliveryInspection'],state['overdeliveryQualified'],state['contractReceived']),previous)
        self.assertGreater(len(seen),100);self.assertGreater(transitions,8000)


if __name__=='__main__': unittest.main()
