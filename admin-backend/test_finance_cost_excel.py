"""Focused cost workbook lifecycle; fixtures are guarded local *_test databases."""
from datetime import datetime
from io import BytesIO
import unittest
import uuid
from unittest.mock import patch

from flask.testing import FlaskClient
from openpyxl import Workbook, load_workbook

import test_finance as fixtures
import finance
import finance_cost_excel as excel
from test_workbench import db, application


class CostExcelTest(unittest.TestCase):
    setUp = fixtures.FinanceTest.setUp
    tearDown = fixtures.FinanceTest.tearDown
    clear = fixtures.FinanceTest.clear
    call = fixtures.FinanceTest.call
    payload = fixtures.FinanceTest.payload
    stock = fixtures.FinanceTest.stock
    post = fixtures.FinanceTest.post
    change = fixtures.FinanceTest.change
    prepared = fixtures.FinanceTest.prepared
    save_cost = fixtures.FinanceTest.save_cost
    query = fixtures.FinanceTest.query

    def book(self, **filters):
        response = self.post('costs/export', filters)
        self.assertEqual(response.status_code, 200, response.json if response.is_json else '')
        self.assertEqual(response.headers['Cache-Control'], 'no-store')
        return load_workbook(BytesIO(response.data))

    def upload(self, book, role='finance', filename='costs.xlsx'):
        stream = BytesIO()
        book.save(stream)
        stream.seek(0)
        return application.app.test_client().post('/api/admin/finance/costs/import/preview',
            data={'file': (stream, filename)}, headers={'Authorization': 'Bearer ' + self.tokens[role]})

    def commit(self, proposal, **extra):
        return self.post('costs/save', {**proposal, 'requestId': str(uuid.uuid4()), **extra})

    def add_style(self, code):
        product = db.create_product({'categoryKey': 'outerwear', 'productCode': code, 'sku': code + '-black',
            'slug': 'excel-' + uuid.uuid4().hex, 'styleCode': code, 'title': 'Not in workbook', 'image': '',
            'sizes': ['M'], 'sizePrices': [{'sizeCode': 'M', 'price': 9, 'stock': 12}]})
        db._fetch_one('UPDATE products SET style_code=%s WHERE id=%s RETURNING id', (code, product['id']))
        return product

    def test_export_all_filtered_styles_text_beijing_zero_and_missing(self):
        for n in range(24):
            self.add_style(f'00{n:03}')
        self.add_style('=literal-style')
        self.save_cost(cost='0', at='2026-10-21T02:30:00+00:00')
        data = self.book(page=2, pageSize=20)
        self.assertEqual(data[excel.SHEET].max_row, 27)
        self.assertEqual(data[excel.META].sheet_state, 'veryHidden')
        self.assertEqual([c.value for c in data[excel.SHEET][1]], excel.HEADERS)
        rows = {r[0].value: r for r in list(data[excel.SHEET])[1:]}
        self.assertEqual(rows['GT-2026'][1].value, datetime(2026, 10, 21, 10, 30))
        self.assertEqual(rows['GT-2026'][2].value, 0)
        self.assertIsNone(rows['GT-2026'][3].value)
        self.assertEqual(rows['00000'][0].data_type, 's')
        self.assertEqual(rows['00000'][0].number_format, '@')
        self.assertIsNone(rows['00000'][1].value)
        self.assertEqual(rows['00000'][2].value, '未设置')
        self.assertEqual(rows['=literal-style'][0].data_type, 's')
        self.assertEqual(self.upload(data).json['summary']['blank'], 26)
        self.assertEqual(self.book(category='outerwear')[excel.SHEET].max_row, 26)
        self.assertEqual(self.book(keyword='GT-2026-1')[excel.SHEET].max_row, 2)
        self.assertEqual(self.book(missingOnly=True)[excel.SHEET].max_row, 26)
        self.assertEqual(self.book(styleCodes='00000,00001')[excel.SHEET].max_row, 3)
        with patch.object(excel, 'MAX_STYLES', 20):
            self.assertEqual(self.post('costs/export', {}).status_code, 400)

    def test_preview_blank_same_zero_reference_edits_and_cancel_are_read_only(self):
        self.add_style('OTHER')
        self.add_style('MISSING')
        self.save_cost(cost='12')
        data = self.book()
        rows = {r[0].value: r[0].row for r in list(data[excel.SHEET])[1:]}
        sheet = data[excel.SHEET]
        sheet.cell(rows['GT-2026'], 4, '12.00')
        sheet.cell(rows['OTHER'], 4, 0)
        sheet.cell(rows['OTHER'], 3, 9999)  # Reference cells never supply old values.
        sheet.cell(rows['OTHER'], 2, 'changed reference time')
        result = self.upload(data)
        self.assertEqual(result.status_code, 200, result.json)
        proposal = result.json
        self.assertEqual(proposal['summary'], {'update': 1, 'blank': 1, 'same': 1, 'error': 0})
        self.assertEqual(proposal['changes'][0]['costCny'], '0.00')
        self.assertIsNone(proposal['changes'][0]['oldCostCny'])
        self.assertEqual(db._fetch_one('SELECT count(*) AS n FROM finance_style_cost_audit')['n'], 1)
        self.assertIsNone(db._fetch_one("SELECT cost_cny FROM finance_style_costs WHERE style_code='OTHER'"))

    def test_import_uses_existing_save_timeline_idempotency_and_one_batch_timestamp(self):
        self.add_style('OTHER')
        before = self.stock()
        data = self.book()
        for row in data[excel.SHEET].iter_rows(min_row=2):
            row[3].value = 12
        proposal = self.upload(data).json
        self.assertEqual(proposal['summary']['update'], 2)
        self.assertEqual(proposal['affectedOrders'], 3)
        rid = str(uuid.uuid4())
        with patch.object(finance, 'cost_save_time', return_value=datetime.fromisoformat('2026-09-18T09:00:00+08:00')):
            result = self.commit(proposal, requestId=rid)
        self.assertEqual(result.status_code, 200, result.json)
        self.assertEqual(self.query()['summary']['totalCostCny'], '108.00')
        self.assertEqual(db._fetch_one('SELECT count(DISTINCT effective_at) AS n FROM finance_style_cost_audit')['n'], 1)
        replay = self.commit(proposal, requestId=rid)
        self.assertEqual(replay.status_code, 200, replay.json)
        self.assertTrue(replay.json['replayed'])
        next_book = self.book()
        for row in next_book[excel.SHEET].iter_rows(min_row=2):
            row[3].value = 18
        next_proposal = self.upload(next_book).json
        self.assertEqual(next_proposal['affectedOrders'], 0)
        with patch.object(finance, 'cost_save_time', return_value=datetime.fromisoformat('2026-10-21T10:00:00+08:00')):
            self.assertEqual(self.commit(next_proposal).status_code, 200)
        self.assertEqual(self.query()['summary']['totalCostCny'], '108.00')
        self.assertEqual(self.stock(), before)
        self.assertEqual(self.call('finance/costs/history?styleCode=GT-2026', role='finance').json['total'], 2)

    def test_conflicts_before_preview_and_after_preview_prevent_the_entire_batch(self):
        self.add_style('OTHER')
        data = self.book()
        for row in data[excel.SHEET].iter_rows(min_row=2):
            row[3].value = 18
        proposal = self.upload(data).json
        self.save_cost(cost='12')
        conflict = self.upload(data)
        self.assertEqual(conflict.status_code, 200, conflict.json)
        self.assertEqual(conflict.json['summary']['error'], 1)
        self.assertIn('版本冲突', next(r for r in conflict.json['rows'] if r['status'] == 'error')['message'])
        self.assertFalse(conflict.json['changes'])
        self.assertIsNone(conflict.json['previewHash'])
        self.assertEqual(self.commit(proposal).status_code, 409)
        self.assertIsNone(db._fetch_one("SELECT * FROM finance_style_costs WHERE style_code='OTHER'"))
        self.assertEqual(db._fetch_one('SELECT count(*) AS n FROM finance_style_cost_audit')['n'], 1)
        # Blank rows aren't updates and need not match an obsolete cost revision.
        for row in data[excel.SHEET].iter_rows(min_row=2):
            if row[0].value == 'GT-2026':
                row[3].value = None
        self.assertEqual(self.upload(data).json['summary']['error'], 0)

    def test_invalid_money_formula_duplicates_missing_styles_and_hidden_version(self):
        for value in (-1, '1.001', '=2+3', 'not money', True, 'NaN', '10000000000'):
            with self.subTest(value=value):
                book = self.book()
                book[excel.SHEET]['D2'] = value
                response = self.upload(book)
                self.assertEqual(response.status_code, 200, response.json)
                self.assertEqual(response.json['summary']['error'], 1)
                self.assertEqual(response.json['rows'][0]['rowNumber'], 2)
                self.assertFalse(response.json['changes'])
        data = self.book()
        data[excel.SHEET].append(['GT-2026', None, None, 13])
        self.assertEqual(self.upload(data).json['summary']['error'], 2)
        data = self.book()
        data[excel.SHEET]['A2'] = 'DOES-NOT-EXIST'
        self.assertEqual(self.upload(data).json['summary']['error'], 1)
        data = self.book()
        data[excel.META]['B3'] = '123'
        self.assertEqual(self.upload(data).status_code, 400)
        data = self.book()
        data[excel.META]['A1'] = 'old-template'
        self.assertEqual(self.upload(data).status_code, 400)
        self.assertEqual(self.upload(Workbook()).status_code, 400)
        self.assertEqual(self.upload(self.book(), filename='costs.xlsm').status_code, 400)
        with patch.object(excel, 'MAX_BYTES', 100):
            self.assertEqual(self.upload(self.book()).status_code, 400)
        with patch.object(excel, 'MAX_EXPANDED_BYTES', 100):
            self.assertEqual(self.upload(self.book()).status_code, 400)
        self.assertEqual(db._fetch_one('SELECT count(*) AS n FROM finance_style_cost_audit')['n'], 0)

    def test_removed_catalog_style_is_a_row_error(self):
        product = self.add_style('REMOVED')
        data = self.book(styleCodes='REMOVED')
        data[excel.SHEET]['D2'] = 8
        db._fetch_one('DELETE FROM products WHERE id=%s RETURNING id', (product['id'],))
        result = self.upload(data)
        self.assertEqual(result.status_code, 200, result.json)
        self.assertIn('商品已移除', result.json['rows'][0]['message'])
        self.assertFalse(result.json['changes'])

    def test_auth_csrf_and_export_upload_quota_boundaries(self):
        data = self.book()
        for role in ('sales', 'warehouse', 'customer'):
            self.assertEqual(self.post('costs/export', {}, role).status_code, 403)
            self.assertEqual(self.upload(data, role).status_code, 403)
        raw = FlaskClient(application.app, application.app.response_class)
        self.assertEqual(raw.post('/api/admin/finance/costs/export', json={}).status_code, 401)
        raw.set_cookie(application.security.cookie, self.tokens['finance'])
        self.assertEqual(raw.post('/api/admin/finance/costs/export', json={}, headers={'Origin': 'http://localhost'}).status_code, 403)
        security = application.app.extensions['security']
        with patch.object(security, 'reserve', return_value=7) as reserve:
            rejected = self.upload(data)
            self.assertEqual(rejected.status_code, 429)
            self.assertEqual(rejected.headers['Retry-After'], '7')
            self.assertEqual(reserve.call_args.args[0], 'upload')
        with patch.object(security, 'reserve', return_value=7) as reserve:
            self.assertEqual(self.post('costs/export', {}).status_code, 429)
            self.assertEqual(reserve.call_args.args[0], 'export')


if __name__ == '__main__':
    unittest.main()
