"""Order quotes, FX and inventory invariants on the guarded local *_test DB."""
import copy
from datetime import datetime
from decimal import Decimal
from io import BytesIO
import unittest
import uuid
from openpyxl import load_workbook
from test_workbench import seed, application, db
import order_pricing


class OrderPricingTest(unittest.TestCase):
    def setUp(self):
        self.tokens, _ = seed()
        application.app.config['TESTING'] = True
        db._fetch_one('DELETE FROM finance_settings RETURNING id')
        db._fetch_one('INSERT INTO finance_settings(id,exchange_rate) VALUES(1,7.2) RETURNING id')

    def call(self, path, method='GET', body=None, role='admin'):
        with application.app.test_client() as client:
            return client.open('/api/admin/'+path, method=method, json=body,
                               headers={'Authorization':'Bearer '+self.tokens[role]})

    def payload(self, **values):
        return {'requestId':str(uuid.uuid4()),'userId':1,'contactName':'Quoted Customer','phone':'123456',
            'country':'CN','address':'Quote Street','shippingFee':10,
            'currency':'USD','exchangeRate':None,
            'items':[{'productId':1,'sizeCode':'M','quantity':2,'unitPrice':80,'priceSource':'manual'}],**values}

    def create(self, **values):
        response=self.call('orders','POST',self.payload(**values))
        self.assertEqual(response.status_code,200,response.json)
        return self.call(f"orders/{response.json['order']['id']}").json['order']

    def historical_cny(self, **values):
        # Simulate a persisted record from the retained feature version. New
        # API orders are USD-only, but saved money must stay interpretable.
        order = self.create(**values)
        db._fetch_one("UPDATE orders SET currency='CNY',exchange_rate=7.2 WHERE id=%s RETURNING id",(order['id'],))
        return self.call(f"orders/{order['id']}").json['order']

    def test_native_price_edit_does_not_change_stock_and_clears_link(self):
        order=self.historical_cny()
        self.assertEqual(order['currency'],'CNY');self.assertEqual(order['totalAmount'],170)
        stock=db._fetch_one("SELECT stock FROM product_size_prices WHERE product_id=1 AND size_code='M'")['stock']
        db._fetch_one("UPDATE orders SET payment_link='https://pay.example.test/old' WHERE id=%s RETURNING id",(order['id'],))
        body=self.call(f"orders/{order['id']}").json['order'];body['items'][0]['unitPrice']=75
        response=self.call(f"orders/{order['id']}/details",'PUT',body)
        self.assertEqual(response.status_code,200,response.json)
        new=response.json['order'];self.assertEqual(new['totalAmount'],160);self.assertEqual(new['paymentLink'],'')
        self.assertEqual(new['items'][0]['unitPrice'],75)
        self.assertEqual(db._fetch_one("SELECT stock FROM product_size_prices WHERE product_id=1 AND size_code='M'")['stock'],stock)
        history=self.call(f"orders/{order['id']}/operations").json['items'][0]
        self.assertEqual(history['beforeCurrency'],'CNY');self.assertEqual(history['afterCurrency'],'CNY')
        self.assertTrue(any(c['field']=='unitPrice' for c in history['changes']))

    def test_catalog_snapshot_conflict_and_manual_price_allowed(self):
        body=self.payload(items=[{'productId':1,'sizeCode':'M','quantity':1,'unitPrice':29.9,'priceSource':'catalog','referencePriceUsd':29.9}])
        self.assertEqual(self.call('orders','POST',body).status_code,200)
        db._fetch_one("UPDATE product_size_prices SET price=30 WHERE product_id=1 AND size_code='M' RETURNING id")
        body['requestId']=str(uuid.uuid4())
        before=db._fetch_one('SELECT COUNT(*) AS n FROM orders')['n']
        response=self.call('orders','POST',body)
        self.assertEqual(response.status_code,409,response.json)
        self.assertEqual(response.json['changes'][0]['unitPrice'],30)
        self.assertEqual(response.json['changes'][0]['referencePriceUsd'],30)
        self.assertEqual(db._fetch_one('SELECT COUNT(*) AS n FROM orders')['n'],before)
        body['items'][0]['priceSource']='manual'
        self.assertEqual(self.call('orders','POST',body).status_code,200)

    def test_draft_preserves_manual_prices_zero_and_stale_catalog(self):
        body=self.payload(items=[{'productId':1,'sizeCode':'M','quantity':2,'unitPrice':80,'priceSource':'manual'},
            {'productId':1,'sizeCode':'L','quantity':0,'unitPrice':0,'priceSource':'manual'},
            {'productId':2,'sizeCode':'M','quantity':1,'unitPrice':29.9,'priceSource':'catalog','referencePriceUsd':29.9}])
        result=self.call('orders/drafts','POST',{'draftKey':str(uuid.uuid4()),'draft':body})
        self.assertEqual(result.status_code,200,result.json)
        db._fetch_one("UPDATE product_size_prices SET price=33 WHERE size_code='M' RETURNING id")
        resumed=self.call(f"orders/drafts/{result.json['id']}").json
        self.assertEqual(resumed['draft']['currency'],'USD');self.assertIsNone(resumed['draft']['exchangeRate'])
        self.assertEqual([r['unitPrice'] for r in resumed['draft']['items']],[80,0,29.9])
        self.assertTrue(resumed['draft']['items'][2]['catalogPriceChanged'])
        submit={**resumed['draft'],'draftId':resumed['id'],'draftVersion':resumed['version'],'requestId':resumed['submitRequestId']}
        submit['items']=[r for r in submit['items'] if r['quantity']>0]
        self.assertEqual(self.call('orders','POST',submit).status_code,409)
        submit['items'][1]['priceSource']='manual'
        response=self.call('orders','POST',submit)
        self.assertEqual(response.status_code,200,response.json)

    def test_paid_currency_rate_lock_and_price_edit(self):
        order=self.historical_cny()
        body=self.call(f"orders/{order['id']}").json['order'];body['status']='paid'
        self.assertEqual(self.call(f"orders/{order['id']}/details",'PUT',body).status_code,200)
        current=self.call(f"orders/{order['id']}").json['order']
        for extra in [{'currency':'USD'},{'exchangeRate':7}]:
            bad={**copy.deepcopy(current),**extra}
            self.assertEqual(self.call(f"orders/{order['id']}/details",'PUT',bad).status_code,400)
        current['items'][0]['unitPrice']=90
        response=self.call(f"orders/{order['id']}/details",'PUT',current)
        self.assertEqual(response.status_code,200,response.json)
        self.assertEqual(response.json['order']['status'],'paid')
        self.assertEqual(response.json['order']['totalAmount'],190)
        for state in ['shipped','completed']:
            db._fetch_one('UPDATE orders SET status=%s WHERE id=%s RETURNING id',(state,order['id']))
            body=self.call(f"orders/{order['id']}").json['order']
            body['items'][0]['unitPrice'] += 1
            body['trackingNo']='PRICE-QA-TRACK'
            response=self.call(f"orders/{order['id']}/details",'PUT',body)
            self.assertEqual(response.status_code,200,response.json)
            self.assertEqual(response.json['order']['status'],state)
        db._fetch_one("UPDATE orders SET status='cancelled' WHERE id=%s RETURNING id",(order['id'],))
        bad=self.call(f"orders/{order['id']}").json['order'];bad['items'][0]['unitPrice'] += 1
        self.assertEqual(self.call(f"orders/{order['id']}/details",'PUT',bad).status_code,400)

    def test_reorder_is_usd_current_catalog_with_saved_freight_conversion(self):
        order=self.historical_cny()
        db._fetch_one('UPDATE finance_settings SET exchange_rate=7 WHERE id=1 RETURNING id')
        db._fetch_one("UPDATE product_size_prices SET price=30 WHERE product_id=1 AND size_code='M' RETURNING id")
        response=self.call(f"orders/{order['id']}/reorder-draft")
        self.assertEqual(response.status_code,200,response.json)
        draft=response.json['draft']
        self.assertEqual(draft['currency'],'USD');self.assertIsNone(draft['exchangeRate']);self.assertEqual(draft['shippingFee'],1.39)
        self.assertEqual(draft['items'][0]['unitPrice'],30);self.assertEqual(draft['items'][0]['priceSource'],'catalog')
        draft.update(sourceOrderId=order['id'],requestId=str(uuid.uuid4()))
        draft['items'][0].update(unitPrice=200,priceSource='manual')
        self.assertEqual(self.call('orders','POST',draft).status_code,200)

    def test_legacy_usd_and_public_settings_permissions(self):
        db._fetch_one('DELETE FROM finance_settings RETURNING id')
        self.assertEqual(self.call('orders/pricing-settings').json['exchangeRate'],None)
        body=self.payload(currency='USD',exchangeRate=None)
        del body['currency'];del body['exchangeRate']
        response=self.call('orders','POST',body);self.assertEqual(response.status_code,200,response.json)
        self.assertEqual(response.json['order']['currency'],'USD')
        self.assertIsNone(response.json['order']['exchangeRate'])
        for role in ['warehouse','customer']:
            self.assertEqual(self.call('orders/pricing-settings',role=role).status_code,403)
        warehouse=self.call(f"orders/{response.json['order']['id']}",role='warehouse').json['order']
        for key in ['currency','exchangeRate','totalAmount']:self.assertNotIn(key,warehouse)
        for key in ['priceSource','referencePriceUsd','unitPrice']:self.assertNotIn(key,warehouse['items'][0])

    def test_mixed_currency_finance_and_workbench(self):
        self.historical_cny(shippingFee=0,items=[{'productId':1,'sizeCode':'M','quantity':2,'unitPrice':100,'priceSource':'manual'}])
        self.create(currency='USD',exchangeRate=None,shippingFee=0,items=[{'productId':1,'sizeCode':'M','quantity':2,'unitPrice':10,'priceSource':'manual'}])
        today=datetime.now().date().isoformat()
        query={'dateFrom':today,'dateTo':today,'exchangeRate':7}
        response=self.call('finance/reports/query','POST',query);self.assertEqual(response.status_code,200,response.json)
        report=response.json
        self.assertEqual(report['summary']['salesUsd'],'20.00');self.assertEqual(report['summary']['nativeSalesCny'],'200.00')
        self.assertEqual(report['summary']['salesCny'],'340.00');self.assertIsNone(report['summary']['profitCny'])
        self.assertEqual({line['currency'] for row in report['items'] for line in row['lines']},{'USD','CNY'})
        style=report['items'][0]['styleCode']
        db._fetch_one("INSERT INTO finance_style_cost_audit(style_code,cost_cny,actor,request_id,is_initial) VALUES(%s,10,'{}',%s,true) RETURNING id",(style,uuid.uuid4()))
        known=self.call('finance/reports/query','POST',query).json
        self.assertEqual(known['summary']['totalCostCny'],'40.00');self.assertEqual(known['summary']['profitCny'],'300.00')
        dash=self.call(f'dashboard?view=workbench&dateFrom={today}&dateTo={today}')
        self.assertEqual(dash.status_code,200,dash.json);self.assertAlmostEqual(dash.json['metrics']['amount'],47.78)
        response=self.call('finance/reports/export','POST',query)
        self.assertEqual(response.status_code,200,response.json)
        book=load_workbook(BytesIO(response.data));headers=[c.value for c in book['订单明细'][1]]
        self.assertIn('成交币种',headers);self.assertIn('原币商品金额',headers)
        self.assertEqual(book['订单明细'].cell(2,headers.index('原币商品金额')+1).data_type,'n')

    def test_native_exports_and_shipping_compatibility(self):
        order=self.historical_cny()
        for route in ['orders/export','orders/export-by-sheet']:
            response=self.call(route+f"?orderIds={order['id']}&includeImages=false")
            self.assertEqual(response.status_code,200,response.json)
            sheet=load_workbook(BytesIO(response.data)).active
            headers=[c.value for c in sheet[5]]
            price_col=headers.index('单价 CNY')+1
            self.assertEqual(sheet.cell(6,price_col).value,80)
            self.assertIn('CNY',sheet.cell(6,price_col).number_format)
        db._fetch_one('UPDATE finance_settings SET exchange_rate=9 WHERE id=1 RETURNING id')
        response=self.call(f"orders/{order['id']}/invoice")
        self.assertEqual(response.status_code,200,response.json)
        book=load_workbook(BytesIO(response.data))
        self.assertTrue(any(c.value==11.11 and 'USD' in c.number_format for row in book.active for c in row))
        self.assertIn('(USD)',book.active['J11'].value)
        self.assertEqual(book.active['I12'].value,'(USD)')
        template = load_workbook(application.PROFORMA_TEMPLATE_PATH)['PI']
        self.assertTrue(any(c.value == template['A29'].value for row in book.active for c in row))
        self.assertTrue(any(c.value == template['A28'].value for row in book.active for c in row))
        self.assertFalse(any('CNY payment details:' in str(c.value) for row in book.active for c in row))
        saved = self.call(f"orders/{order['id']}").json['order']
        self.assertEqual(saved['currency'], 'CNY')
        self.assertEqual(saved['items'][0]['unitPrice'], 80)
        self.assertEqual(saved['totalAmount'], 170)
        db._fetch_one("UPDATE orders SET payment_link='https://pay.example.test/old' WHERE id=%s RETURNING id",(order['id'],))
        current=self.call(f"orders/{order['id']}").json['order']
        response=self.call(f"orders/{order['id']}",'PUT',{'version':current['version'],'status':current['status'],'shippingFee':20,'paymentLink':current['paymentLink']})
        self.assertEqual(response.status_code,200,response.json)
        self.assertEqual(response.json['order']['totalAmount'],180);self.assertEqual(response.json['order']['paymentLink'],'')

    def test_validation_and_idempotent_migration(self):
        for extra in [{'currency':'EUR'},{'currency':'CNY','exchangeRate':None},{'exchangeRate':0},{'exchangeRate':'7.1234567'}]:
            self.assertEqual(self.call('orders','POST',self.payload(**extra)).status_code,400)
        order=self.create();before=db._fetch_one('SELECT order_no,total_amount,currency,exchange_rate FROM orders WHERE id=%s',(order['id'],))
        with db._connect() as conn:
            order_pricing.migrate(conn.cursor());order_pricing.migrate(conn.cursor())
        self.assertEqual(db._fetch_one('SELECT order_no,total_amount,currency,exchange_rate FROM orders WHERE id=%s',(order['id'],)),before)

    def test_pi_retains_real_nonstandard_size_quantity_and_price(self):
        order=self.historical_cny(items=[{'productId':1,'sizeCode':'Tall XL','quantity':2,'unitPrice':82,'priceSource':'manual'}])
        response=self.call(f"orders/{order['id']}/invoice")
        self.assertEqual(response.status_code,200,response.json)
        sheet=load_workbook(BytesIO(response.data)).active
        row=next(cell.row for cells in sheet for cell in cells if cell.column==1 and 'Size: Tall XL' in str(cell.value))
        self.assertEqual(sheet.cell(row,8).value,2);self.assertEqual(sheet.cell(row,9).value,11.39)
        self.assertIn('USD',sheet.cell(row,9).number_format)

    def test_old_usd_draft_with_unused_fx_can_resume_save_and_submit(self):
        body = self.payload()
        response = self.call('orders/drafts','POST',{'draftKey':str(uuid.uuid4()),'draft':body})
        self.assertEqual(response.status_code,200)
        draft_id = response.json['id']
        db._fetch_one("UPDATE admin_order_drafts SET document=jsonb_set(document,'{exchangeRate}','\"7.2\"'::jsonb) WHERE id=%s RETURNING id",(draft_id,))
        restored = self.call(f'orders/drafts/{draft_id}').json
        self.assertIsNone(restored['draft']['exchangeRate'])
        self.assertEqual(restored['draft']['items'][0]['unitPrice'],80)
        saved = self.call(f'orders/drafts/{draft_id}','PUT',{'version':restored['version'],'draft':restored['draft']})
        self.assertEqual(saved.status_code,200,saved.json)
        restored = self.call(f'orders/drafts/{draft_id}').json
        result = self.call('orders','POST',{**restored['draft'],'draftId':draft_id,
            'draftVersion':restored['version'],'requestId':restored['submitRequestId']})
        self.assertEqual(result.status_code,200,result.json)
        self.assertEqual(result.json['order']['currency'],'USD')
        self.assertEqual(result.json['order']['totalAmount'],170)

    def test_new_orders_drafts_and_edits_reject_currency_and_fx_changes_atomically(self):
        before = db._fetch_all('SELECT id,total_amount FROM orders ORDER BY id')
        stock = db._fetch_all('SELECT id,stock FROM product_size_prices ORDER BY id')
        for extra in [{'currency':'CNY','exchangeRate':'7.2'}, {'exchangeRate':'7.2'}, {'currency':'EUR'}]:
            body = self.payload(**extra)
            self.assertEqual(self.call('orders','POST',body).status_code,400)
            response = self.call('orders/drafts','POST',{'draftKey':str(uuid.uuid4()),'draft':body})
            self.assertEqual(response.status_code,400)
        self.assertEqual(db._fetch_all('SELECT id,total_amount FROM orders ORDER BY id'),before)
        self.assertEqual(db._fetch_all('SELECT id,stock FROM product_size_prices ORDER BY id'),stock)
        order = self.create()
        for extra in [{'currency':'CNY','exchangeRate':'7.2'}, {'exchangeRate':'7.2'}]:
            response = self.call(f"orders/{order['id']}/details",'PUT',{**order,**extra})
            self.assertEqual(response.status_code,400)
        current = self.call(f"orders/{order['id']}").json['order']
        self.assertEqual(current['currency'],'USD')
        self.assertEqual(current['totalAmount'],order['totalAmount'])
        self.assertIsNone(current['exchangeRate'])
        settings = self.call('orders/pricing-settings').json
        self.assertEqual(settings,{'defaultCurrency':'USD','currencies':['USD'],'exchangeRate':None})
        self.assertEqual(db._fetch_one('SELECT exchange_rate FROM finance_settings WHERE id=1')['exchange_rate'],Decimal('7.2'))


if __name__=='__main__':unittest.main()
