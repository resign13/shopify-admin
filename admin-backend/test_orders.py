"""Order editing tests use the guarded synthetic database from test_workbench."""
import copy
import os
import unittest
import uuid
from concurrent.futures import ThreadPoolExecutor
from test_workbench import seed, application, db, fixture_file


class AdminOrdersTest(unittest.TestCase):
    def test_notes_nine_images_and_atomic_validation(self):
        order = self.create()
        body = copy.deepcopy(order)
        body['note'] = 'Packing note: keep all nine photos'
        body['labelImageUrls'] = [fixture_file(f'/uploads/photo-{i}.jpg') for i in range(9)]
        body['status'] = 'allocated'
        response = self.call(f"orders/{order['id']}/details", 'PUT', body)
        self.assertEqual(response.status_code, 200, response.json)
        updated = self.call(f"orders/{order['id']}").json['order']
        self.assertEqual(updated['labelImageUrls'], body['labelImageUrls'])
        self.assertEqual(updated['note'], body['note'])
        self.assertEqual(updated['status'], 'allocated')
        stock = self.stock()['stock']
        invalid = copy.deepcopy(updated)
        invalid['labelImageUrls'].append('/uploads/extra.jpg')
        invalid['items'][0]['quantity'] += 1
        invalid['note'] = 'must roll back'
        self.assertEqual(self.call(f"orders/{order['id']}/details", 'PUT', invalid).status_code, 400)
        self.assertEqual(self.stock()['stock'], stock)
        self.assertEqual(self.call(f"orders/{order['id']}").json['order']['note'], body['note'])
        for urls in [['javascript:photo.jpg'], ['/uploads/not-image.pdf'], ['//evil.test/photo.jpg'], ['/uploads/a.jpg'] * 2]:
            invalid['labelImageUrls'] = urls
            self.assertEqual(self.call(f"orders/{order['id']}/details", 'PUT', invalid).status_code, 400)
        invalid = copy.deepcopy(updated); invalid['status'] = 'shipped'; invalid['trackingNo'] = ''
        self.assertEqual(self.call(f"orders/{order['id']}/details", 'PUT', invalid).status_code, 400)
        self.assertEqual(self.call(f"orders/{order['id']}/details", 'PUT', body).status_code, 409)

    def test_image_removal_and_legacy_file_preservation(self):
        order = self.create(); oid = order['id']
        fixture_file('/uploads/legacy.pdf', 'order', oid)
        fixture_file('/uploads/old.jpg', 'order', oid)
        db._fetch_one('UPDATE orders SET label_pdf_url=%s,label_image_urls=%s WHERE id=%s RETURNING id',
                      ('/uploads/legacy.pdf', '["/uploads/legacy.pdf","/uploads/old.jpg"]', oid))
        body = self.call(f'orders/{oid}').json['order']; body['labelImageUrls'] = [fixture_file('/uploads/new.jpg')]
        self.assertEqual(self.call(f'orders/{oid}/details', 'PUT', body).status_code, 200)
        updated = self.call(f'orders/{oid}').json['order']
        self.assertEqual(updated['labelImageUrls'], [fixture_file('/uploads/new.jpg'), fixture_file('/uploads/legacy.pdf','order',oid)])
        updated['labelImageUrls'] = []
        self.assertEqual(self.call(f'orders/{oid}/details', 'PUT', updated).status_code, 200)
        self.assertEqual(self.call(f'orders/{oid}').json['order']['labelImageUrls'], [fixture_file('/uploads/legacy.pdf','order',oid)])
        fixture_file('/uploads/legacy.pdf', 'order', oid)
        fixture_file('/uploads/old.jpg', 'order', oid)
        db._fetch_one('UPDATE orders SET label_pdf_url=%s,label_image_urls=%s WHERE id=%s RETURNING id',
                      ('/uploads/old.jpg', '[]', oid))
        body = self.call(f'orders/{oid}').json['order']; body['labelImageUrls'] = []
        self.assertEqual(self.call(f'orders/{oid}/details', 'PUT', body).status_code, 200)
        self.assertEqual(self.call(f'orders/{oid}').json['order']['labelImageUrls'], [])

    def test_matrix_exports_use_local_images_and_preserve_size_quantities(self):
        from tempfile import TemporaryDirectory
        from pathlib import Path
        from unittest.mock import patch
        from PIL import Image
        from openpyxl import load_workbook
        with TemporaryDirectory() as directory:
            Image.new('RGB', (600, 900), '#335577').save(Path(directory) / 'matrix.jpg')
            order = {'orderNo': 'MATRIX-FIXTURE', 'items': [
                {'productId': 1, 'sku': 'BLUE', 'sizeCode': size, 'quantity': qty,
                 'image': 'https://img.smawell.shop/uploads/matrix.jpg'}
                for size, qty in [('S', 2), ('30/M', 3), ('XXL', 4)]]}
            application.media.register_legacy('https://img.smawell.shop/uploads/matrix.jpg',(Path(directory)/'matrix.jpg').read_bytes(),purpose='product',entity_type='product',entity_id=1)
            for builder in [application.build_orders_export, application.build_orders_sheet_export]:
                with patch.dict(application.media.uploads, {'admin': Path(directory)}), patch.dict(os.environ, {'SECURITY_ADMIN_LEGACY_ORIGINS': 'https://img.smawell.shop'}), application.app.test_request_context('/api/admin/orders/export'):
                    sheet = load_workbook(builder([order])).active
                    self.assertEqual(len(sheet._images), 1)
                    self.assertEqual([sheet.cell(6, c).value for c in [3, 4, 7, 8]], [2, 3, 4, 9])

    def setUp(self):
        self.tokens, _ = seed()
        application.app.config['TESTING'] = True

    def call(self, path, method='GET', body=None, role='admin'):
        with application.app.test_client() as client:
            return client.open('/api/admin/'+path, method=method, json=body,
                               headers={'Authorization':'Bearer '+self.tokens[role]})

    def payload(self):
        return {'requestId':str(uuid.uuid4()), 'userId':1, 'contactName':'Alex', 'phone':'123456',
                'country':'Germany', 'address':'18 Market Street', 'city':'Berlin', 'shippingFee':12.5,
                'items':[{'productId':1,'sizeCode':'M','quantity':2,'unitPrice':20},
                         {'productId':1,'sizeCode':'L','quantity':3,'unitPrice':25}]}

    def create(self):
        response=self.call('orders','POST',self.payload())
        self.assertEqual(response.status_code,200,response.json)
        return self.call(f"orders/{response.json['order']['id']}").json['order']

    def stock(self):
        return self.call('inventory/1').json['product']

    def test_create_and_concurrent_replay(self):
        body=self.payload()
        with ThreadPoolExecutor(max_workers=2) as pool:
            results=list(pool.map(lambda _:self.call('orders','POST',body),range(2)))
        self.assertTrue(all(r.status_code==200 for r in results),[r.json for r in results])
        self.assertEqual(len({r.json['order']['id'] for r in results}),1)
        self.assertEqual(sum(r.json['replayed'] for r in results),1)
        self.assertEqual(results[0].json['order']['totalAmount'],127.5)
        self.assertEqual(self.stock()['stock'],25)
        self.assertEqual(sum(s['contractPending'] for s in self.stock()['sizePrices']),21)
        self.assertEqual(sum(s['pendingInbound'] for s in self.stock()['sizePrices']),6)
        self.assertEqual(db._fetch_one("SELECT COUNT(*) AS n FROM admin_audit_logs WHERE entity_table='orders' AND action='INSERT'")['n'],1)
        body['note']='changed body'
        self.assertEqual(self.call('orders','POST',body).status_code,400)

    def test_edit_delta_price_contact_and_stale_version(self):
        order=self.create();body=copy.deepcopy(order)
        body['items'][0]['quantity']=4
        body['items'][1]['quantity']=1
        body['items'][0]['unitPrice']=22
        body['address']='New address';body['contactName']='Updated contact';body['note']='updated'
        response=self.call(f"orders/{order['id']}/details",'PUT',body)
        self.assertEqual(response.status_code,200,response.json)
        updated=response.json['order']
        self.assertEqual(updated['totalAmount'],125.5)
        self.assertIn('New address',updated['shippingAddress'])
        self.assertEqual(updated['contactName'],'Updated contact')
        sizes={r['sizeCode']:r['stock'] for r in self.stock()['sizePrices']}
        self.assertEqual(sizes,{'M':6,'L':9,'Tall XL':10})
        self.assertEqual(self.stock()['stock'],25)
        self.assertEqual(self.call(f"orders/{order['id']}/details",'PUT',body).status_code,409)
        body=self.call(f"orders/{order['id']}").json['order'];body['items']=body['items'][:1]
        self.assertEqual(self.call(f"orders/{order['id']}/details",'PUT',body).status_code,200)
        self.assertEqual(self.stock()['stock'],26)

    def test_overflow_and_invalid_input_rollback(self):
        for modify in [lambda b:b['items'][1].update(quantity=2147483648),
                       lambda b:b['items'][0].update(quantity=1.5),
                       lambda b:b['items'][0].update(sizeCode='UNKNOWN'),
                       lambda b:b['items'][0].update(unitPrice=-1),
                       lambda b:b['items'].append(dict(b['items'][0])),
                       lambda b:b.update(userId=999),lambda b:b.update(shippingFee=-1)]:
            body=self.payload();modify(body)
            response=self.call('orders','POST',body)
            self.assertEqual(response.status_code,400,response.json)
            self.assertEqual(self.stock()['stock'],30)
        self.assertEqual(db._fetch_one('SELECT COUNT(*) AS n FROM orders')['n'],4)
        self.assertEqual(db._fetch_one('SELECT COUNT(*) AS n FROM admin_audit_logs')['n'],0)

    def test_allocated_status_filter_export_and_migration(self):
        order=self.create()
        response=self.call(f"orders/{order['id']}",'PUT',{'status':'allocated','shippingFee':12.5,'version':order['version']},role='warehouse')
        self.assertEqual(response.status_code,200,response.json)
        self.assertEqual(self.stock()['stock'],25)
        result=self.call('orders?page=1&status=allocated').json
        self.assertEqual(result['total'],1)
        self.assertEqual(result['statusCounts']['allocated'],1)
        self.assertEqual(self.call('orders/export?view=workbench&status=allocated&includeImages=0', 'POST', role='warehouse').status_code,200)
        self.assertEqual(self.call('orders/export-by-sheet?view=workbench&status=allocated&includeImages=0', 'POST', role='warehouse').status_code,200)
        self.assertEqual(self.call(f"orders/{order['id']}",'PUT',{'status':'paid','shippingFee':99,'paymentLink':'https://changed.example'},role='warehouse').status_code,200)
        updated=self.call(f"orders/{order['id']}").json['order']
        self.assertEqual(float(updated['shippingFee']),12.5)
        self.assertEqual(updated['paymentLink'],'')
        with db.get_connection() as conn:
            with conn.cursor() as cur: db._migrate_order_status_values(cur)
        self.assertEqual(self.call(f"orders/{order['id']}").json['order']['status'],'paid')
        self.assertEqual(self.call(f"orders/{order['id']}",'PUT',{'status':'allocated','shippingFee':-1}).status_code,400)

    def test_partial_allocation_filter_exports_audit_and_migration(self):
        from io import BytesIO
        from openpyxl import load_workbook
        order = self.create()
        before = self.stock()
        response = self.call(f"orders/{order['id']}", 'PUT', {
            'status': 'partially_allocated', 'version': order['version'],
            'shippingFee': 999, 'paymentLink': 'https://ignored.invalid',
        }, 'warehouse')
        self.assertEqual(response.status_code, 200, response.json)
        self.assertTrue({'totalAmount', 'shippingFee', 'paymentLink'}.isdisjoint(response.json['order']))
        current = self.call(f"orders/{order['id']}").json['order']
        self.assertEqual(current['status'], 'partially_allocated')
        self.assertEqual((current['shippingFee'], current['totalAmount']), (12.5, 127.5))
        self.assertEqual(self.stock(), before)
        listing = self.call('orders?page=1&status=partially_allocated').json
        self.assertEqual((listing['total'], listing['statusCounts']['partially_allocated']), (1, 1))
        self.assertEqual(listing['items'][0]['id'], order['id'])
        for export in ['export', 'export-by-sheet']:
            result = self.call(f'orders/{export}?view=workbench&status=partially_allocated&includeImages=0', 'POST')
            self.assertEqual(result.status_code, 200, result.json)
            workbook = load_workbook(BytesIO(result.data))
            values = [cell.value for sheet in workbook for row in sheet for cell in row]
            self.assertTrue(any(order['orderNo'] in str(value) for value in values))
            self.assertFalse(any('GT20260918' in str(value) for value in values))
        history = self.call(f"orders/{order['id']}/operations").json['items']
        self.assertTrue(any(c['field'] == 'status' and c['after'] == 'partially_allocated'
                            for h in history for c in h['changes']))
        with db.get_connection() as conn:
            db._migrate_order_status_values(conn.cursor())
            db._migrate_order_status_values(conn.cursor())
        self.assertEqual(self.call(f"orders/{order['id']}").json['order']['status'], 'partially_allocated')
        self.assertEqual(self.stock(), before)
        # The old version must not overwrite a later status change.
        self.assertEqual(self.call(f"orders/{order['id']}", 'PUT', {
            'status': 'allocated', 'version': order['version'],
        }).status_code, 409)

    def test_partial_allocation_details_delta_cancel_restore_and_scope(self):
        order = self.create()
        self.assertEqual(self.call(f"orders/{order['id']}", 'PUT', {'status': 'partially_allocated'}, 'sales').status_code, 404)
        self.assertEqual(self.call(f"orders/{order['id']}", 'PUT', {'status': 'partially_allocated'}, 'customer').status_code, 403)
        order['status'] = 'partially_allocated'
        order['items'][0]['quantity'] = 4
        result = self.call(f"orders/{order['id']}/details", 'PUT', order)
        self.assertEqual(result.status_code, 200, result.json)
        current = result.json['order']
        self.assertEqual(current['status'], 'partially_allocated')
        self.assertEqual(self.stock()['stock'], 23)
        self.assertEqual([r['currentStock'] for r in current['items']], [6, 7])
        for target in ['allocated', 'paid', 'partially_allocated']:
            current = self.change_status(current, target)
            self.assertEqual(self.stock()['stock'], 23)
        current = self.change_status(current, 'cancelled')
        self.assertEqual(self.stock()['stock'], 30)
        current = self.change_status(current, 'partially_allocated')
        self.assertEqual(self.stock()['stock'], 23)
        current = self.change_status(current, 'shipped', trackingNo='PARTIAL-FIXTURE')
        self.assertEqual(self.stock()['stock'], 23)
        self.assertEqual(current['status'], 'shipped')

    def test_partial_allocation_protects_unfulfilled_true_sizes(self):
        current = self.change_status(self.create(), 'partially_allocated')
        with db.get_connection() as conn:
            conn.execute('UPDATE product_size_prices SET contract_pending=0,pending_inbound=0,pending_inspection=0 WHERE product_id=1')
            # Remove unrelated historical references, isolating this new state.
            conn.execute('DELETE FROM order_items WHERE order_id<>%s', (current['id'],))
        product = self.call('products/1').json['product']
        product.update(title=product['name']['zh'], familyCode=product['colorGroup'])
        product['sizePrices'] = [s for s in product['sizePrices'] if s['sizeCode'] != 'M']
        product['sizes'] = [s['sizeCode'] for s in product['sizePrices']]
        with self.assertRaisesRegex(ValueError, '未履约订单'):
            db.update_product(1, product)
        self.assertEqual(self.stock()['stock'], 25)

    def test_warehouse_order_responses_hide_financial_fields(self):
        order = self.create()
        financial_fields = {'totalAmount', 'goodsAmount', 'shippingFee', 'paymentLink'}

        listing = self.call('orders?page=1', role='warehouse').json
        listed = next(item for item in listing['items'] if item['id'] == order['id'])
        self.assertTrue(financial_fields.isdisjoint(listed))
        self.assertTrue(all({'unitPrice', 'totalPrice'}.isdisjoint(item) for item in listed['items']))

        detail = self.call(f"orders/{order['id']}", role='warehouse').json['order']
        self.assertTrue(financial_fields.isdisjoint(detail))
        self.assertTrue(all({'unitPrice', 'totalPrice'}.isdisjoint(item) for item in detail['items']))

        updated = self.call(
            f"orders/{order['id']}",
            'PUT',
            {'status': 'allocated', 'shippingFee': 999, 'paymentLink': 'https://example.invalid'},
            role='warehouse',
        ).json['order']
        self.assertTrue(financial_fields.isdisjoint(updated))

    def test_permissions_and_fulfilled_lines(self):
        for role in ['warehouse','customer']:
            self.assertEqual(self.call('orders','POST',self.payload(),role).status_code,403)
            self.assertEqual(self.call('orders/1/details','PUT',self.payload(),role).status_code,403)
            self.assertEqual(self.call('orders/customers',role=role).status_code,403)
        self.assertEqual(self.call('orders/customers?keyword=Northline',role='sales').json['total'],1)
        order=self.create()
        self.assertEqual(self.call(f"orders/{order['id']}",'PUT',{'status':'shipped','trackingNo':'TEST-TRACK','shippingFee':12.5}).status_code,200)
        body=self.call(f"orders/{order['id']}").json['order'];body['items'][0]['quantity']=5
        response=self.call(f"orders/{order['id']}/details",'PUT',body)
        self.assertEqual(response.status_code,200,response.json)
        self.assertEqual(response.json['order']['status'],'shipped')
        self.assertEqual(response.json['order']['trackingNo'],'TEST-TRACK')
        self.assertEqual(self.stock()['stock'],22)
        body=response.json['order'];body['items'][0]['quantity']=2;body['note']='Contact update after shipment'
        self.assertEqual(self.call(f"orders/{order['id']}/details",'PUT',body).status_code,200)
        self.assertEqual(self.stock()['stock'],25)

    def change_status(self, order, status, role='admin', **fields):
        response = self.call(f"orders/{order['id']}", 'PUT', {
            'version': order['version'], 'status': status, 'shippingFee': order['shippingFee'],
            'paymentLink': order['paymentLink'], 'trackingNo': order['trackingNo'], **fields,
        }, role)
        self.assertEqual(response.status_code, 200, response.json)
        return self.call(f"orders/{order['id']}").json['order']

    def test_fulfilled_items_add_remove_price_delta_and_stale_save(self):
        for state in ['shipped', 'completed']:
            with self.subTest(state=state):
                order = self.change_status(self.create(), state, trackingNo='FULFILLED-QA')
                before = {r['sizeCode']: r['stock'] for r in self.stock()['sizePrices']}
                db._fetch_one("UPDATE orders SET payment_link='https://pay.example/old' WHERE id=%s RETURNING id", (order['id'],))
                body = self.call(f"orders/{order['id']}").json['order']
                old_completed_at, old_shipped_at = body['completedAt'], body['shippedAt']
                body['items'] = [{**body['items'][0], 'quantity': 4, 'unitPrice': 22, 'priceSource': 'manual'},
                    {'productId': 1, 'sizeCode': 'Tall XL', 'quantity': 2, 'unitPrice': 29.9, 'priceSource': 'catalog', 'referencePriceUsd': 29.9},
                    {'productId': 2, 'sizeCode': 'M', 'quantity': 1, 'unitPrice': 15, 'priceSource': 'manual'}]
                response = self.call(f"orders/{order['id']}/details", 'PUT', body)
                self.assertEqual(response.status_code, 200, response.json)
                updated = response.json['order']
                self.assertEqual((updated['status'], updated['trackingNo']), (state, 'FULFILLED-QA'))
                self.assertEqual((updated['completedAt'], updated['shippedAt']), (old_completed_at, old_shipped_at))
                self.assertEqual(updated['totalAmount'], 175.3)
                self.assertEqual(updated['paymentLink'], '')
                after = {r['sizeCode']: r['stock'] for r in self.stock()['sizePrices']}
                self.assertEqual(after, {'M': before['M'] - 2, 'L': before['L'] + 3, 'Tall XL': before['Tall XL'] - 2})
                self.assertEqual(db._fetch_one("SELECT stock FROM product_size_prices WHERE product_id=2 AND size_code='M'")['stock'], 9)
                self.assertEqual(self.call(f"orders/{order['id']}/details", 'PUT', body).status_code, 409)
                self.assertEqual({r['sizeCode']: r['stock'] for r in self.stock()['sizePrices']}, after)
                history = self.call(f"orders/{order['id']}/operations").json['items'][0]
                self.assertTrue(any(c['field'] == 'unitPrice' for c in history['changes']))
                # Cancellation after editing returns exactly the saved quantities.
                cancelled = self.change_status(updated, 'cancelled', cancelNote='Fixture cancellation')
                self.assertEqual({r['sizeCode']: r['stock'] for r in self.stock()['sizePrices']}, {'M': 10, 'L': 10, 'Tall XL': 10})
                bad = copy.deepcopy(cancelled); bad['items'][0]['quantity'] += 1
                self.assertEqual(self.call(f"orders/{order['id']}/details", 'PUT', bad).status_code, 400)

    def test_fulfilled_items_permissions_and_validation_rollback(self):
        body = self.payload(); body['ownerAdminId'] = 2
        made = self.call('orders', 'POST', body).json['order']
        order = self.change_status(made, 'completed')
        before_stock = self.stock()
        for role in ['warehouse', 'customer']:
            self.assertEqual(self.call(f"orders/{order['id']}/details", 'PUT', order, role).status_code, 403)
        for changes in [dict(sizeCode='REMOVED-SIZE'), dict(quantity=-1), dict(unitPrice=-1)]:
            bad = copy.deepcopy(order); bad['items'][0].update(changes)
            self.assertEqual(self.call(f"orders/{order['id']}/details", 'PUT', bad).status_code, 400)
            self.assertEqual(self.call(f"orders/{order['id']}").json['order'], order)
            self.assertEqual(self.stock(), before_stock)
        order['items'][0]['quantity'] += 1
        response = self.call(f"orders/{order['id']}/details", 'PUT', order, 'sales')
        self.assertEqual(response.status_code, 200, response.json)
        self.assertEqual(response.json['order']['status'], 'completed')
        other = self.change_status(self.create(), 'shipped', trackingNo='OTHER-QA')
        other['items'][0]['quantity'] += 1
        self.assertEqual(self.call(f"orders/{other['id']}/details", 'PUT', other, 'sales').status_code, 404)

    def test_completed_can_reopen_to_every_status_with_one_operation_and_no_new_event(self):
        for target in ['pending_payment', 'partially_allocated', 'allocated', 'paid', 'shipped', 'cancelled']:
            with self.subTest(target=target):
                stock_before_creation = self.stock()['stock']
                order = self.create()
                completed = self.change_status(order, 'completed')
                self.assertTrue(completed['completedAt'])
                stages = [(r['sizeCode'], r['contractPending'], r['pendingInspection'],
                           r['pendingInbound'], r['defectivePending'], r['temporaryInbound'])
                          for r in self.stock()['sizePrices']]
                events = db._fetch_one('SELECT COUNT(*) AS n FROM order_created_events')['n']
                updated = self.change_status(completed, target, trackingNo='REOPEN-FIXTURE' if target == 'shipped' else '')
                self.assertEqual(updated['status'], target)
                self.assertFalse(updated['completedAt'])
                self.assertEqual([{k:v for k,v in r.items() if k!='currentStock'} for r in updated['items']],
                                 [{k:v for k,v in r.items() if k!='currentStock'} for r in completed['items']])
                for old, now in zip(completed['items'], updated['items']):
                    self.assertEqual(now['currentStock'],old['currentStock']+(old['quantity'] if target=='cancelled' else 0))
                self.assertEqual(updated['totalAmount'], 127.5)
                self.assertEqual(updated['shippingFee'], 12.5)
                expected = stock_before_creation if target == 'cancelled' else stock_before_creation - 5
                self.assertEqual(self.stock()['stock'], expected)
                self.assertEqual([(r['sizeCode'], r['contractPending'], r['pendingInspection'],
                                   r['pendingInbound'], r['defectivePending'], r['temporaryInbound'])
                                  for r in self.stock()['sizePrices']], stages)
                history = self.call(f"orders/{order['id']}/operations").json
                self.assertEqual(history['total'], 2)
                changes = {c['field']: c for c in history['items'][0]['changes']}
                self.assertEqual((changes['status']['before'], changes['status']['after']), ('completed', target))
                self.assertEqual(db._fetch_one('SELECT COUNT(*) AS n FROM order_created_events')['n'], events)
                self.assertEqual(self.call(f"orders/{order['id']}", 'PUT', {
                    'version': completed['version'], 'status': target, 'trackingNo': 'REOPEN-FIXTURE',
                }).status_code, 409)
                self.assertEqual(self.stock()['stock'], expected)
                self.assertEqual(self.call(f"orders/{order['id']}/operations").json['total'], 2)

    def test_completed_details_reopen_with_quantity_delta(self):
        completed = self.change_status(self.create(), 'completed')
        bad = copy.deepcopy(completed)
        bad['status'] = 'paid'
        bad['items'][0]['quantity'] += 1
        response = self.call(f"orders/{completed['id']}/details", 'PUT', bad)
        self.assertEqual(response.status_code, 200, response.json)
        self.assertEqual(self.stock()['stock'], 24)
        self.assertEqual(self.call(f"orders/{completed['id']}/operations").json['total'], 2)
        body = response.json['order']
        body.update(status='paid', note='Reopened by staff')
        response = self.call(f"orders/{completed['id']}/details", 'PUT', body)
        self.assertEqual(response.status_code, 200, response.json)
        reopened = self.call(f"orders/{completed['id']}").json['order']
        self.assertEqual((reopened['status'], reopened['note']), ('paid', 'Reopened by staff'))
        self.assertFalse(reopened['completedAt'])
        self.assertEqual(self.stock()['stock'], 24)
        recompleted = self.change_status(reopened, 'completed')
        self.assertTrue(recompleted['completedAt'])
        recompleted['status'] = 'cancelled'
        response = self.call(f"orders/{recompleted['id']}/details", 'PUT', recompleted)
        self.assertEqual(response.status_code, 200, response.json)
        self.assertEqual(self.stock()['stock'], 30)
        self.assertEqual(self.call(f"orders/{recompleted['id']}/operations").json['total'], 5)

    def test_completed_invalid_and_concurrent_cancel_are_atomic(self):
        completed = self.change_status(self.create(), 'completed')
        for fields in [{'status': 'invalid'}, {'status': 'shipped', 'trackingNo': ''}]:
            response = self.call(f"orders/{completed['id']}", 'PUT', {'version': completed['version'], **fields})
            self.assertEqual(response.status_code, 400, response.json)
        unchanged = self.call(f"orders/{completed['id']}").json['order']
        self.assertEqual(unchanged, completed)
        self.assertEqual(self.call(f"orders/{completed['id']}/operations").json['total'], 1)
        body = {'version': completed['version'], 'status': 'cancelled', 'shippingFee': 12.5}
        with ThreadPoolExecutor(max_workers=2) as pool:
            responses = list(pool.map(lambda _: self.call(f"orders/{completed['id']}", 'PUT', body), range(2)))
        self.assertEqual(sorted(r.status_code for r in responses), [200, 409])
        self.assertEqual(self.stock()['stock'], 30)
        cancelled = self.call(f"orders/{completed['id']}").json['order']
        self.change_status(cancelled, 'cancelled')
        self.assertEqual(self.stock()['stock'], 30)
        self.assertEqual(self.call(f"orders/{completed['id']}/operations").json['total'], 2)
        restored = self.change_status(self.call(f"orders/{completed['id']}").json['order'], 'paid')
        self.assertEqual(restored['status'], 'paid')
        self.assertEqual(self.stock()['stock'], 25)

    def test_completed_cancel_overflow_rolls_back_status_stock_and_history(self):
        completed = self.change_status(self.create(), 'completed')
        db.update_product_inventory(1, {'M': 2147483645, 'L': 0, 'Tall XL': 0})
        before = self.stock()
        # Inventory was intentionally changed after the earlier detail snapshot.
        before_order = self.call(f"orders/{completed['id']}").json['order']
        response = self.call(f"orders/{completed['id']}", 'PUT', {
            'version': completed['version'], 'status': 'cancelled', 'shippingFee': 12.5,
        })
        self.assertEqual(response.status_code, 400, response.json)
        self.assertEqual(self.stock(), before)
        self.assertEqual(self.call(f"orders/{completed['id']}").json['order'], before_order)
        self.assertEqual(self.call(f"orders/{completed['id']}/operations").json['total'], 1)

    def test_completed_reopen_respects_owner_role_module_and_warehouse_redaction(self):
        import inventory_policy
        own_response = self.call('orders', 'POST', self.payload(), 'sales')
        self.assertEqual(own_response.status_code, 200, own_response.json)
        own = self.call(f"orders/{own_response.json['order']['id']}").json['order']
        completed = self.change_status(own, 'completed')
        response = self.call(f"orders/{own['id']}", 'PUT', {
            'version': completed['version'], 'status': 'paid', 'shippingFee': 999,
            'paymentLink': 'https://fixture.invalid/changed',
        }, 'warehouse')
        self.assertEqual(response.status_code, 200, response.json)
        self.assertTrue({'totalAmount', 'goodsAmount', 'shippingFee', 'paymentLink'}.isdisjoint(response.json['order']))
        self.assertTrue(all({'unitPrice', 'totalPrice'}.isdisjoint(r) for r in response.json['order']['items']))
        updated = self.call(f"orders/{own['id']}").json['order']
        self.assertEqual((updated['shippingFee'], updated['totalAmount'], updated['paymentLink']), (12.5, 127.5, ''))
        completed = self.change_status(updated, 'completed')
        self.change_status(completed, 'allocated', 'sales')
        other = self.change_status(self.create(), 'completed')
        self.assertEqual(self.call(f"orders/{other['id']}", 'PUT', {'status': 'paid'}, 'sales').status_code, 404)
        self.assertEqual(self.call(f"orders/{own['id']}", 'PUT', {'status': 'paid'}, 'customer').status_code, 403)
        db._fetch_one("UPDATE admin_users SET permissions='[]' WHERE role='warehouse' RETURNING id")
        self.assertEqual(self.call(f"orders/{other['id']}", 'PUT', {'status': 'paid'}, 'warehouse').status_code, 403)
        for target in ['paid', 'cancelled', 'shipped']:
            with self.assertRaises(ValueError):
                inventory_policy.validate_transition('completed', target)


if __name__=='__main__':
    unittest.main()
