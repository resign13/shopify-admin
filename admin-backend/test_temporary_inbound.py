"""Independent temporary receipts; only run against the guarded local *_test database."""
import hashlib
import importlib.util
import io
import json
from pathlib import Path
import unittest
import uuid
from concurrent.futures import ThreadPoolExecutor

from psycopg.errors import CheckViolation
from psycopg.types.json import Jsonb
from openpyxl import load_workbook
import test_workbench as fixtures
import test_orders as order_fixtures
import workbench
application, db = fixtures.application, fixtures.db


class TemporaryInboundTest(unittest.TestCase):
    setUp = fixtures.WorkbenchTest.setUp
    call = fixtures.WorkbenchTest.call
    detail = fixtures.WorkbenchTest.detail
    workbook = fixtures.WorkbenchTest.workbook
    upload = fixtures.WorkbenchTest.upload
    product_payload = fixtures.WorkbenchTest.product_payload

    def stage(self, values, **extra):
        return self.call('inventory/1', 'PUT', {'version': self.detail()['version'], 'sizeStocks': {},
                                              'temporaryInboundBySize': values, **extra})

    def receive(self, body=None, source='temporary', role='admin'):
        if body is None: body = {'version': self.detail()['version'], 'requestId': str(uuid.uuid4())}
        suffix = '/temporary/receive' if source == 'temporary' else '/receive'
        return self.call('inventory/1' + suffix, 'POST', body, role)

    def sizes(self):
        return {s['sizeCode']: s for s in self.detail()['sizePrices']}

    def test_persistent_independent_registration_and_clear(self):
        before = self.sizes()
        self.assertEqual(self.stage({'M': 8, 'L': 3}).status_code, 200)
        after = self.sizes()
        for size in before:
            for key in ['stock', 'contractPending', 'pendingInspection', 'pendingInbound']:
                self.assertEqual(before[size][key], after[size][key])
        self.assertEqual((after['M']['temporaryInbound'], after['L']['temporaryInbound']), (8, 3))
        self.assertEqual(self.stage({'M': 5, 'L': 0}).status_code, 200)
        self.assertEqual(self.sizes()['M']['temporaryInbound'], 5)
        self.assertEqual(self.stage({'M': 0}).status_code, 200)
        self.assertEqual(sum(s['temporaryInbound'] for s in self.sizes().values()), 0)
        self.assertEqual(self.receive().status_code, 400)

    def test_validation_and_atomic_registration_rollback(self):
        before = self.detail()
        for invalid in [-1, 2147483648, 1.5, True, None, '', '1e2']:
            response = self.stage({'M': invalid}, sizeStocks={'L': 999})
            self.assertEqual(response.status_code, 400, (invalid, response.json))
            self.assertEqual(self.detail()['stock'], before['stock'])
        for values in [{'UNKNOWN': 2}, [], '3']:
            self.assertEqual(self.stage(values).status_code, 400)
        self.assertEqual(self.call('audit-logs?page=1').json['total'], 0)

    def test_receipts_add_signed_balance_and_do_not_consume_procurement(self):
        for stock in [-5, 0, 10]:
            with self.subTest(stock=stock):
                self.assertEqual(self.stage({'M': 8, 'L': 2}, sizeStocks={'M': stock}).status_code, 200)
                before = self.sizes()
                body = {'version': self.detail()['version'], 'requestId': str(uuid.uuid4()),
                        'temporaryInboundBySize': {'M': 999}, 'quantity': 999, 'source': 'normal'}
                response = self.receive(body)
                self.assertEqual(response.status_code, 200, response.json)
                self.assertEqual((response.json['source'], response.json['receivedUnits'], response.json['receivedSizes']), ('temporary', 10, 2))
                after = self.sizes()
                self.assertEqual(after['M']['stock'], stock + 8)
                for size in after:
                    self.assertEqual(after[size]['temporaryInbound'], 0)
                    for key in ['contractPending', 'pendingInspection', 'pendingInbound']:
                        self.assertEqual(after[size][key], before[size][key])
                self.assertEqual(self.detail()['stock'], sum(s['stock'] for s in after.values()))
                self.assertTrue(self.receive(body).json['replayed'])
                self.assertEqual(self.receive().status_code, 400)
                stored = db._fetch_one('SELECT result FROM inventory_receipts WHERE request_id=%s', (body['requestId'],))['result']
                self.assertEqual(stored['changes'][0]['temporaryInboundAfter'], 0)
                self.assertEqual(stored['changes'][0]['pendingInboundAfter'], before[stored['changes'][0]['sizeCode']]['pendingInbound'])

    def test_normal_receipt_preserves_temporary_and_legacy_replay(self):
        self.stage({'M': 7})
        normal = {'version': self.detail()['version'], 'requestId': str(uuid.uuid4())}
        response = self.receive(normal, source='normal')
        self.assertEqual(response.status_code, 200, response.json)
        self.assertEqual(response.json['receivedUnits'], 6)
        self.assertEqual(self.sizes()['M']['temporaryInbound'], 7)
        self.assertEqual(sum(s['pendingInbound'] for s in self.sizes().values()), 0)
        self.assertEqual(self.receive(normal).status_code, 400)  # same UUID, different source
        # Old ordinary receipt, before source metadata existed, remains replayable.
        old = {'version': 'old-version', 'requestId': str(uuid.uuid4())}
        result = {'productId': 1, 'receivedUnits': 4, 'receivedSizes': 1, 'changes': []}
        digest = hashlib.sha256(json.dumps({'productId': 1, 'version': old['version']}, sort_keys=True).encode()).hexdigest()
        db._fetch_one('INSERT INTO inventory_receipts(request_id,actor_id,product_id,request_hash,result) VALUES(%s,1,1,%s,%s) RETURNING request_id', (old['requestId'], digest, Jsonb(result)))
        self.assertEqual(self.receive(old, source='normal').json, {**result, 'replayed': True})

    def test_duplicate_concurrent_receipts_and_source_namespace(self):
        self.stage({'M': 8})
        body = {'version': self.detail()['version'], 'requestId': str(uuid.uuid4())}
        with ThreadPoolExecutor(max_workers=2) as pool:
            responses = list(pool.map(lambda _: self.receive(body), range(2)))
        self.assertEqual([r.status_code for r in responses], [200, 200], [r.json for r in responses])
        self.assertEqual(sum(bool(r.json.get('replayed')) for r in responses), 1)
        self.assertEqual(self.sizes()['M']['stock'], 18)
        self.assertEqual(self.receive(body, source='normal').status_code, 400)
        self.assertEqual(self.receive(body, role='warehouse').status_code, 400)
        self.stage({'M': 4})
        version = self.detail()['version']
        with ThreadPoolExecutor(max_workers=2) as pool:
            responses = list(pool.map(lambda _: self.receive({'version': version, 'requestId': str(uuid.uuid4())}), range(2)))
        self.assertEqual(sorted(r.status_code for r in responses), [200, 409])
        self.assertEqual(self.sizes()['M']['stock'], 22)

    def test_concurrent_normal_and_temporary_require_fresh_confirmation(self):
        self.stage({'M': 8})
        version = self.detail()['version']
        with ThreadPoolExecutor(max_workers=2) as pool:
            responses = list(pool.map(lambda source: self.receive({'version': version, 'requestId': str(uuid.uuid4())}, source), ['normal', 'temporary']))
        self.assertEqual(sorted(r.status_code for r in responses), [200, 409])
        success = next(r.json for r in responses if r.status_code == 200)
        after = self.sizes()
        if success['source'] == 'normal':
            self.assertEqual(after['M']['temporaryInbound'], 8)
        else:
            self.assertEqual(sum(s['pendingInbound'] for s in after.values()), 6)
        self.assertEqual(self.detail()['stock'], 30 + success['receivedUnits'])

    def test_concurrent_order_registration_and_receipt_conserve_stock(self):
        self.stage({'M': 8})
        version = self.detail()['version']
        order = order_fixtures.AdminOrdersTest.payload(self)
        order['items'] = [{'productId': 1, 'sizeCode': 'M', 'quantity': 2, 'unitPrice': 20}]
        with ThreadPoolExecutor(max_workers=2) as pool:
            a = pool.submit(self.receive, {'version': version, 'requestId': str(uuid.uuid4())})
            b = pool.submit(self.call, 'orders', 'POST', order)
            received, created = a.result(timeout=30), b.result(timeout=30)
        self.assertEqual(created.status_code, 200, created.json)
        self.assertIn(received.status_code, [200, 409], received.json)
        self.assertEqual(self.sizes()['M']['stock'], 10 - 2 + (8 if received.status_code == 200 else 0))
        self.assertEqual(self.sizes()['M']['temporaryInbound'], 0 if received.status_code == 200 else 8)

    def test_concurrent_product_edit_and_registration_do_not_lose_quantity(self):
        p = self.product_payload(1)
        version = self.detail()['version']
        p['title'] = 'Edited fixture title'
        with ThreadPoolExecutor(max_workers=2) as pool:
            a = pool.submit(self.call, 'inventory/1', 'PUT', {'version': version, 'sizeStocks': {}, 'temporaryInboundBySize': {'M': 8}})
            b = pool.submit(self.call, 'products/save-group', 'POST', {'products': [p], 'versions': {'1': version}})
            responses = [a.result(timeout=30), b.result(timeout=30)]
        self.assertEqual(sorted(r.status_code for r in responses), [200, 409], [r.json for r in responses])
        self.assertEqual(self.sizes()['M']['temporaryInbound'], 8 if responses[0].status_code == 200 else 0)

    def test_product_edit_removal_and_inactive_inventory(self):
        self.stage({'M': 7})
        p = self.product_payload(1)
        p['sizePrices'][0]['temporaryInbound'] = 999
        self.assertEqual(self.call('products/save-group', 'POST', {'products': [p], 'versions': {'1': p['version']}}).status_code, 200)
        self.assertEqual(self.sizes()['M']['temporaryInbound'], 7)
        db.update_product_inventory(1, {}, {'Tall XL': 0}, {'Tall XL': 0}, {'Tall XL': 0}, {'Tall XL': 5})
        p = self.product_payload(1)
        p['sizes'] = ['M', 'L']; p['sizePrices'] = p['sizePrices'][:2]
        self.assertEqual(self.call('products/save-group', 'POST', {'products': [p], 'versions': {'1': p['version']}}).status_code, 400)
        db._fetch_one('UPDATE product_size_prices SET contract_pending=0,pending_inspection=0,pending_inbound=0,temporary_inbound=3 WHERE product_id=3 RETURNING id')
        response = self.call('products/3', 'DELETE')
        self.assertEqual(response.status_code, 200, response.json)
        self.assertIsNotNone(db._fetch_one('SELECT id FROM products WHERE id=3 AND is_active=FALSE'))
        self.assertIn(3, [p['id'] for p in self.call('inventory?pageSize=50').json['items']])

    def test_excel_v4_and_legacy_templates_preserve_independence(self):
        raw = self.workbook(lambda s: setattr(s['Q2'], 'value', 7))
        preview = self.upload('preview', raw).json
        self.assertEqual(preview['errors'], [])
        self.assertEqual(preview['summary']['changedFieldCount'], 1)
        self.assertTrue(preview['rows'][0]['temporaryInboundChanged'])
        self.assertEqual(self.upload('confirm', raw, preview['fileHash']).status_code, 200)
        self.assertEqual((self.sizes()['M']['stock'], self.sizes()['M']['temporaryInbound']), (10, 7))
        self.assertTrue(self.upload('confirm', raw, preview['fileHash']).json['alreadyProcessed'])
        for version, columns in [('inventory-v3', 16), ('inventory-v2', 14), ('inventory-v1', 12)]:
            def legacy(sheet):
                sheet.delete_cols(columns + 1, sheet.max_column - columns)
                for column, label in enumerate(application.INVENTORY_HEADER_VERSIONS[version], 1):
                    sheet.cell(1, column, label)
                for row in range(2, sheet.max_row + 1): sheet.cell(row, 1, version)
                sheet['H2'] = int(sheet['H2'].value) + 1
            old = self.workbook(legacy)
            data = self.upload('preview', old).json
            self.assertEqual(data['errors'], [], data)
            self.assertNotIn('temporaryInbound', data['rows'][0])
            self.assertEqual(self.upload('confirm', old, data['fileHash']).status_code, 200)
            self.assertEqual(self.sizes()['M']['temporaryInbound'], 7)

    def test_excel_conflicts_formula_limits_and_atomic_rollback(self):
        raw = self.workbook(lambda s: setattr(s['Q2'], 'value', 7))
        preview = self.upload('preview', raw).json
        self.stage({'M': 5})
        self.assertEqual(self.upload('confirm', raw, preview['fileHash']).status_code, 400)
        self.assertEqual(self.sizes()['M']['temporaryInbound'], 5)
        rows, errors = application.parse_inventory_workbook(raw, db.list_products(include_contract_pending=True))
        self.assertEqual(errors, [])
        rows[0]['stock'] = 99
        with self.assertRaises(ValueError): db.apply_inventory_import(rows)
        self.assertEqual(self.sizes()['M']['stock'], 10)
        for invalid in ['=1+1', -1, 2147483648]:
            bad = self.workbook(lambda s: setattr(s['Q2'], 'value', invalid))
            self.assertTrue(self.upload('preview', bad).json['errors'])

    def test_summaries_public_boundaries_and_role_module_permissions(self):
        self.stage({'M': 7})
        page = self.call('inventory?page=1&pageSize=25').json
        self.assertEqual((page['summary']['temporaryInbound'], page['summary']['availableStock']), (7, 90))
        dashboard = self.call('dashboard?view=workbench').json
        self.assertEqual(dashboard['snapshot']['temporaryInbound'], 7)
        styles = self.call('dashboard?view=style-performance').json
        self.assertEqual(styles['summary']['temporaryInbound'], 7)
        detail = self.call('dashboard?view=style-detail&styleCode=GT-2026').json
        self.assertEqual(detail['items'][0]['temporaryInbound'], 7)
        for path in ['inventory', 'inventory/1', 'products', 'products/1']:
            response = self.call(path, role='customer')
            if response.status_code == 200: self.assertNotIn('temporaryInbound', response.get_data(as_text=True))
        for role in ['admin', 'sales', 'warehouse']:
            self.assertEqual(self.stage({'M': 7}).status_code, 200)
            self.assertEqual(self.receive(role=role).status_code, 200)
        self.assertEqual(self.receive(role='customer').status_code, 403)
        db._fetch_one("UPDATE admin_users SET permissions='[]' WHERE role='warehouse' RETURNING id")
        self.assertEqual(self.receive(role='warehouse').status_code, 403)

    def test_integer_overflow_rolls_back_all_sizes_and_audit(self):
        self.stage({'M': 8, 'L': 3}, sizeStocks={'M': 2147483647, 'L': -100, 'Tall XL': 0})
        before = self.detail()
        audit_count = self.call('audit-logs?page=1').json['total']
        self.assertEqual(self.receive().status_code, 400)
        self.assertEqual(self.detail()['stock'], before['stock'])
        self.assertEqual(self.sizes()['L']['temporaryInbound'], 3)
        self.assertEqual(self.call('audit-logs?page=1').json['total'], audit_count)
        self.stage({'M': 1, 'L': 0}, sizeStocks={'M': 2147483647, 'L': -1})
        self.assertEqual(self.receive().status_code, 400)  # per-size overflow despite valid parent sum
        self.stage({'M': 1, 'L': 0}, sizeStocks={'M': 2147483646, 'L': 1})
        self.assertEqual(self.receive().status_code, 400)  # parent overflow despite valid per-size sum

    def test_storefront_purchases_only_actual_stock_before_and_after_receipt(self):
        store_path = Path(__file__).resolve().parents[2] / 'shopify/storefront-backend/db.py'
        if not store_path.exists():
            self.skipTest('Sibling storefront checkout required for cross-service regression')
        spec = importlib.util.spec_from_file_location('temporary_checkout_store_db', store_path)
        store = importlib.util.module_from_spec(spec); spec.loader.exec_module(store)
        self.assertEqual(self.stage({'M': 23}, sizeStocks={'M': -20, 'L': 10, 'Tall XL': 0}).status_code, 200)
        public = next(p for p in store.list_products() if p['id'] == 1)
        self.assertEqual(public['stock'], 10)
        self.assertNotIn('temporaryInbound', repr(public))
        payload = {'userId': 1, 'contactName': 'Test Buyer', 'phone': '123456',
                   'shippingAddress': 'Test address', 'items': [{'productId': 1, 'sizeCode': 'M', 'quantity': 1}],
                   'allowNegativeStock': True, 'temporaryInbound': 999}
        with self.assertRaisesRegex(RuntimeError, 'Insufficient stock'):
            store.create_order(payload)
        self.assertEqual(self.sizes()['M']['temporaryInbound'], 23)
        payload['items'][0]['sizeCode'] = 'L'
        store.create_order(payload)  # Positive size remains buyable with a negative parent balance.
        self.assertEqual(self.detail()['stock'], -11)
        self.assertEqual(self.receive().status_code, 200)
        self.assertEqual(self.sizes()['M']['stock'], 3)
        payload['items'][0].update(sizeCode='M', quantity=3)
        store.create_order(payload)
        self.assertEqual((self.sizes()['M']['stock'], self.sizes()['M']['temporaryInbound']), (0, 0))
        self.assertEqual(self.detail()['stock'], sum(s['stock'] for s in self.sizes().values()))

    def test_both_migrations_keep_values_and_old_schema_starts_with_zero(self):
        store_path = Path(__file__).resolve().parents[2] / 'shopify/storefront-backend/db.py'
        migrations = [db._apply_schema_migrations]
        # Single-repository CI still exercises the admin migration; paired local
        # checkouts additionally exercise both services against the same database.
        if store_path.exists():
            spec = importlib.util.spec_from_file_location('temporary_store_db', store_path)
            store = importlib.util.module_from_spec(spec); spec.loader.exec_module(store)
            migrations.append(store._apply_schema_migrations)
            self.assertEqual(store_path.with_name('inventory_policy.py').read_bytes(), Path(db.inventory_policy.__file__).read_bytes())
        self.stage({'M': 9}, sizeStocks={'M': -9})
        before = db._fetch_all('SELECT product_id,size_code,stock,contract_pending,pending_inspection,pending_inbound,temporary_inbound FROM product_size_prices ORDER BY product_id,size_code')
        with db._connect() as conn:
            cur = conn.cursor()
            for migrate in migrations * 2: migrate(cur)
        self.assertEqual(db._fetch_all('SELECT product_id,size_code,stock,contract_pending,pending_inspection,pending_inbound,temporary_inbound FROM product_size_prices ORDER BY product_id,size_code'), before)
        with self.assertRaises(CheckViolation):
            with db._connect() as conn: conn.execute('UPDATE product_size_prices SET temporary_inbound=-1 WHERE product_id=1')
        for migrate in migrations:
            with db._connect() as conn:
                cur = conn.cursor(); cur.execute('ALTER TABLE product_size_prices DROP COLUMN temporary_inbound')
                migrate(cur)
                self.assertEqual(conn.execute('SELECT SUM(temporary_inbound) AS n FROM product_size_prices').fetchone()['n'], 0)
                self.assertEqual(conn.execute("SELECT stock FROM product_size_prices WHERE product_id=1 AND size_code='M'").fetchone()['stock'], -9)
                conn.rollback()


if __name__ == '__main__': unittest.main()
