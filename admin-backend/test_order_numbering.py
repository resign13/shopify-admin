import unittest
from datetime import datetime, timezone
from concurrent.futures import ThreadPoolExecutor
from openpyxl import load_workbook
import order_numbering as numbering
import test_orders as helpers
from test_workbench import db, application

class NumberFormatTest(unittest.TestCase):
    def test_dates_unicode_spaces_and_sequence(self):
        for instant, date in [('2026-10-12T15:59:59+00:00','1012'),('2026-10-12T16:00:00+00:00','1013'),('2026-12-31T16:00:00+00:00','0101')]:
            name='张三 Élodie-Marie  Vargo'
            self.assertEqual(numbering.make_order_number(1000001,datetime.fromisoformat(instant),name),'LM-1000001')
        self.assertEqual(numbering.sheet_title('PO1013-Élodie-Marie-000188'),'PO1013-000188')
        self.assertEqual(numbering.sheet_title('LM-000188'),'LM-000188')
        self.assertEqual(numbering.make_order_number(1,None,None),'LM-000001')

class OrderNumberIntegrationTest(unittest.TestCase):
    setUp=helpers.AdminOrdersTest.setUp
    call=helpers.AdminOrdersTest.call
    payload=helpers.AdminOrdersTest.payload
    create=helpers.AdminOrdersTest.create
    stock=helpers.AdminOrdersTest.stock

    def check_number(self, oid, name):
        row=db._fetch_one('SELECT id,order_no,created_at FROM orders WHERE id=%s',(oid,))
        self.assertEqual(row['order_no'],numbering.make_order_number(oid,row['created_at'],name))
        return row['order_no']

    def test_admin_long_number_search_immutable_and_exports(self):
        body=self.payload(); body['contactName']='Jonathan  Vargo Élodie 张三-' + '长姓名'*25
        body['secondaryOrderNo']='PO-'+body['contactName']
        response=self.call('orders','POST',body); self.assertEqual(response.status_code,200,response.json)
        oid=response.json['order']['id']; number=self.check_number(oid,body['contactName'])
        found=self.call('orders?keyword='+__import__('urllib.parse',fromlist=['quote']).quote(number)).json
        self.assertIn(oid,[o['id'] for o in found['items']])
        order=self.call(f'orders/{oid}').json['order']
        changed={**order,'contactName':'New name'}
        self.assertEqual(self.call(f'orders/{oid}/details','PUT',changed).status_code,200)
        self.assertEqual(self.call(f'orders/{oid}').json['order']['orderNo'],number)
        for builder in [application.build_orders_export,application.build_orders_sheet_export]:
            with application.app.test_request_context('/api/admin/orders/export'):
                sheet=load_workbook(builder([order],include_images=False)).active
                self.assertIn(number,sheet['A1'].value)
                self.assertTrue(sheet['A1'].alignment.wrap_text)
                self.assertGreater(sheet.row_dimensions[1].height,28)
                if builder==application.build_orders_sheet_export:self.assertEqual(sheet.title,numbering.sheet_title(number))
        with application.app.test_request_context('/api/admin/orders/export'):
            sheet=load_workbook(application.build_order_invoice_export(order)).active
            self.assertEqual(sheet['B5'].value,number+'\n'+body['secondaryOrderNo'])
            self.assertTrue(sheet['B5'].alignment.wrap_text)
            self.assertGreater(sheet.row_dimensions[5].height,24)

    def test_legacy_entry_and_migration_keep_historical_numbers(self):
        old=self.create(); db._fetch_one("UPDATE orders SET order_no='LM-000188' WHERE id=%s RETURNING id",(old['id'],))
        with db._connect() as conn:
            conn.execute('ALTER TABLE orders ALTER COLUMN order_no TYPE VARCHAR(40)')
            numbering.migrate(conn.cursor()); numbering.migrate(conn.cursor())
            self.assertEqual(conn.execute("SELECT character_maximum_length FROM information_schema.columns WHERE table_schema=current_schema() AND table_name='orders' AND column_name='order_no'").fetchone()['character_maximum_length'],255)
        self.assertEqual(self.call(f"orders/{old['id']}").json['order']['orderNo'],'LM-000188')
        result=db.create_order({'userId':1,'contactName':'Legacy 客户','phone':'123','shippingAddress':'Street','productId':1,'sizeCode':'M','quantity':1})
        self.check_number(result['id'],'Legacy 客户')

    def test_concurrent_distinct_and_replay(self):
        bodies=[self.payload() for _ in range(3)]
        with ThreadPoolExecutor(max_workers=3) as pool:
            results=list(pool.map(lambda b:self.call('orders','POST',b),bodies))
        self.assertTrue(all(r.status_code==200 for r in results),[r.json for r in results])
        numbers=[self.check_number(r.json['order']['id'],'Alex') for r in results]
        self.assertEqual(len(set(numbers)),3)
        replay=self.call('orders','POST',bodies[0]);self.assertEqual(replay.json['order']['orderNo'],numbers[0])

    def test_reorder_and_draft_use_final_customer(self):
        source=self.create(); result=self.call(f"orders/{source['id']}/reorder-draft").json
        body={**result['draft'],'sourceOrderId':source['id'],'requestId':self.payload()['requestId'],'ownerConfirmed':True,'contactName':'返单 客户'}
        response=self.call('orders','POST',body);self.assertEqual(response.status_code,200,response.json)
        self.check_number(response.json['order']['id'],'返单 客户')
        before=db._fetch_one('SELECT last_value FROM orders_id_seq')['last_value']
        body=self.payload(); saved=self.call('orders/drafts','POST',{'draftKey':body['requestId'],'draft':body}).json
        self.assertEqual(db._fetch_one('SELECT last_value FROM orders_id_seq')['last_value'],before)
        loaded=self.call(f"orders/drafts/{saved['id']}").json['draft']
        submitted=self.call('orders','POST',{**loaded,'contactName':'Final Customer','draftId':saved['id'],'draftVersion':saved['version'],'requestId':saved['submitRequestId']})
        self.assertEqual(submitted.status_code,200,submitted.json)
        self.check_number(submitted.json['order']['id'],'Final Customer')

    def test_po_migration_retains_secondary_and_business_state(self):
        order=self.create();oid=order['id'];original=f'PO1013-Alex Buyer-{oid:06d}'
        db._fetch_one('UPDATE orders SET order_no=%s WHERE id=%s RETURNING id',(original,oid))
        stocks=self.stock()
        sequence=db._fetch_one('SELECT last_value FROM orders_id_seq')['last_value']
        notifications=db._fetch_one('SELECT count(*) AS n FROM order_created_events')['n']
        with db._connect() as conn:
            numbering.migrate(conn.cursor());numbering.migrate(conn.cursor())
        order=self.call(f'orders/{oid}').json['order']
        self.assertEqual(order['orderNo'],f'LM-{oid:06d}')
        self.assertEqual(order['secondaryOrderNo'],original)
        self.assertEqual(self.stock(),stocks)
        self.assertEqual(db._fetch_one('SELECT last_value FROM orders_id_seq')['last_value'],sequence)
        self.assertEqual(db._fetch_one('SELECT count(*) AS n FROM order_created_events')['n'],notifications)
        found=self.call('orders?keyword='+original).json
        self.assertEqual(found['items'][0]['id'],oid)

    def test_secondary_create_edit_clear_and_permissions(self):
        body=self.payload();body['secondaryOrderNo']='po客户 A-2026'
        response=self.call('orders','POST',body)
        self.assertEqual(response.status_code,200,response.json)
        order=response.json['order'];oid=order['id']
        order=self.call(f'orders/{oid}').json['order']
        self.assertEqual(order['secondaryOrderNo'],'PO客户 A-2026')
        self.assertTrue(self.call('orders','POST',body).json['replayed'])
        stocks=self.stock()
        legacy={**order,'note':'Legacy edit'};legacy.pop('secondaryOrderNo')
        self.assertEqual(self.call(f'orders/{oid}/details','PUT',legacy).status_code,200)
        order=self.call(f'orders/{oid}').json['order']
        self.assertEqual(order['secondaryOrderNo'],'PO客户 A-2026')
        self.assertEqual(self.call(f'orders/{oid}/details','PUT',{**order,'secondaryOrderNo':'POdifferent'},'warehouse').status_code,403)
        self.assertEqual(self.call(f'orders/{oid}/details','PUT',{**order,'secondaryOrderNo':''}).status_code,200)
        cleared=self.call(f'orders/{oid}').json['order']
        self.assertEqual(cleared['secondaryOrderNo'],'');self.assertEqual(cleared['orderNo'],order['orderNo'])
        self.assertEqual(self.stock(),stocks)
        operations=self.call(f'orders/{oid}/operations').json['items']
        self.assertTrue(any(c['field']=='secondaryOrderNo' for r in operations for c in r['changes']))

    def test_secondary_draft_restores_but_reorder_starts_empty(self):
        body=self.payload();body['secondaryOrderNo']='PO-DRAFT-1'
        saved=self.call('orders/drafts','POST',{'draftKey':body['requestId'],'draft':body}).json
        loaded=self.call(f"orders/drafts/{saved['id']}").json['draft']
        self.assertEqual(loaded['secondaryOrderNo'],'PO-DRAFT-1')
        loaded.pop('secondaryOrderNo')
        updated=self.call(f"orders/drafts/{saved['id']}",'PUT',{'version':saved['version'],'draft':loaded})
        self.assertEqual(updated.status_code,200,updated.json)
        self.assertEqual(self.call(f"orders/drafts/{saved['id']}").json['draft']['secondaryOrderNo'],'PO-DRAFT-1')
        response=self.call('orders','POST',{**loaded,'draftId':saved['id'],'draftVersion':updated.json['version'],'requestId':saved['submitRequestId']})
        self.assertEqual(response.status_code,200,response.json)
        oid=response.json['order']['id']
        self.assertEqual(response.json['order']['secondaryOrderNo'],'PO-DRAFT-1')
        self.assertEqual(self.call(f'orders/{oid}/reorder-draft').json['draft']['secondaryOrderNo'],'')

    def test_bad_secondary_fails_without_order_or_stock_write(self):
        stocks=self.stock();count=db._fetch_one('SELECT count(*) AS n FROM orders')['n']
        for value in (123,True,'PO'+('a'*254),'POline\nother'):
            body=self.payload();body['secondaryOrderNo']=value
            self.assertEqual(self.call('orders','POST',body).status_code,400)
        self.assertEqual(self.stock(),stocks)
        self.assertEqual(db._fetch_one('SELECT count(*) AS n FROM orders')['n'],count)

if __name__=='__main__':unittest.main()
