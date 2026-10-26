"""Independent measurements and workbook lifecycle; guarded *_test fixtures only."""
from copy import deepcopy
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
import unittest
import uuid
from unittest.mock import patch
from openpyxl import Workbook, load_workbook
from psycopg.types.json import Jsonb

import test_workbench as fixtures
import size_charts as charts
import size_chart_excel as excel

db, application = fixtures.db, fixtures.application


class SizeChartsTest(unittest.TestCase):
    call = fixtures.WorkbenchTest.call

    def setUp(self):
        self.tokens, self.products = fixtures.seed()
        self.app = application.app
        self.app.config['TESTING'] = True
        with db.get_connection() as c: c.execute('TRUNCATE product_size_charts,product_size_chart_audit,product_size_chart_receipts RESTART IDENTITY')

    def detail(self):
        response = self.call('size-charts/detail?styleCode=GT-2026')
        self.assertEqual(response.status_code, 200, response.json)
        return response.json

    def save(self, doc, **extra):
        return self.call('size-charts/save', 'POST', {'requestId': str(uuid.uuid4()), 'changes': [doc], **extra})

    def workbook(self, mode='maintenance', codes=None, unit=None):
        response = self.call('size-charts/export', 'POST', {'styleCodes': codes or ['GT-2026'], 'mode': mode, 'unit':unit})
        self.assertEqual(response.status_code, 200, response.json if response.is_json else '')
        return load_workbook(BytesIO(response.data))

    def preview(self, book):
        stream=BytesIO();book.save(stream);stream.seek(0)
        response=self.app.test_client().post('/api/admin/size-charts/import/preview',data={'file':(stream,'sizes.xlsx')},headers={'Authorization':'Bearer '+self.tokens['admin']})
        return response

    def commit(self, data, rid=None):
        return self.call('size-charts/import/confirm','POST',{'proposal':data['proposal'],'previewToken':data['previewToken'],'requestId':rid or str(uuid.uuid4())})

    def test_grouped_sizes_and_default_permissions(self):
        data=self.call('size-charts').json
        self.assertEqual(data['total'],1)
        self.assertEqual(self.detail()['sizes'],['M','L','Tall XL'])
        self.assertEqual(self.call('size-charts',role='sales').status_code,200)
        for role in ['warehouse','customer']:
            self.assertEqual(self.call('size-charts',role=role).status_code,403)
        db._fetch_one("UPDATE admin_users SET permissions=%s WHERE role='admin' RETURNING id",(Jsonb(['products']),))
        self.assertEqual(self.call('size-charts').status_code,403)

    def test_explicit_grant_and_revocation(self):
        db._fetch_one("UPDATE admin_users SET permissions=%s WHERE role='sales' RETURNING id",(Jsonb(['size-charts']),))
        self.assertEqual(self.call('size-charts',role='sales').status_code,200)
        doc=self.detail();doc['values']['hip']['M']='96'
        response=self.call('size-charts/save','POST',{'requestId':str(uuid.uuid4()),'changes':[doc]},role='sales')
        self.assertEqual(response.status_code,200,response.json)
        self.assertEqual(self.call('products',role='sales').status_code,403)

    def test_admin_and_sales_defaults_and_explicit_revocation(self):
        import module_permissions
        for role in ['admin','sales']:
            self.assertIn('size-charts',module_permissions.effective({'role':role,'permissions':None}))
            self.assertNotIn('size-charts',module_permissions.effective({'role':role,'permissions':[]}))
        for role in ['warehouse','finance','customer']:
            self.assertNotIn('size-charts',module_permissions.effective({'role':role,'permissions':None}))
        doc=self.detail();doc['values']['hip']['M']='90'
        self.assertEqual(self.call('size-charts/save','POST',{'requestId':str(uuid.uuid4()),'changes':[doc]},role='sales').status_code,200)
        db._fetch_one("UPDATE admin_users SET permissions=%s WHERE role='sales' RETURNING id",(Jsonb(['orders']),))
        self.assertEqual(self.call('size-charts',role='sales').status_code,403)

    def test_save_zero_noop_idempotency_and_product_unchanged(self):
        before=db._fetch_all('SELECT id,price,stock,size_chart_image_url FROM products ORDER BY id')
        doc=self.detail();doc['values']['hip']['M']='0';doc['values']['hip']['L']='96.25'
        rid=str(uuid.uuid4());r=self.save(doc,requestId=rid)
        self.assertEqual(r.status_code,200,r.json);self.assertEqual(r.json['items'][0]['values']['hip']['M'],'0.00')
        self.assertEqual(self.save(doc,requestId=rid).json['replayed'],True)
        new=self.detail();date=new['updatedDate'];version=new['version']
        self.assertEqual(self.save(new).json['updated'],[])
        self.assertEqual(self.detail()['version'],version);self.assertEqual(self.detail()['updatedDate'],date)
        self.assertEqual(before,db._fetch_all('SELECT id,price,stock,size_chart_image_url FROM products ORDER BY id'))
        self.assertEqual(db._fetch_one('SELECT count(*) AS n FROM product_size_chart_audit')['n'],1)

    def test_blank_unsaved_import_is_noop(self):
        r=self.preview(self.workbook());self.assertEqual(r.status_code,200,r.json)
        self.assertEqual(r.json['changedCount'],0)
        self.assertEqual(self.commit(r.json).json['updated'],[])
        self.assertEqual(self.detail()['updatedDate'],'')

    def test_import_overwrite_clear_zero_and_date(self):
        doc=self.detail();doc['values']['waist-relaxed']['M']='30';self.save(doc)
        book=self.workbook();sheet=book.worksheets[0]
        self.assertEqual(sheet['D9'].value,30)
        self.assertEqual(len(sheet._images),1)
        sheet['D9']=None;sheet['E9']=0;sheet['F9']=32.25;sheet['C4']='1999-01-01'
        result=self.preview(book);self.assertEqual(result.status_code,200,result.json);self.assertEqual(result.json['clearCount'],1)
        rid=str(uuid.uuid4());saved=self.commit(result.json,rid);self.assertEqual(saved.status_code,200,saved.json)
        self.assertEqual(self.commit(result.json,rid).json['replayed'],True)
        detail=self.detail();self.assertEqual(detail['values']['waist-relaxed'],{'M':None,'L':'0.00','Tall XL':'32.25'})
        self.assertNotEqual(detail['updatedDate'],'1999-01-01')

    def test_english_hides_dimensions_without_deleting_values(self):
        doc=self.detail();doc.update(titleEn='Slim fit trousers',material='聚酯纤维',materialEn='100% Polyester')
        doc['values']['knee']['M']='40'
        for d in doc['dimensions']:d['showEn']=d['id']=='hip'
        self.assertEqual(self.save(doc).status_code,200)
        book=self.workbook('english');self.assertEqual(len(book.sheetnames),1)
        words=[str(c.value or '') for row in book.active for c in row]
        self.assertIn('Hip',words);self.assertNotIn('Knee',words);self.assertNotIn('膝围',words)
        self.assertEqual(self.detail()['values']['knee']['M'],'40.00')
        self.assertEqual(self.preview(book).status_code,400)

    def test_custom_dimensions_and_exact_mixed_sizes(self):
        with db.get_connection() as c:
            c.execute("INSERT INTO product_sizes(product_id,size_code,sort_order) VALUES(1,'S/28',9),(2,'28',8)")
        doc=self.detail();self.assertEqual(doc['sizes'],['M','L','Tall XL','S/28','28'])
        doc['dimensions'].insert(0,{'id':'outseam','name':'外长','nameEn':'Outseam','unit':'cm','showEn':True})
        doc['values']['outseam']={s:None for s in doc['sizes']};doc['values']['outseam']['S/28']='100';doc['values']['outseam']['28']='90'
        self.assertEqual(self.save(doc).status_code,200)
        self.assertEqual(self.detail()['dimensions'][0]['id'],'outseam')
        self.assertEqual(self.detail()['values']['outseam']['28'],'90.00')

    def test_removed_size_retained_and_old_template_rejected(self):
        doc=self.detail();doc['values']['hip']['M']='90';self.save(doc);book=self.workbook()
        with db.get_connection() as c:c.execute("DELETE FROM product_sizes WHERE size_code='M'")
        new=self.detail();self.assertNotIn('M',new['sizes'])
        new['material']='cotton';self.assertEqual(self.save(new).status_code,200)
        self.assertEqual(db._fetch_one("SELECT document FROM product_size_charts WHERE style_code='GT-2026'")['document']['values']['hip']['M'],'90.00')
        result=self.preview(book);self.assertTrue(result.json['errors']);self.assertIsNone(result.json['proposal'])

    def test_excel_errors_and_structure(self):
        for value in ['=1+2',-1,'bad',1.123]:
            book=self.workbook();book.active['D9']=value;r=self.preview(book)
            self.assertEqual(r.status_code,200,r.json);self.assertTrue(r.json['errors']);self.assertIsNone(r.json['proposal'])
        book=self.workbook();book.active.delete_rows(9)
        deleted=self.preview(book).json
        self.assertEqual(deleted['errors'],[])
        self.assertEqual(deleted['removedDimensionCount'],1)
        book=self.workbook();book.active['D8']='WRONG';self.assertTrue(self.preview(book).json['errors'])

    def test_conflict_on_confirm_and_forged_preview(self):
        book=self.workbook();book.active['D9']=30;r=self.preview(book).json
        doc=self.detail();doc['values']['hip']['M']='99';self.save(doc)
        self.assertEqual(self.commit(r).status_code,409)
        r['proposal']['changes'][0]['values']['hip']['M']='500'
        self.assertEqual(self.commit(r).status_code,400)
        self.assertEqual(self.detail()['values']['waist-relaxed']['M'],None)

    def test_atomic_multi_style_and_duplicate_rejection(self):
        with db.get_connection() as c:c.execute("UPDATE products SET style_code='OTHER' WHERE id=3")
        a=self.detail();b=self.call('size-charts/detail?styleCode=OTHER').json
        a['values']['hip']['M']='90';b['values']['hip']['M']='-1'
        r=self.call('size-charts/save','POST',{'requestId':str(uuid.uuid4()),'changes':[a,b]})
        self.assertEqual(r.status_code,400);self.assertEqual(self.detail()['version'],'0')
        b['values']['hip']['M']='91'
        r=self.call('size-charts/save','POST',{'requestId':str(uuid.uuid4()),'changes':[a,b]});self.assertEqual(r.status_code,200,r.json)
        book=self.workbook(codes=['GT-2026','OTHER']);self.assertEqual(len(book.worksheets),3)
        book.remove(book.worksheets[1]);book.active['D9']=25
        p=self.preview(book);self.assertEqual(p.json['styleCount'],1)
        self.assertEqual(self.commit(p.json).status_code,200)
        self.assertEqual(self.call('size-charts/detail?styleCode=OTHER').json['values']['hip']['M'],'91.00')

    def test_image_validation_and_read_failure(self):
        doc=self.detail();first=doc['image'];doc['image']='http://127.0.0.1/private'
        self.assertEqual(self.save(doc).status_code,200)
        self.assertEqual(self.detail()['image'],first)
        with patch.object(application,'fetch_image_bytes',return_value=None):
            r=self.call('size-charts/export','POST',{'styleCodes':['GT-2026']});self.assertEqual(r.status_code,400)

    def test_representative_image_always_follows_first_catalog_image(self):
        doc=self.detail();doc['material']='cotton';self.save(doc)
        with db.get_connection() as c:
            c.execute("UPDATE product_size_charts SET document=jsonb_set(document,'{image}', '\"/uploads/old-choice.jpg\"')")
            c.execute("UPDATE products SET main_image_url='/uploads/new-first.jpg' WHERE id=1")
        self.assertEqual(self.detail()['image'],'/uploads/new-first.jpg')
        self.assertEqual(self.call('size-charts').json['items'][0]['image'],'/uploads/new-first.jpg')
        with patch.object(application,'fetch_image_bytes',wraps=application.fetch_image_bytes) as fetch:
            self.call('size-charts/export','POST',{'styleCodes':['GT-2026']})
            fetch.assert_called_once_with('/uploads/new-first.jpg')

    def test_concurrent_first_save_has_one_winner(self):
        a=self.detail();b=deepcopy(a)
        a['values']['hip']['M']='90';b['values']['hip']['M']='95'
        with ThreadPoolExecutor(max_workers=2) as pool:
            results=list(pool.map(self.save,[a,b]))
        self.assertEqual(sorted(r.status_code for r in results),[200,409])
        self.assertEqual(self.detail()['version'],'1')
        self.assertEqual(db._fetch_one('SELECT count(*) AS n FROM product_size_chart_audit')['n'],1)

    def test_failed_audit_rolls_back_chart_write(self):
        doc=self.detail();doc['values']['hip']['M']='90'
        original=db._fetch_one
        def fail_audit(query,*args,**kwargs):
            if 'INSERT INTO product_size_chart_audit' in query:raise ValueError('fixture transaction failure')
            return original(query,*args,**kwargs)
        with patch.object(db,'_fetch_one',side_effect=fail_audit):self.assertEqual(self.save(doc).status_code,400)
        self.assertEqual(self.detail()['version'],'0')
        self.assertEqual(db._fetch_one('SELECT count(*) AS n FROM product_size_chart_receipts')['n'],0)

    def test_excel_gridlines_visible_without_frozen_or_split_panes(self):
        doc=self.detail();doc['titleEn']='Size chart fixture'
        self.assertEqual(self.save(doc).status_code,200)
        for mode in ('maintenance', 'english'):
            with self.subTest(mode=mode):
                book=self.workbook(mode)
                for sheet in book:
                    if sheet.sheet_state != 'visible': continue
                    self.assertTrue(sheet.sheet_view.showGridLines)
                    self.assertIsNone(sheet.freeze_panes)
                    self.assertIsNone(sheet.sheet_view.pane)
                    self.assertEqual(len(sheet._images),1)
                if mode == 'maintenance':
                    preview=self.preview(book)
                    self.assertEqual(preview.status_code,200,preview.json)
                    self.assertEqual(preview.json['errors'],[])
                    self.assertEqual(preview.json['changedCount'],0)

    def test_session_and_csrf(self):
        from flask.testing import FlaskClient
        raw=FlaskClient(self.app)
        self.assertEqual(raw.get('/api/admin/size-charts').status_code,401)
        raw.set_cookie(application.security.cookie,self.tokens['admin'])
        self.assertEqual(raw.post('/api/admin/size-charts/save',json={},headers={'Origin':'http://localhost'}).status_code,403)

    def test_unified_units_default_validation_and_legacy_read(self):
        doc=self.detail()
        self.assertEqual(doc['unit'],'cm')
        self.assertTrue(all(d['unit']=='cm' for d in doc['dimensions']))
        doc['dimensions'][0]['unit']='in'
        self.assertEqual(self.save(doc).status_code,400)
        self.assertEqual(self.detail()['version'],'0')
        # Legacy mixed-unit records are interpreted in cm without a database write.
        legacy=charts.defaults(charts.catalog()['GT-2026']);legacy.pop('unit')
        legacy['dimensions'][0]['unit']='in'
        legacy['values']['waist-relaxed'].update(M='30',L='0',**{'Tall XL':None})
        legacy['values']['hip']['M']='96'
        db._fetch_one('INSERT INTO product_size_charts(style_code,document) VALUES(%s,%s) RETURNING revision',('GT-2026',Jsonb(legacy)))
        doc=self.detail()
        self.assertEqual(doc['values']['waist-relaxed'],{'M':'76.20','L':'0.00','Tall XL':None})
        self.assertEqual(doc['values']['hip']['M'],'96.00')
        self.assertTrue(all(d['unit']=='cm' for d in doc['dimensions']))
        self.assertEqual(db._fetch_one('SELECT document FROM product_size_charts')['document'],legacy)
        self.assertEqual(self.save(doc).status_code,200)
        self.assertEqual(self.detail()['values'],doc['values'])
        self.assertEqual(self.save(self.detail()).json['updated'],[])

    def test_unit_conversion_save_retained_sizes_and_excel_conflict(self):
        doc=self.detail();doc['values']['hip']['M']='2.54';doc['values']['hip']['L']='25.4'
        self.save(doc);book=self.workbook()
        with db.get_connection() as c:c.execute("DELETE FROM product_sizes WHERE size_code='M'")
        doc=charts.unified_document(self.detail(),'in')
        self.assertEqual(doc['values']['hip']['L'],'10.00')
        self.assertEqual(self.save(doc).status_code,200)
        stored=db._fetch_one('SELECT document FROM product_size_charts')['document']
        self.assertEqual(stored['values']['hip']['M'],'1.00')
        self.assertEqual(self.detail()['unit'],'in')
        self.assertTrue(self.preview(book).json['errors'])
        new=self.workbook();preview=self.preview(new)
        self.assertEqual(preview.json['errors'],[])
        self.assertEqual(preview.json['changedCount'],0)
        self.assertEqual(charts.convert_value('0.25','in','cm'),'0.64')
        with self.assertRaises(ValueError):charts.convert_value('9999999999.99','in','cm')

    def test_compact_excel_layout_and_old_layout_import(self):
        doc=self.detail();doc['titleEn']='Size chart';self.save(doc)
        for mode in ('maintenance','english'):
            book=self.workbook(mode);sheet=book.active
            self.assertEqual(sheet['C2'].value,'GT-2026')
            self.assertEqual(sheet['C5'].value,'cm')
            self.assertEqual(sheet['A8'].value,'Measurement' if mode=='english' else '测量维度')
            self.assertEqual(sheet['D9'].border.bottom.style,'thin')
            self.assertLess(sheet.row_dimensions[9].height,32)
            self.assertEqual(sheet._images[0].anchor._from.row,0)
            self.assertNotIn('END',[c.value for row in sheet for c in row])
        book=self.workbook();book.active.delete_rows(17)
        self.assertEqual(self.preview(book).json['removedDimensionCount'],1)
        book=self.workbook();book.active['C5']='in'
        self.assertTrue(self.preview(book).json['errors'])
        # Previously downloaded v1 files keep their old cell coordinates.
        doc=self.detail();legacy=Workbook();s=legacy.active;s.title='001-GT-2026';s['B2']=doc['styleCode']
        for col,value in enumerate(['测量维度','English','单位',*doc['sizes']],4):s.cell(8,col,value)
        for row,d in enumerate(doc['dimensions'],9):
            for col,value in enumerate([d['name'],d['nameEn'],d['unit']],4):s.cell(row,col,value)
        s.cell(9+len(doc['dimensions']),4,'END');s['G9']=12.5
        meta=legacy.create_sheet(excel.META);meta['A1']=excel.LEGACY_VERSION
        item={'sheet':s.title,**excel.structure(doc)};meta['A2']=excel.packed(item)
        with self.app.app_context():meta['B2']=excel.sign(item)
        p=self.preview(legacy);self.assertEqual(p.json['errors'],[],p.json)
        self.assertEqual(p.json['rows'][0]['cell'],'G9')
        self.assertEqual(self.commit(p.json).status_code,200)
        self.assertEqual(self.detail()['values']['waist-relaxed']['M'],'12.50')

    def test_excel_overwrite_preserves_each_whole_style_unit(self):
        for unit in ['in', 'cm']:
            with self.subTest(unit=unit):
                doc=charts.unified_document(self.detail(),unit)
                doc['values']['waist-relaxed'].update(M='12.34',L='1.25',**{'Tall XL':None})
                self.assertEqual(self.save(doc).status_code,200)
                book=self.workbook();sheet=book.active
                self.assertEqual(sheet['C5'].value,unit)
                sheet['D9']=None;sheet['E9']=0;sheet['F9']=23.45
                preview=self.preview(book)
                self.assertEqual(preview.json['errors'],[],preview.json)
                self.assertEqual(preview.json['clearCount'],1)
                rid=str(uuid.uuid4())
                self.assertEqual(self.commit(preview.json,rid).status_code,200)
                self.assertTrue(self.commit(preview.json,rid).json['replayed'])
                saved=self.detail()
                self.assertEqual(saved['unit'],unit)
                self.assertTrue(all(d['unit']==unit for d in saved['dimensions']))
                self.assertEqual(saved['values']['waist-relaxed'],{'M':None,'L':'0.00','Tall XL':'23.45'})
                unchanged=self.preview(self.workbook()).json
                self.assertEqual(unchanged['changedCount'],0)
                self.assertEqual(self.commit(unchanged).json['updated'],[])
                self.assertEqual(self.detail()['version'],saved['version'])
                self.assertEqual(self.detail()['updatedDate'],saved['updatedDate'])

    def test_multi_style_excel_errors_and_late_conflict_never_partially_apply(self):
        with db.get_connection() as c:c.execute("UPDATE products SET style_code='OTHER' WHERE id=3")
        book=self.workbook(codes=['GT-2026','OTHER'])
        book.worksheets[0]['D9']=21.25;book.worksheets[1]['D9']='=1+2'
        preview=self.preview(book).json
        self.assertTrue(preview['errors'])
        self.assertIn('D9',preview['errors'][0]['message'])
        self.assertIsNone(preview['proposal'])
        self.assertEqual(self.detail()['version'],'0')
        self.assertEqual(self.call('size-charts/detail?styleCode=OTHER').json['version'],'0')
        # Saving an openpyxl workbook consumes its embedded-image stream.
        # Upload a fresh export, as a real second upload would do.
        book=self.workbook(codes=['GT-2026','OTHER'])
        book.worksheets[0]['D9']=21.25;book.worksheets[1]['D9']=22.25
        preview=self.preview(book).json
        self.assertEqual(preview['errors'],[])
        other=self.call('size-charts/detail?styleCode=OTHER').json
        other['values']['hip']['M']='99'
        self.assertEqual(self.save(other).status_code,200)
        self.assertEqual(self.commit(preview).status_code,409)
        self.assertEqual(self.detail()['version'],'0')
        self.assertIsNone(self.detail()['values']['waist-relaxed']['M'])
        other=self.call('size-charts/detail?styleCode=OTHER').json
        self.assertEqual(other['values']['hip']['M'],'99.00')
        self.assertIsNone(other['values']['waist-relaxed']['M'])

    def test_preview_unit_exports_do_not_save_or_lose_original_precision(self):
        doc=self.detail();doc['titleEn']='Size chart';doc['values']['hip'].update(M='2.54',L='100.50',**{'Tall XL':'0'})
        self.save(doc);original=self.detail()
        before=db._fetch_all('SELECT * FROM product_size_chart_audit ORDER BY id')
        for mode in ['maintenance','english']:
            book=self.workbook(mode,unit='in');sheet=book.active
            self.assertEqual(sheet['C5'].value,'in')
            self.assertEqual(sheet['D11' if mode=='maintenance' else 'C11'].value,1)
            self.assertEqual(sheet['E11' if mode=='maintenance' else 'D11'].value,39.57)
            self.assertEqual(len(sheet._images),1)
            if mode=='maintenance':
                result=self.preview(book).json
                self.assertEqual(result['errors'],[],result)
                self.assertEqual(result['changedCount'],0)
                self.assertEqual(self.commit(result).json['updated'],[])
            else:self.assertEqual(len(book.worksheets),1)
        self.assertEqual(self.detail(),original)
        self.assertEqual(db._fetch_all('SELECT * FROM product_size_chart_audit ORDER BY id'),before)

    def test_converted_template_changes_return_to_saved_unit_with_clear_and_zero(self):
        for stored,exported in [('cm','in'),('in','cm')]:
            with self.subTest(stored=stored):
                doc=charts.unified_document(self.detail(),stored)
                doc['values']['hip'].update(M='100.50',L='2.54',**{'Tall XL':'3'})
                self.assertEqual(self.save(doc).status_code,200)
                book=self.workbook(unit=exported);sheet=book.active
                sheet['D11']=12.34;sheet['E11']=None;sheet['F11']=0
                preview=self.preview(book).json
                self.assertEqual(preview['errors'],[],preview)
                self.assertEqual(preview['clearCount'],1)
                self.assertTrue(all(row['unit']==stored for row in preview['rows']))
                self.assertEqual(self.commit(preview).status_code,200)
                saved=self.detail()
                self.assertEqual(saved['unit'],stored)
                self.assertEqual(saved['values']['hip'],{'M':charts.convert_value('12.34',exported,stored),'L':None,'Tall XL':'0.00'})

    def test_preview_export_unit_and_template_tampering_are_rejected(self):
        for unit in ['mm',{},True]:
            result=self.call('size-charts/export','POST',{'styleCodes':['GT-2026'],'unit':unit})
            self.assertEqual(result.status_code,400,result.json)
        book=self.workbook(unit='in');book.active['C5']='cm'
        self.assertTrue(self.preview(book).json['errors'])
        book=self.workbook(unit='in');book.active['C9']='cm'
        self.assertTrue(self.preview(book).json['errors'])
        book=self.workbook(unit='in')
        meta=book[excel.META];item=__import__('json').loads(meta['A2'].value);item['exportUnit']='cm';meta['A2']=excel.packed(item)
        self.assertEqual(self.preview(book).status_code,400)

    def test_excel_deleted_dimension_rows_preview_and_commit_preserve_remaining_data(self):
        doc=self.detail();doc['values']['waist-relaxed']['M']='30';doc['values']['knee']['M']='42'
        self.assertEqual(self.save(doc).status_code,200)
        before=self.detail()
        book=self.workbook();book.active.delete_rows(9,4)
        preview=self.preview(book).json
        self.assertEqual(preview['errors'],[],preview)
        self.assertEqual(preview['removedDimensionCount'],4)
        self.assertEqual(preview['dimensionChangedCount'],4)
        self.assertEqual(preview['valueChangedCount'],0)
        self.assertEqual(preview['dimensionRows'][0]['filledCount'],1)
        self.assertEqual(self.detail(),before,'Preview must not apply deletion')
        rid=str(uuid.uuid4())
        self.assertEqual(self.commit(preview,rid).status_code,200)
        self.assertTrue(self.commit(preview,rid).json['replayed'])
        saved=self.detail()
        self.assertEqual([d['id'] for d in saved['dimensions']],['knee','leg-opening','front-rise','back-rise','inseam'])
        self.assertEqual(saved['values']['knee']['M'],'42.00')
        stored=db._fetch_one('SELECT document FROM product_size_charts')['document']
        self.assertEqual(stored['values']['waist-relaxed']['M'],'30.00','Removed data remains in history')
        self.assertEqual(self.preview(self.workbook()).json['changedCount'],0)
        self.assertEqual(db._fetch_one('SELECT count(*) AS n FROM product_size_chart_audit')['n'],2)

    def test_excel_add_rename_reorder_and_retained_english_visibility(self):
        doc=self.detail();doc['titleEn']='Size chart';doc['values']['hip']['M']='96.25'
        next(d for d in doc['dimensions'] if d['id']=='hip')['showEn']=False
        self.assertEqual(self.save(doc).status_code,200)
        book=self.workbook();sheet=book.active
        # Move Hip above the waist rows and rename its Chinese label only.
        data=[c.value for c in sheet[11]];sheet.delete_rows(11);sheet.insert_rows(9)
        data[0]='臀围成衣'
        for col,value in enumerate(data,1):sheet.cell(9,col,value)
        sheet['B10']='Waist at rest'
        sheet.append(['外长','Outseam',None,100.5,0,None])
        preview=self.preview(book).json
        self.assertEqual(preview['errors'],[],preview)
        self.assertEqual(sorted(r['kind'] for r in preview['dimensionRows']),['add','rename','rename','reorder'])
        self.assertEqual(self.commit(preview).status_code,200)
        saved=self.detail()
        self.assertEqual(saved['dimensions'][0]['id'],'hip')
        self.assertEqual(saved['dimensions'][0]['name'],'臀围成衣')
        self.assertFalse(saved['dimensions'][0]['showEn'])
        self.assertEqual(saved['values']['hip']['M'],'96.25')
        self.assertEqual(saved['dimensions'][1]['nameEn'],'Waist at rest')
        new=saved['dimensions'][-1]
        self.assertEqual(new['nameEn'],'Outseam');self.assertTrue(new['showEn'])
        self.assertEqual(saved['values'][new['id']],{'M':'100.50','L':'0.00','Tall XL':None})
        english=self.workbook('english')
        labels=[c.value for row in english.active for c in row]
        self.assertNotIn('Hip',labels);self.assertIn('Outseam',labels)
        self.assertEqual(len(english.active._images),1)

    def test_excel_both_names_replaced_are_explicit_remove_and_add(self):
        doc=self.detail();doc['values']['hip']['M']='96';self.save(doc)
        book=self.workbook();sheet=book.active
        sheet['A11']='胸围';sheet['B11']='Chest';sheet['D11']=102
        preview=self.preview(book).json
        self.assertEqual(preview['errors'],[],preview)
        self.assertEqual([r['kind'] for r in preview['dimensionRows']],['remove','add'])
        self.assertEqual(self.commit(preview).status_code,200)
        saved=self.detail();chest=next(d for d in saved['dimensions'] if d['name']=='胸围')
        self.assertNotEqual(chest['id'],'hip')
        self.assertEqual(saved['values'][chest['id']]['M'],'102.00')
        self.assertEqual(db._fetch_one('SELECT document FROM product_size_charts')['document']['values']['hip']['M'],'96.00')

    def test_excel_dimension_validation_and_limits_never_write(self):
        for kind in ['empty','duplicate','missing-name','mixed-unit','invalid-name','too-many','formula','size-column']:
            with self.subTest(kind=kind):
                book=self.workbook();sheet=book.active
                if kind=='empty':sheet.delete_rows(9,9)
                elif kind=='duplicate':sheet.append(['臀围','Hip','cm',1,2,3])
                elif kind=='missing-name':sheet.append([None,'Outseam','cm',100])
                elif kind=='mixed-unit':sheet.append(['外长','Outseam','in',100])
                elif kind=='invalid-name':sheet.append(['a'*81,'Outseam','cm',100])
                elif kind=='too-many':
                    for i in range(56):sheet.append([f'自定义{i}',f'Custom {i}','cm',1])
                elif kind=='formula':sheet.append(['外长','Outseam','cm','=1+2'])
                else:sheet['D8']='OTHER SIZE'
                preview=self.preview(book).json
                self.assertTrue(preview['errors'],preview)
                self.assertIsNone(preview['proposal'])
                self.assertEqual(self.detail()['version'],'0')
        self.assertEqual(db._fetch_one('SELECT count(*) AS n FROM product_size_chart_audit')['n'],0)

    def test_converted_template_dimension_edits_preserve_precision_and_saved_unit(self):
        doc=self.detail();doc['values']['hip']['M']='100.50';self.save(doc)
        book=self.workbook(unit='in');sheet=book.active
        sheet.delete_rows(9,2);sheet['A9']='臀围成衣'
        sheet.append(['外长','Outseam','in',10,0,None])
        preview=self.preview(book).json
        self.assertEqual(preview['errors'],[],preview)
        self.assertEqual(preview['removedDimensionCount'],2)
        self.assertEqual(self.commit(preview).status_code,200)
        saved=self.detail()
        self.assertEqual(saved['unit'],'cm')
        self.assertEqual(saved['values']['hip']['M'],'100.50')
        new=saved['dimensions'][-1]
        self.assertEqual(saved['values'][new['id']]['M'],'25.40')
        self.assertTrue(all(d['unit']=='cm' for d in saved['dimensions']))

    def test_dimension_import_multi_style_conflict_and_tampering_are_atomic(self):
        with db.get_connection() as c:c.execute("UPDATE products SET style_code='OTHER' WHERE id=3")
        book=self.workbook(codes=['GT-2026','OTHER'])
        book.worksheets[0].delete_rows(9,4);book.worksheets[1].append(['外长','Outseam','cm',100])
        preview=self.preview(book).json
        self.assertEqual(preview['errors'],[],preview)
        modified=deepcopy(preview);modified['proposal']['changes'][0]['dimensions'][0]['name']='篡改'
        self.assertEqual(self.commit(modified).status_code,400)
        other=self.call('size-charts/detail?styleCode=OTHER').json;other['material']='cotton';self.save(other)
        self.assertEqual(self.commit(preview).status_code,409)
        self.assertEqual(self.detail()['version'],'0')
        self.assertEqual(len(self.detail()['dimensions']),9)
        self.assertEqual(len(self.call('size-charts/detail?styleCode=OTHER').json['dimensions']),9)


if __name__=='__main__':unittest.main()
