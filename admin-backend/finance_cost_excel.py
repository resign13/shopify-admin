"""Bounded, signed XLSX input; all writes still go through finance.save_costs.

Uses the application's existing openpyxl Excel stack. No uploads are persisted.
"""
from collections import Counter
from datetime import datetime
import hmac
from io import BytesIO
import json
from zipfile import ZipFile

from flask import current_app, request, send_file
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

import db
import finance
import safe_excel

VERSION = 'gingtto-style-cost-v1'
SHEET = '商品成本'
META = '_成本模板'
HEADERS = ['款号', '最近保存时间（北京时间）', '当前成本 CNY／件', '新成本 CNY／件']
MAX_STYLES = 2000
# Match the existing inventory Excel import protection (10 MB / 50 MB expanded).
MAX_BYTES = 10 * 1024 * 1024
MAX_EXPANDED_BYTES = 50 * 1024 * 1024
MIME = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def signature(code, revision):
    value = json.dumps([VERSION, code, revision], ensure_ascii=False, separators=(',', ':'))
    return current_app.extensions['security'].digest('finance-cost-excel:' + value)


def export(payload):
    if not isinstance(payload, dict):
        raise ValueError('成本筛选格式无效')
    source, params = finance.cost_query(payload)
    rows = db._fetch_all(source + ' ORDER BY p.style_code LIMIT %s', (*params, MAX_STYLES + 1))
    if len(rows) > MAX_STYLES:
        raise ValueError('成本表单次最多导出 2,000 款，请缩小查询范围后重试；未截断数据')
    book = Workbook()
    sheet = book.active
    sheet.title = SHEET
    sheet.append(HEADERS)
    meta = book.create_sheet(META)
    meta.append([VERSION])
    meta.append(['款号', '成本版本', '校验'])
    for item in rows:
        code, revision = item['style_code'], str(item['revision'])
        saved_at = item['updated_at'].astimezone(finance.TZ).replace(tzinfo=None) if item['updated_at'] else None
        sheet.append([code, saved_at, item['cost_cny'] if item['cost_cny'] is not None else '未设置', None])
        row = sheet.max_row
        sheet.cell(row, 1).number_format = '@'
        sheet.cell(row, 1).data_type = 's'
        sheet.cell(row, 2).number_format = 'yyyy-mm-dd hh:mm:ss'
        for col in (3, 4):
            sheet.cell(row, col).number_format = '#,##0.00'
            sheet.cell(row, col).alignment = Alignment(horizontal='right', vertical='center')
        sheet.cell(row, 4).fill = PatternFill('solid', fgColor='EFF6FF')
        sheet.cell(row, 4).font = Font(name='Microsoft YaHei', size=11, color='1D4ED8')
        sheet.row_dimensions[row].height = 27
        meta.append([code, revision, signature(code, revision)])
    # A literal style beginning with '=' remains text even in the hidden sheet.
    notes = book.create_sheet('填写说明')
    for text in [
        '商品成本填写说明',
        '只填写“商品成本”表的“新成本 CNY／件”列；每个款号一个成本，全部颜色与尺码共用。',
        '留空不修改；0 是明确零成本；与当前成本相同则跳过。不通过 Excel 清空成本。',
        '填写非负金额，最多两位小数，不使用公式。款号保留文本，不新增或重复款号。',
        '当前成本与最近保存时间仅供参考，修改它们不会改变系统数据；时间均为北京时间。',
        '成本表不含页面未保存草稿。导入前请先保存或撤销页面修改，上传后预览并确认才保存。',
        '首次成本补算此前订单；后续成本仅用于保存后的新订单，服务器统一确定保存时间。',
        '请保留系统模板及隐藏版本信息。成本已被其他人修改时，重新导出核对后再导入。',
        '每次最多 2,000 款、10 MB。任一错误或版本冲突均整批不保存。',
    ]:
        notes.append([text])
    notes.column_dimensions['A'].width = 104
    for row in notes:
        row[0].font = Font(name='Microsoft YaHei', size=11, color='334155', bold=row[0].row == 1)
        row[0].alignment = Alignment(vertical='center', wrap_text=True)
        notes.row_dimensions[row[0].row].height = 38 if row[0].row > 1 else 32
    for col, width in zip('ABCD', (28, 31, 25, 25)):
        sheet.column_dimensions[col].width = width
    for cell in sheet[1]:
        cell.font = Font(name='Microsoft YaHei', size=11, bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='243850')
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    sheet.row_dimensions[1].height = 34
    sheet.freeze_panes = 'B2'
    sheet.auto_filter.ref = f'A1:D{sheet.max_row}'
    sheet.sheet_view.showGridLines = notes.sheet_view.showGridLines = False
    validation = DataValidation(type='decimal', operator='between', formula1='0', formula2='9999999999.99', allow_blank=True)
    validation.errorTitle = '请输入有效成本'
    validation.error = '成本须为非负金额，最多两位小数；请勿使用公式。'
    validation.showErrorMessage = True
    sheet.add_data_validation(validation)
    validation.add(f'D2:D{max(2, sheet.max_row)}')
    meta.sheet_state = 'veryHidden'
    out = BytesIO()
    safe_excel.prepare(book).save(out)
    book.close()
    out.seek(0)
    stamp = datetime.now(finance.TZ).strftime('%Y%m%d-%H%M%S')
    return send_file(out, as_attachment=True, download_name=f'product-costs-{stamp}.xlsx', mimetype=MIME)


def _text(cell):
    if cell.data_type == 'f' or not isinstance(cell.value, str):
        raise ValueError('款号须保留模板中的文本，不使用公式')
    value = cell.value.strip()
    if not value or len(value) > 200 or '\x00' in value:
        raise ValueError('款号无效')
    return value


def _metadata(book):
    if SHEET not in book.sheetnames or META not in book.sheetnames:
        raise ValueError('仅接受系统导出的商品成本 .xlsx 模板，请重新导出')
    sheet = book[META]
    if sheet['A1'].data_type == 'f' or sheet['A1'].value != VERSION:
        raise ValueError('成本模板版本不受支持，请重新导出')
    if (sheet.max_row or 0) > MAX_STYLES + 2 or (sheet.max_column or 0) > 3:
        raise ValueError('成本模板版本信息超出限制，请重新导出')
    result = {}
    for index, cells in enumerate(sheet.iter_rows(min_row=3, max_col=3), 3):
        if index > MAX_STYLES + 2:
            raise ValueError('成本模板版本信息超出限制，请重新导出')
        if all(c.value is None for c in cells):
            continue
        try:
            code = _text(cells[0])
            if any(c.data_type == 'f' for c in cells):
                raise ValueError()
            revision = str(finance.integer(cells[1].value, 0))
            if code in result or not isinstance(cells[2].value, str) or not hmac.compare_digest(cells[2].value, signature(code, revision)):
                raise ValueError()
        except (ValueError, TypeError):
            raise ValueError('模板版本信息损坏或不属于当前系统，请重新导出')
        result[code] = revision
    return result


def parse(raw):
    if len(raw) > MAX_BYTES:
        raise ValueError('导入文件不能超过 10 MB')
    try:
        with ZipFile(BytesIO(raw)) as zipped:
            if sum(f.file_size for f in zipped.infolist()) > MAX_EXPANDED_BYTES or len(zipped.infolist()) > 1000:
                raise ValueError('Excel 解压内容过大，请使用系统模板')
            if any(n.endswith('vbaProject.bin') for n in zipped.namelist()):
                raise ValueError('仅接受无宏的 .xlsx 成本模板')
        book = load_workbook(BytesIO(raw), read_only=True, data_only=False, keep_links=False)
    except Exception as exc:
        if isinstance(exc, ValueError) and str(exc).startswith(('Excel ', '仅接受')):
            raise
        raise ValueError('请上传有效的系统 .xlsx 成本模板') from exc
    try:
        versions = _metadata(book)
        sheet = book[SHEET]
        if (sheet.max_row or 0) > MAX_STYLES + 1 or (sheet.max_column or 0) > 4:
            raise ValueError('商品成本表最多 2,000 款，请勿新增列或超出模板范围')
        if [c.value for c in next(sheet.iter_rows(min_row=1, max_row=1, max_col=4))] != HEADERS:
            raise ValueError('成本模板表头已改变，请重新导出')
        result = []
        # Also cap iteration independently of potentially forged XLSX dimensions.
        for index, cells in enumerate(sheet.iter_rows(min_row=2, max_col=4), 2):
            if index > MAX_STYLES + 1:
                raise ValueError('商品成本表最多 2,000 款')
            if all(c.value is None for c in cells):
                continue
            row = {'rowNumber': index, 'styleCode': '', 'costCny': None, 'status': 'blank', 'message': '留空跳过'}
            try:
                row['styleCode'] = _text(cells[0])
                if row['styleCode'] not in versions:
                    raise ValueError('款号不属于本次导出模板或版本信息缺失，请重新导出')
                row['version'] = versions[row['styleCode']]
                cell = cells[3]
                if cell.data_type == 'f':
                    raise ValueError('新成本不支持公式，请填写金额')
                if cell.value is not None and str(cell.value).strip():
                    row['costCny'] = finance.money(finance.number(cell.value))
                    row.update(status='update', message='待更新')
            except ValueError as exc:
                row.update(status='error', message=str(exc))
            result.append(row)
        duplicates = {code for code, count in Counter(r['styleCode'] for r in result if r['styleCode']).items() if count > 1}
        for row in result:
            if row['styleCode'] in duplicates:
                row.update(status='error', message='款号重复，请保留一行后重试')
        return result
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError('成本模板内容损坏，请重新导出 .xlsx 文件') from exc
    finally:
        book.close()


def preview_upload():
    uploaded = request.files.get('file')
    if not uploaded or not uploaded.filename or not uploaded.filename.lower().endswith('.xlsx'):
        raise ValueError('请选择系统导出的 .xlsx 成本表')
    rows = parse(uploaded.read(MAX_BYTES + 1))
    codes = list({r['styleCode'] for r in rows if r['styleCode']})
    source, _ = finance.cost_query({})
    saved = {r['style_code']: r for r in db._fetch_all(source + ' WHERE p.style_code=ANY(%s)', (codes,))} if codes else {}
    changes = []
    for row in rows:
        current = saved.get(row['styleCode'])
        row['oldCostCny'] = finance.money(current['cost_cny']) if current else None
        if row['status'] == 'error':
            continue
        if current is None:
            row.update(status='error', message='款号不存在或商品已移除，请重新导出')
        elif row['status'] == 'update':
            if row['version'] != str(current['revision']):
                row.update(status='error', message='版本冲突：导出后成本已更新，请重新导出并核对')
            elif row['costCny'] == row['oldCostCny']:
                row.update(status='same', message='相同跳过')
            else:
                changes.append({k: row[k] for k in ('styleCode', 'costCny', 'version')})
    summary = {key: sum(r['status'] == key for r in rows) for key in ('update', 'blank', 'same', 'error')}
    response = {'rows': rows, 'summary': summary, 'changes': [], 'previewHash': None, 'affectedOrders': 0}
    if changes and not summary['error']:
        response.update(finance.preview({'changes': changes}))
    return response
