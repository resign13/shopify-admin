"""Defect staging/reservation/return; fixtures write only to a guarded *_test DB."""
import io
import importlib.util
from pathlib import Path
import unittest
import uuid
from concurrent.futures import ThreadPoolExecutor
from unittest.mock import patch
from openpyxl import load_workbook
from psycopg.errors import CheckViolation
import test_workbench as fixtures
import test_orders as order_fixtures
import workbench
db, application = fixtures.db, fixtures.application


class DefectiveInventoryTest(unittest.TestCase):
    setUp = fixtures.WorkbenchTest.setUp
    call = fixtures.WorkbenchTest.call
    detail = fixtures.WorkbenchTest.detail
    workbook = fixtures.WorkbenchTest.workbook
    upload = fixtures.WorkbenchTest.upload
    product_payload = fixtures.WorkbenchTest.product_payload

    def sizes(self): return {s['sizeCode']: s for s in self.detail()['sizePrices']}
    def save(self, values, role='admin', **extra):
        return self.call('inventory/1', 'PUT', {'version': self.detail()['version'], 'sizeStocks': {},
                        'defectivePendingBySize': values, **extra}, role)
    def return_saved(self, body=None, role='admin', product_id=1):
        if body is None: body={'version': self.detail()['version'], 'requestId': str(uuid.uuid4())}
        return self.call(f'inventory/{product_id}/defective/return', 'POST', body, role)
    def history(self): return self.call('inventory/1/operations').json

    def test_registration_is_independent_persistent_absolute_and_clearable(self):
        before=self.sizes()
        self.assertEqual(self.save({'M': 3, 'L': 2}).status_code, 200)
        after=self.sizes()
        for code in before:
            for field in ['stock','contractPending','pendingInspection','pendingInbound','temporaryInbound']:
                self.assertEqual(after[code][field],before[code][field])
        self.assertEqual(after['M']['defectivePending'],3)
        self.assertEqual(self.save({'M': 2, 'L': 0}).status_code,200)
        self.assertEqual(self.sizes()['M']['defectivePending'],2)
        self.assertEqual(self.save({'M': 0}).status_code,200)
        self.assertEqual(self.sizes()['M']['defectivePending'],0)
        self.assertEqual(self.history()['total'],3)

    def test_input_validation_and_whole_save_rollback(self):
        before=self.detail()
        for value in [-1,7,2147483648,1.5,True,None,'','1e2']:
            response=self.save({'M':value}, sizeStocks={'L':99})
            self.assertEqual(response.status_code,400,(value,response.json))
            self.assertEqual(self.detail(),before)
        for values in [{'UNKNOWN':2}, [], '3']:
            self.assertEqual(self.save(values).status_code,400)
        self.assertEqual(self.history()['total'],0)

    def test_qualified_transfer_and_inspection_edit_reserve_defects(self):
        self.save({'M':3})
        before=self.detail()
        for extra in [{'pendingInboundBySize':{'M':8}}, {'pendingInspectionBySize':{'M':2}}]:
            response=self.save({},**extra)
            self.assertEqual(response.status_code,400,response.json)
            self.assertEqual(self.detail(),before)
        self.assertEqual(self.save({},pendingInboundBySize={'M':7}).status_code,200)
        self.assertEqual(self.sizes()['M']['pendingInspection'],3)
        self.assertEqual(self.sizes()['M']['contractPending'],12)
        # Releasing the reserve and moving the newly free quantity is atomic.
        self.assertEqual(self.save({'M':1},pendingInboundBySize={'M':9}).status_code,200)
        self.assertEqual(self.sizes()['M']['pendingInspection'],1)

    def test_multi_size_return_and_all_sources_remain_independent(self):
        for stock in [-5,0,10]:
            with self.subTest(stock=stock):
                self.save({'M':3,'L':2},sizeStocks={'M':stock},temporaryInboundBySize={'M':8})
                before=self.sizes()
                body={'version': self.detail()['version'], 'requestId':str(uuid.uuid4()),
                      'defectivePendingBySize':{'M':999},'quantity':999,'source':'temporary'}
                response=self.return_saved(body)
                self.assertEqual(response.status_code,200,response.json)
                self.assertEqual((response.json['returnedUnits'],response.json['returnedSizes'],response.json['source']),(5,2,'defective-return'))
                after=self.sizes()
                for code in before:
                    quantity=before[code]['defectivePending']
                    self.assertEqual(after[code]['defectivePending'],0)
                    self.assertEqual(after[code]['pendingInspection'],before[code]['pendingInspection']-quantity)
                    self.assertEqual(after[code]['contractPending'],before[code]['contractPending']+quantity)
                    for field in ['stock','pendingInbound','temporaryInbound']:
                        self.assertEqual(after[code][field],before[code][field])
                self.assertEqual(self.detail()['stock'],sum(s['stock'] for s in before.values()))
                log=self.history()['items'][0]
                self.assertTrue(all(c['operation']=='defective_return' for c in log['changes']))
                count=self.history()['total']
                self.assertTrue(self.return_saved(body).json['replayed'])
                self.assertEqual(self.history()['total'],count)
                self.assertEqual(self.return_saved().status_code,400)
                # Prepare inspection again for the next signed-stock scenario.
                self.save({},pendingInspectionBySize={'M':6,'L':3})

    def test_entire_inspection_can_be_returned_and_normal_receipts_preserve_defects(self):
        self.save({'M':6,'L':3},temporaryInboundBySize={'M':5})
        for suffix in ['/receive','/temporary/receive']:
            body={'version':self.detail()['version'],'requestId':str(uuid.uuid4())}
            self.assertEqual(self.call('inventory/1'+suffix,'POST',body).status_code,200)
            self.assertEqual(self.sizes()['M']['defectivePending'],6)
            self.assertEqual(self.sizes()['M']['pendingInspection'],6)
        self.assertEqual(self.return_saved().status_code,200)
        self.assertEqual(self.sizes()['M']['pendingInspection'],0)

    def test_duplicate_requests_and_source_actor_namespaces(self):
        self.save({'M':3})
        body={'version':self.detail()['version'],'requestId':str(uuid.uuid4())}
        with ThreadPoolExecutor(max_workers=2) as pool:
            responses=list(pool.map(lambda _:self.return_saved(body),range(2)))
        self.assertEqual([r.status_code for r in responses],[200,200],[r.json for r in responses])
        self.assertEqual(sum(bool(r.json.get('replayed')) for r in responses),1)
        self.assertEqual(self.history()['total'],2)
        self.assertEqual(self.return_saved(body,role='warehouse').status_code,400)
        self.assertEqual(self.call('inventory/1/receive','POST',body).status_code,400)
        self.assertEqual(self.return_saved({**body,'version':'other'}).status_code,400)

    def test_different_requests_same_version_conflict(self):
        self.save({'M':3})
        version=self.detail()['version']
        with ThreadPoolExecutor(max_workers=2) as pool:
            responses=list(pool.map(lambda _:self.return_saved({'version':version,'requestId':str(uuid.uuid4())}),range(2)))
        self.assertEqual(sorted(r.status_code for r in responses),[200,409])
        self.assertEqual(self.sizes()['M']['contractPending'],15)
        self.assertEqual(self.history()['total'],2)

    def test_concurrent_registration_and_return_one_version_one_effect(self):
        self.save({'M':3})
        version=self.detail()['version']
        with ThreadPoolExecutor(max_workers=2) as pool:
            a=pool.submit(self.return_saved,{'version':version,'requestId':str(uuid.uuid4())})
            b=pool.submit(self.call,'inventory/1','PUT',{'version':version,'sizeStocks':{},'defectivePendingBySize':{'M':5}})
            responses=[a.result(timeout=40),b.result(timeout=40)]
        self.assertEqual(sorted(r.status_code for r in responses),[200,409],[r.json for r in responses])
        after=self.sizes()['M']
        self.assertEqual(after['defectivePending'],0 if responses[0].status_code==200 else 5)
        self.assertEqual(after['stock'],10)
        self.assertEqual(self.history()['total'],2)

    def test_log_failure_rolls_back_stages_version_receipt_and_audit(self):
        self.save({'M':3,'L':2})
        before=self.detail();audit=self.call('audit-logs').json['total']
        body={'version':before['version'],'requestId':str(uuid.uuid4())}
        record=workbench.record_inventory_registration
        def fail(*args,**kwargs):
            record(*args,**kwargs)
            raise RuntimeError('Injected defect log failure')
        with patch.object(workbench,'record_inventory_registration',side_effect=fail):
            with self.assertRaisesRegex(RuntimeError,'defect log failure'):self.return_saved(body)
        self.assertEqual(self.detail(),before)
        self.assertEqual(self.history()['total'],1)
        self.assertEqual(self.call('audit-logs').json['total'],audit)
        self.assertIsNone(db._fetch_one('SELECT request_id FROM inventory_receipts WHERE request_id=%s',(body['requestId'],)))
        self.assertEqual(self.return_saved(body).status_code,200)

    def test_concurrent_product_edit_and_return_preserve_saved_defects(self):
        self.save({'M':3})
        product=self.product_payload(1);version=self.detail()['version']
        with ThreadPoolExecutor(max_workers=2) as pool:
            a=pool.submit(self.return_saved,{'version':version,'requestId':str(uuid.uuid4())})
            b=pool.submit(self.call,'products/save-group','POST',{'products':[product],'versions':{'1':version}})
            responses=[a.result(timeout=40),b.result(timeout=40)]
        self.assertEqual(sorted(r.status_code for r in responses),[200,409],[r.json for r in responses])
        self.assertEqual(self.sizes()['M']['defectivePending'],0 if responses[0].status_code==200 else 3)
        self.assertEqual(self.sizes()['M']['stock'],10)

    def test_concurrent_order_and_return_do_not_adjust_stock_twice(self):
        self.save({'M':3})
        body={'version':self.detail()['version'],'requestId':str(uuid.uuid4())}
        order=order_fixtures.AdminOrdersTest.payload(self)
        order['items']=[{'productId':1,'sizeCode':'M','quantity':2,'unitPrice':20}]
        with ThreadPoolExecutor(max_workers=2) as pool:
            a=pool.submit(self.return_saved,body)
            b=pool.submit(self.call,'orders','POST',order)
            returned,created=a.result(timeout=40),b.result(timeout=40)
        self.assertEqual(created.status_code,200,created.json)
        self.assertIn(returned.status_code,[200,409],returned.json)
        self.assertEqual(self.sizes()['M']['stock'],8)
        self.assertEqual(self.sizes()['M']['defectivePending'],0 if returned.status_code==200 else 3)

    def test_concurrent_normal_receipt_and_return_have_independent_sources(self):
        self.save({'M':3},temporaryInboundBySize={'M':7})
        version=self.detail()['version']
        with ThreadPoolExecutor(max_workers=2) as pool:
            a=pool.submit(self.return_saved,{'version':version,'requestId':str(uuid.uuid4())})
            b=pool.submit(self.call,'inventory/1/receive','POST',{'version':version,'requestId':str(uuid.uuid4())})
            returned,received=a.result(timeout=40),b.result(timeout=40)
        self.assertEqual(sorted([returned.status_code,received.status_code]),[200,409])
        self.assertEqual(self.sizes()['M']['stock'],14 if received.status_code==200 else 10)
        self.assertEqual(self.sizes()['M']['defectivePending'],0 if returned.status_code==200 else 3)
        self.assertEqual(self.sizes()['M']['temporaryInbound'],7)

    def test_overflow_invalid_size_and_missing_request_do_not_partially_return(self):
        self.save({'M':3,'L':2},contractPendingBySize={'L':2147483647})
        before=self.detail()
        self.assertEqual(self.return_saved().status_code,400)
        self.assertEqual(self.detail(),before)
        self.save({},contractPendingBySize={'L':6})
        db._fetch_one("DELETE FROM product_sizes WHERE product_id=1 AND size_code='L' RETURNING id")
        self.assertEqual(self.return_saved().status_code,400)
        self.assertEqual(self.sizes()['M']['defectivePending'],3)
        for body in [{},{'version':self.detail()['version'],'requestId':'invalid'}]:
            self.assertEqual(self.return_saved(body).status_code,400)

    def test_product_edit_preserves_and_protects_defects_and_inactive_return(self):
        self.save({'M':3})
        p=self.product_payload(1);p['sizePrices'][0]['defectivePending']=999
        self.assertEqual(self.call('products/save-group','POST',{'products':[p],'versions':{'1':p['version']}}).status_code,200)
        self.assertEqual(self.sizes()['M']['defectivePending'],3)
        p=self.product_payload(1);p['sizes']=p['sizes'][1:];p['sizePrices']=p['sizePrices'][1:]
        self.assertEqual(self.call('products/save-group','POST',{'products':[p],'versions':{'1':p['version']}}).status_code,400)
        self.assertEqual(self.call('products/1','DELETE').status_code,200)
        self.assertIn(1,[p['id'] for p in self.call('inventory?pageSize=50').json['items']])
        self.assertEqual(self.return_saved().status_code,200)

    def test_excel_v5_readonly_reserve_and_v1_v4_compatibility(self):
        self.save({'M':3})
        raw=self.workbook()
        book=load_workbook(io.BytesIO(raw));sheet=book.active
        self.assertEqual((sheet['A2'].value,sheet['M1'].value,sheet['N1'].value,sheet['S1'].value,sheet['S2'].value),('inventory-v5','合格','原始合格','次品（只读）',3))
        ignored=self.workbook(lambda s:setattr(s['S2'],'value',999))
        data=self.upload('preview',ignored).json
        self.assertEqual(data['errors'],[])
        self.assertEqual(data['summary']['changedRowCount'],0)
        self.assertEqual(self.upload('confirm',ignored,data['fileHash']).status_code,200)
        self.assertEqual(self.sizes()['M']['defectivePending'],3)
        for version,columns in [('inventory-v1',12),('inventory-v2',14),('inventory-v3',16),('inventory-v4',18),('inventory-v5',19)]:
            def legacy(s):
                if version!='inventory-v5':
                    s.delete_cols(columns+1,s.max_column-columns)
                    for col,label in enumerate(application.INVENTORY_HEADER_VERSIONS[version],1):s.cell(1,col,label)
                    for row in range(2,s.max_row+1):s.cell(row,1,version)
                s['H2']=int(s['H2'].value)+1
            raw=self.workbook(legacy);data=self.upload('preview',raw).json
            self.assertEqual(data['errors'],[],(version,data))
            self.assertEqual(self.upload('confirm',raw,data['fileHash']).status_code,200)
            self.assertEqual(self.sizes()['M']['defectivePending'],3)
        for cell,value in [('M2',8),('O2',2)]:
            raw=self.workbook(lambda s:setattr(s[cell],'value',value))
            self.assertTrue(self.upload('preview',raw).json['errors'])

    def test_excel_confirmation_rechecks_new_reserve_and_rolls_back_other_rows(self):
        raw=self.workbook(lambda s:(setattr(s['M2'],'value',8),setattr(s['H3'],'value',99)))
        data=self.upload('preview',raw).json;self.assertEqual(data['errors'],[])
        self.save({'M':3})
        before=self.detail()
        self.assertEqual(self.upload('confirm',raw,data['fileHash']).status_code,400)
        self.assertEqual(self.detail(),before)

    def test_summaries_customer_boundaries_and_permissions(self):
        self.save({'M':3})
        self.assertEqual(self.call('inventory?page=1&pageSize=25').json['summary']['defectivePending'],3)
        self.assertEqual(self.call('inventory?page=1&pageSize=25').json['summary']['availableStock'],90)
        self.assertEqual(self.call('dashboard?view=workbench').json['snapshot']['defectivePending'],3)
        self.assertEqual(self.call('dashboard?view=style-performance').json['summary']['defectivePending'],3)
        self.assertEqual(self.call('dashboard?view=style-detail&styleCode=GT-2026').json['items'][0]['defectivePending'],3)
        for path in ['inventory','inventory/1','products','products/1']:
            response=self.call(path,role='customer')
            if response.status_code==200:self.assertNotIn('defectivePending',response.get_data(as_text=True))
        for role in ['admin','sales','warehouse']:
            self.save({'M':1},role=role)
            self.assertEqual(self.return_saved(role=role).status_code,200)
        self.assertEqual(self.save({'M':1},role='customer').status_code,403)
        self.assertEqual(self.return_saved(role='customer').status_code,403)
        db._fetch_one("UPDATE admin_users SET permissions='[]' WHERE role='warehouse' RETURNING id")
        self.assertEqual(self.return_saved(role='warehouse').status_code,403)

    def test_both_migrations_are_repeat_safe_and_storefront_hides_defects(self):
        self.save({'M':3},sizeStocks={'M':-5},temporaryInboundBySize={'M':7})
        store_path=Path(__file__).resolve().parents[2]/'shopify/storefront-backend/db.py'
        migrations=[db._apply_schema_migrations]
        if store_path.exists():
            spec=importlib.util.spec_from_file_location('defective_store_db',store_path)
            store=importlib.util.module_from_spec(spec);spec.loader.exec_module(store)
            migrations.append(store._apply_schema_migrations)
            self.assertEqual(store_path.with_name('inventory_policy.py').read_bytes(),Path(db.inventory_policy.__file__).read_bytes())
            public=next(p for p in store.list_products() if p['id']==1)
            self.assertNotIn('defectivePending',repr(public))
            self.assertEqual(public['stock'],20)
        before=self.sizes()
        with db._connect() as conn:
            for migrate in migrations*2:migrate(conn.cursor())
        self.assertEqual(self.sizes(),before)
        for sql in ['SET defective_pending=-1','SET defective_pending=99','SET pending_inspection=0']:
            with self.assertRaises(CheckViolation):
                with db._connect() as conn:conn.execute('UPDATE product_size_prices '+sql+' WHERE product_id=1')
        # Simulate an old schema in a rolled-back transaction, never the mirror.
        with db._connect() as conn:
            conn.execute('ALTER TABLE product_size_prices DROP COLUMN defective_pending')
            for migrate in migrations*2:migrate(conn.cursor())
            self.assertEqual(conn.execute('SELECT SUM(defective_pending) AS n FROM product_size_prices').fetchone()['n'],0)
            conn.rollback()


if __name__=='__main__':unittest.main()
