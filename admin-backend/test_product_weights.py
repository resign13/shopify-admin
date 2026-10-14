"""Current net-weight tests; all fixtures are guarded by test_workbench's *_test check."""
import copy
import importlib.util
from io import BytesIO
from pathlib import Path
import unittest
import uuid
from unittest.mock import patch
from openpyxl import load_workbook
from psycopg.errors import CheckViolation
import test_workbench as fixtures
import test_orders as orders
import product_weights

db, application = fixtures.db, fixtures.application


class ProductWeightTest(unittest.TestCase):
    setUp = fixtures.WorkbenchTest.setUp
    call = fixtures.WorkbenchTest.call
    product_payload = fixtures.WorkbenchTest.product_payload

    def sizes(self, pid=1):
        return {s['sizeCode']: s for s in self.call(f'products/{pid}').json['product']['sizePrices']}

    def save(self, weights, pid=1, *, legacy=False):
        row = self.product_payload(pid)
        for size in row['sizePrices']:
            size.pop('weightGrams', None)
            if size['sizeCode'] in weights:
                size['weightGrams'] = weights[size['sizeCode']]
        if legacy:
            return self.call(f'products/{pid}', 'PUT', row)
        return self.call('products/save-group', 'POST', {'products': [row], 'versions': {str(pid): row['version']}})

    def detail(self, oid=1, role='admin'):
        r = self.call(f'orders/{oid}', role=role)
        self.assertEqual(r.status_code, 200, r.json)
        return r.json['order']

    def balances(self):
        return db._fetch_all('SELECT product_id,size_code,stock,contract_pending,pending_inspection,pending_inbound,temporary_inbound,defective_pending FROM product_size_prices ORDER BY product_id,size_code')

    def test_optional_save_clear_and_omission_preserve(self):
        self.assertTrue(all(s['weightGrams'] is None for s in self.sizes().values()))
        before = self.balances()
        self.assertEqual(self.save({'M':350,'L':380}).status_code, 200)
        self.assertEqual(self.save({}, legacy=True).status_code, 200)
        self.assertEqual(self.sizes()['M']['weightGrams'], 350)
        self.assertEqual(self.sizes()['L']['weightGrams'], 380)
        self.assertEqual(self.save({'L':None}).status_code, 200)
        self.assertEqual(self.sizes()['M']['weightGrams'], 350)
        self.assertIsNone(self.sizes()['L']['weightGrams'])
        self.assertEqual(self.balances(), before)
        self.assertIsNotNone(db._fetch_one("SELECT id FROM admin_audit_logs WHERE entity_table='product_size_prices' AND after_data->>'weight_grams'='350' LIMIT 1"))

    def test_invalid_and_group_atomicity_stale_versions(self):
        for bad in (0,-1,1.5,True,'', '1.2', 'NaN', 'Infinity', {}, [], 2147483648):
            with self.subTest(value=bad):
                self.assertEqual(self.save({'M':bad}).status_code, 400)
                self.assertIsNone(self.sizes()['M']['weightGrams'])
        a,b = self.product_payload(1), self.product_payload(2)
        a['sizePrices'][0]['weightGrams']=300
        b['sizePrices'][0]['weightGrams']=0
        body={'products':[a,b], 'versions':{str(p['id']):p['version'] for p in (a,b)}}
        self.assertEqual(self.call('products/save-group','POST',body).status_code,400)
        self.assertIsNone(self.sizes()['M']['weightGrams'])
        b['sizePrices'][0]['weightGrams']=350
        self.assertEqual(self.call('products/save-group','POST',body).status_code,200)
        self.assertEqual(self.call('products/save-group','POST',body).status_code,409)

    def test_new_variants_and_compatibility_batch(self):
        row=self.product_payload(1)
        row.pop('id'); row.pop('version')
        row.update(productCode='WEIGHT-BLUE',sku='WEIGHT-BLUE',slug='weight-blue')
        row['sizePrices'][0]['weightGrams']=350
        response=self.call('products','POST',row)
        self.assertEqual(response.status_code,201,response.json)
        pid=response.json['product']['id']
        self.assertEqual(self.sizes(pid)['M']['weightGrams'],350)
        self.assertIsNone(self.sizes(pid)['L']['weightGrams'])
        row.pop('sizePrices')
        row.update(familyCode='BATCH-WEIGHT',sizes=['30/M','Tall XL'])
        row['variants']=[{'sku':'BATCH-WEIGHT-BLACK','productCode':'BATCH-WEIGHT-BLACK','slug':'batch-weight-black',
            'colorName':'Black','colorHex':'#111111','imageUrls':['/uploads/fixture.jpg']*3,
            'sizePrices':[{'sizeCode':'30/M','price':10,'stock':0,'weightGrams':410}, {'sizeCode':'Tall XL','price':10,'stock':0}]}]
        response=self.call('products','POST',row)
        self.assertEqual(response.status_code,201,response.json)
        pid=response.json['product'][0]['id']
        self.assertEqual(self.sizes(pid)['30/M']['weightGrams'],410)
        group=self.product_payload(1)
        group.update(id=None,productCode='WEIGHT-GROUP',sku='WEIGHT-GROUP',slug='weight-group',familyCode='WEIGHT-GROUP',colorGroup='WEIGHT-GROUP')
        group['sizePrices'][0]['weightGrams']=460
        response=self.call('products/save-group','POST',{'products':[group],'versions':{}})
        self.assertEqual(response.status_code,200,response.json)
        pid=db._fetch_one("SELECT id FROM products WHERE sku='WEIGHT-GROUP'")['id']
        self.assertEqual(self.sizes(pid)['M']['weightGrams'],460)

    def test_current_historical_orders_and_no_business_writes(self):
        before=self.balances()
        version=self.detail()['version']
        self.save({'M':350}); self.save({'M':380},pid=2)
        first=self.detail()
        self.assertEqual(first['totalWeightGrams'],'1080')
        self.assertEqual(first['items'][0]['lineWeightGrams'],'700')
        self.assertEqual(first['missingWeightQuantity'],0)
        self.save({'M':500})
        second=self.detail()
        self.assertEqual(second['totalWeightGrams'],'1380')
        self.assertEqual(second['version'],version)
        self.assertEqual(second['totalAmount'],first['totalAmount'])
        self.assertEqual(self.balances(),before)
        self.assertEqual(self.call('orders/1/operations').json['total'],0)
        self.assertEqual(db._fetch_one('SELECT count(*) AS n FROM order_created_events')['n'],0)
        self.assertFalse(any('weight' in key.lower() for key in self.call('orders?view=workbench').json['items'][0]))
        self.assertNotIn('unitWeightGrams',db.list_orders()[0]['items'][0])

    def test_inventory_changes_keep_weights_and_do_not_expose_them(self):
        self.save({'M':350, 'L':380})
        detail=self.call('inventory/1').json['product']
        self.assertNotIn('weightGrams',detail['sizePrices'][0])
        response=self.call('inventory/1','PUT',{'version':detail['version'], 'sizeStocks':{}, 'temporaryInboundBySize':{'M':5}})
        self.assertEqual(response.status_code,200,response.json)
        self.assertEqual(self.sizes()['M']['weightGrams'],350)
        self.assertEqual(self.sizes()['L']['weightGrams'],380)
        response=self.call('inventory/1/temporary/receive','POST',{'version':self.call('inventory/1').json['product']['version'], 'requestId':str(uuid.uuid4())})
        self.assertEqual(response.status_code,200,response.json)
        self.assertEqual(self.sizes()['M']['weightGrams'],350)
        self.assertEqual(self.detail()['knownWeightGrams'],'700')

    def test_draft_preview_batches_current_weights_without_business_writes(self):
        self.save({'M':350, 'L':380}); self.save({'M':400}, pid=2)
        before = self.balances()
        version = self.detail()['version']
        audit_count = db._fetch_one('SELECT COUNT(*) AS n FROM admin_audit_logs')['n']
        body = {'items': [{'productId':1,'sizeCode':'M','quantity':2,'unitWeightGrams':999},
                          {'productId':1,'sizeCode':'L','quantity':1},
                          {'productId':2,'sizeCode':'M','quantity':1},
                          {'productId':2,'sizeCode':'Tall XL','quantity':0}]}
        with patch.object(db, '_fetch_all', wraps=db._fetch_all) as queries:
            response = self.call('orders/weights/preview', 'POST', body)
            self.assertEqual(queries.call_count, 1)
        self.assertEqual(response.status_code, 200, response.json)
        self.assertEqual(response.json['totalWeightGrams'], '1480')
        self.assertEqual(response.json['missingItems'], [])
        self.assertEqual(self.balances(), before)
        self.assertEqual(self.detail()['version'], version)
        self.assertEqual(db._fetch_one('SELECT COUNT(*) AS n FROM admin_audit_logs')['n'], audit_count)
        self.assertEqual(db._fetch_one('SELECT count(*) AS n FROM order_created_events')['n'], 0)
        self.save({'L':None})
        result = self.call('orders/weights/preview', 'POST', body).json
        self.assertIsNone(result['totalWeightGrams'])
        self.assertEqual(result['knownWeightGrams'], '1100')
        self.assertEqual(result['missingWeightQuantity'], 1)
        self.assertEqual(result['missingItems'], [{'productId':1,'sku':'GT-2026-1','sizeCode':'L','quantity':1}])

    def test_draft_preview_exact_sizes_missing_removed_and_empty(self):
        self.save({'M':350})
        body = {'items':[{'productId':1,'sizeCode':size,'quantity':2} for size in ('M','30/M','')]}
        result = self.call('orders/weights/preview', 'POST', body).json
        self.assertEqual(result['knownWeightGrams'], '700')
        self.assertEqual(result['missingWeightQuantity'], 4)
        db._fetch_one("DELETE FROM product_size_prices WHERE product_id=1 AND size_code='M' RETURNING id")
        result = self.call('orders/weights/preview', 'POST', body).json
        self.assertEqual(result['knownWeightGrams'], '0')
        self.assertEqual(result['missingWeightQuantity'], 6)
        result = self.call('orders/weights/preview', 'POST', {'items':[{'productId':99999,'sizeCode':'M','quantity':1}]}).json
        self.assertEqual(result['missingWeightQuantity'], 1)
        self.assertEqual(result['missingItems'][0]['productId'], 99999)
        result = self.call('orders/weights/preview', 'POST', {'items':[]}).json
        self.assertEqual(result['knownWeightGrams'], '0')
        self.assertEqual(result['missingWeightQuantity'], 0)

    def test_draft_preview_large_quantities_use_exact_grams(self):
        self.save({'M':2147483647})
        result = self.call('orders/weights/preview', 'POST', {'items':[{'productId':1,'sizeCode':'M','quantity':2147483647}]}).json
        self.assertEqual(result['totalWeightGrams'], str(2147483647 ** 2))

    def test_draft_preview_validation_and_orders_only_permission(self):
        row = {'productId':1,'sizeCode':'M','quantity':1}
        invalid = [None, [], {}, {'items':'invalid'}, {'items':[None]}, {'items':[row]*501}, {'items':[row,row]}]
        for field, bad_values in [('productId',[0,-1,True,'1',1.5,9223372036854775808]),
                                  ('quantity',[-1,True,'1',1.5,2147483648]),
                                  ('sizeCode',[None,[], 'x'*201])]:
            invalid.extend({'items':[{**row,field:bad}]} for bad in bad_values)
        for body in invalid:
            with self.subTest(body=body):
                self.assertEqual(self.call('orders/weights/preview','POST',body).status_code,400)
        for role in ('warehouse','customer'):
            self.assertEqual(self.call('orders/weights/preview','POST',{'items':[row]},role=role).status_code,403)
        user = db.get_admin_user_by_email('sales@gingtto.test')
        values = {'name':user['name'],'email':user['email'],'role':'sales','status':'active','permissions':['orders']}
        db.update_admin_user(user['id'], values)
        self.assertEqual(self.call('orders/weights/preview','POST',{'items':[row]},role='sales').status_code,401)
        self.tokens['sales']=db.create_admin_session(user['id'])
        self.assertEqual(self.call('orders/weights/preview','POST',{'items':[row]},role='sales').status_code,200)
        db.update_admin_user(user['id'], {**values,'permissions':[]})
        self.assertEqual(self.call('orders/weights/preview','POST',{'items':[row]},role='sales').status_code,401)
        self.tokens['sales']=db.create_admin_session(user['id'])
        self.assertEqual(self.call('orders/weights/preview','POST',{'items':[row]},role='sales').status_code,403)
        with self.app.test_client() as client:
            self.assertEqual(client.post('/api/admin/orders/weights/preview',json={'items':[row]}).status_code,401)

    def test_missing_deleted_and_exact_size_no_sku_fallback(self):
        self.save({'M':350})
        result=self.detail()
        self.assertIsNone(result['totalWeightGrams'])
        self.assertEqual(result['knownWeightGrams'],'700')
        self.assertEqual(result['missingWeightQuantity'],1)
        db._fetch_one("UPDATE order_items SET size_code='30/M' WHERE order_id=1 AND product_id=1 RETURNING id")
        self.assertIsNone(self.detail()['items'][0]['unitWeightGrams'])
        self.assertEqual(self.detail()['knownWeightGrams'],'0')
        db._fetch_one('UPDATE order_items SET size_code=NULL WHERE order_id=1 RETURNING id')
        self.assertEqual(self.detail()['missingWeightQuantity'],3)
        db._fetch_one("UPDATE order_items SET size_code='' WHERE order_id=1 RETURNING id")
        self.assertEqual(self.detail()['missingWeightQuantity'],3)
        db._fetch_one("UPDATE order_items SET size_code='M' WHERE order_id=1 RETURNING id")
        db._fetch_one("DELETE FROM product_size_prices WHERE product_id=1 AND size_code='M' RETURNING id")
        self.assertIsNone(self.detail()['items'][0]['unitWeightGrams'])

    def test_order_edit_quantity_and_missing_weight_still_accept_creation(self):
        case=orders.AdminOrdersTest(); case.tokens=self.tokens
        created=case.create()
        self.assertIsNone(created['totalWeightGrams'])
        self.save({'M':350,'L':380})
        created=self.detail(created['id'])
        self.assertEqual(created['totalWeightGrams'],'1840')
        body=copy.deepcopy(created); body['items'][0]['quantity']=3
        self.assertEqual(self.call(f"orders/{created['id']}/details",'PUT',body).status_code,200)
        self.assertEqual(self.detail(created['id'])['totalWeightGrams'],'2190')

    def test_permissions_warehouse_financial_hide_and_public_boundary(self):
        self.save({'M':350}); self.save({'M':380},pid=2)
        warehouse=self.detail(role='warehouse')
        self.assertEqual(warehouse['totalWeightGrams'],'1080')
        self.assertNotIn('totalAmount',warehouse); self.assertNotIn('unitPrice',warehouse['items'][0])
        self.assertEqual(self.call('orders/1',role='customer').status_code,403)
        self.assertEqual(self.call('orders/1',role='sales').status_code,404)
        db._fetch_one('UPDATE orders SET owner_admin_id=2 WHERE id=1 RETURNING id')
        self.assertEqual(self.detail(role='sales')['totalWeightGrams'],'1080')
        for p in db.list_products():
            self.assertTrue(all('weightGrams' not in s for s in p['sizePrices']))
        for role in ('admin','sales','warehouse','customer'):
            value=self.call('inventory/1',role=role).json['product']
            self.assertTrue(all('weightGrams' not in s for s in value['sizePrices']))
        user=db.get_admin_user_by_email('sales@gingtto.test')
        db.update_admin_user(user['id'], {'name':user['name'],'email':user['email'],'role':'sales','status':'active','permissions':[]})
        self.assertEqual(self.call('orders/1',role='sales').status_code,401)
        self.tokens['sales']=db.create_admin_session(user['id'])
        self.assertEqual(self.call('orders/1',role='sales').status_code,403)

    def test_repeated_both_migrations_and_constraint(self):
        self.save({'M':350})
        before=self.balances()
        db.ensure_database_ready(); db.ensure_database_ready()
        root=Path(__file__).resolve().parents[2].parent/'shopify/storefront-backend'
        # Both repositories carry the identical standalone migration.
        if not root.exists(): root=Path(__file__).resolve().parents[2]/'shopify/storefront-backend'
        source=root/'product_weights.py'
        if source.exists():
            spec=importlib.util.spec_from_file_location('store_weights',source)
            module=importlib.util.module_from_spec(spec); spec.loader.exec_module(module)
            with db.get_connection() as conn,conn.cursor() as cur: module.migrate(cur)
        self.assertEqual(self.sizes()['M']['weightGrams'],350)
        self.assertEqual(self.balances(),before)
        with self.assertRaises(CheckViolation):
            with db.get_connection() as conn:
                conn.execute('UPDATE product_size_prices SET weight_grams=0 WHERE product_id=1')

    def test_all_exports_and_pi_format_with_current_weights(self):
        self.save({'M':350}); self.save({'M':380},pid=2)
        before=self.balances()
        with patch.object(application,'prefetch_export_images'),patch.object(application,'fetch_image_bytes',return_value=None):
            for path in ('orders/export','orders/export-by-sheet','orders/1/invoice'):
                response=self.call(path)
                self.assertEqual(response.status_code,200,response.json)
                sheet=load_workbook(BytesIO(response.data)).active
                values=[str(c.value) for row in sheet for c in row]
                self.assertTrue(any('1.080' in v for v in values))
                if path.endswith('invoice'):
                    self.assertEqual(sheet['J23'].value,'=J21+J22')
                    self.assertIn('$',sheet['I13'].number_format)
                    self.assertEqual(copy.copy(sheet['I13'].font),copy.copy(sheet['J13'].font))
                    self.assertIn('Net Weight',sheet['A24'].value)
                else: self.assertTrue(any('净重' in v for v in values))
            self.save({'M':None},pid=2)
            self.assertIsNone(self.detail()['totalWeightGrams'])
            order=self.detail(); order['note']='Fixture note'; order['labelImageUrls']=['/uploads/fixture.jpg']
            for builder in (application.build_orders_export,application.build_orders_sheet_export,application.build_order_invoice_export):
                book=load_workbook(builder(order if builder==application.build_order_invoice_export else [order]))
                values=[str(c.value) for ws in book for row in ws for c in row]
                self.assertTrue(any('missing weight' in v or '未设置重量' in v for v in values))
                self.assertTrue(any('Fixture note' in v for v in values))
        self.assertEqual(self.balances(),before)

    def test_exports_use_one_current_weight_snapshot(self):
        self.save({'M':380},pid=2)
        for path in ('orders/export','orders/export-by-sheet','orders/1/invoice'):
            self.save({'M':350})
            # A product save during image preparation must not mix weights
            # inside the file already being generated from a single query.
            def change_during_images(_orders):
                # Simulate a concurrent writer, outside the read-only export snapshot.
                with db._connect() as conn:
                    conn.execute("UPDATE product_size_prices SET weight_grams=500 WHERE product_id=1 AND size_code='M'")
            with patch.object(application,'prefetch_export_images',side_effect=change_during_images),patch.object(application,'fetch_image_bytes',return_value=None):
                response=self.call(path)
            self.assertEqual(response.status_code,200,response.json)
            book=load_workbook(BytesIO(response.data))
            self.assertTrue(any('1.080' in str(c.value) for ws in book for row in ws for c in row))
            self.assertEqual(self.detail()['totalWeightGrams'],'1380')
        with patch.object(db,'_fetch_all',wraps=db._fetch_all) as queries:
            values=db.list_orders(include_weight=True)
        self.assertEqual(queries.call_count,1)
        self.assertEqual(values[0]['totalWeightGrams'],'1380')


class ExactWeightTest(unittest.TestCase):
    def test_parser_summary_precision_and_missing(self):
        for v in (0,-1,1.5,True,'',float('nan'),float('inf'),2147483648):
            with self.assertRaises(ValueError): product_weights.parse(v)
        self.assertIsNone(product_weights.parse(None))
        n=2147483647
        values=product_weights.summary([{'unitWeightGrams':n,'quantity':n}]*500)
        self.assertEqual(values['totalWeightGrams'],str(n*n*500))
        self.assertEqual(product_weights.kilograms('9007199254740993'),'9,007,199,254,740.993')
        partial=product_weights.summary([{'unitWeightGrams':350,'quantity':10},{'unitWeightGrams':None,'quantity':5}])
        self.assertEqual(partial['knownWeightGrams'],'3500'); self.assertEqual(partial['missingWeightQuantity'],5)
        self.assertIsNone(partial['totalWeightGrams'])
