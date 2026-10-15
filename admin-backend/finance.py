"""Private style cost timelines and read-only order-time accounting (Decimal)."""
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP, localcontext
from functools import wraps
from io import BytesIO
import hashlib
import json
import uuid
import unicodedata

from flask import g, jsonify, request, send_file
from psycopg.types.json import Jsonb
import db
import module_permissions
import style_catalog
import workbench

TZ = timezone(timedelta(hours=8))
STATES = ['pending_payment', 'allocated', 'paid', 'shipped', 'completed']
CENT = Decimal('.01')
COST_BASIS = 'style_at_order_time'


def migrate(cur):
    # Logical product identity survives hard deletion and size-row reconstruction.
    cur.execute('''CREATE TABLE IF NOT EXISTS finance_cost_versions (
      product_id BIGINT NOT NULL, size_code TEXT NOT NULL, revision BIGINT NOT NULL DEFAULT 0,
      PRIMARY KEY(product_id,size_code));
    CREATE TABLE IF NOT EXISTS finance_cost_history (
      product_id BIGINT NOT NULL, size_code TEXT NOT NULL, effective_date DATE NOT NULL,
      cost_cny NUMERIC(12,2) CHECK(cost_cny>=0), sku TEXT NOT NULL, style_code TEXT NOT NULL,
      actor JSONB NOT NULL, updated_at TIMESTAMPTZ NOT NULL DEFAULT clock_timestamp(),
      PRIMARY KEY(product_id,size_code,effective_date));
    DO $$ BEGIN
      IF EXISTS(SELECT 1 FROM pg_attribute WHERE attrelid='finance_cost_history'::regclass
                AND attname='cost_cny' AND attnotnull AND NOT attisdropped) THEN
        ALTER TABLE finance_cost_history ALTER COLUMN cost_cny DROP NOT NULL;
      END IF;
    END $$;
    CREATE TABLE IF NOT EXISTS finance_cost_audit (
      id BIGINT GENERATED ALWAYS AS IDENTITY PRIMARY KEY, product_id BIGINT NOT NULL,
      size_code TEXT NOT NULL, effective_date DATE NOT NULL, before_data JSONB NOT NULL,
      after_data JSONB NOT NULL, actor JSONB NOT NULL, request_id UUID NOT NULL,
      created_at TIMESTAMPTZ NOT NULL DEFAULT clock_timestamp());
    CREATE TABLE IF NOT EXISTS finance_cost_receipts (
      request_id UUID PRIMARY KEY, actor_id BIGINT NOT NULL, fingerprint TEXT NOT NULL,
      result JSONB NOT NULL, created_at TIMESTAMPTZ NOT NULL DEFAULT clock_timestamp());
    CREATE INDEX IF NOT EXISTS finance_cost_audit_variant ON finance_cost_audit(product_id,size_code,id DESC);
    CREATE TABLE IF NOT EXISTS finance_settings (
      id SMALLINT PRIMARY KEY CHECK(id=1),
      exchange_rate NUMERIC(16,6) NOT NULL CHECK(exchange_rate>0),
      revision BIGINT NOT NULL DEFAULT 1,
      updated_by BIGINT REFERENCES admin_users(id) ON DELETE SET NULL,
      updated_at TIMESTAMPTZ NOT NULL DEFAULT clock_timestamp());
    -- Keep the legacy variant timelines intact; they are not the active cost source.
    -- No product/size FK: changing or deleting catalog rows must not erase cost history.
    CREATE TABLE IF NOT EXISTS finance_style_costs (
      style_code TEXT PRIMARY KEY CHECK(length(btrim(style_code)) BETWEEN 1 AND 200),
      cost_cny NUMERIC(12,2) CHECK(cost_cny>=0), revision BIGINT NOT NULL DEFAULT 0,
      actor JSONB NOT NULL DEFAULT '{}'::jsonb,
      updated_at TIMESTAMPTZ NOT NULL DEFAULT clock_timestamp());
    CREATE TABLE IF NOT EXISTS finance_style_cost_audit (
      id BIGINT GENERATED ALWAYS AS IDENTITY PRIMARY KEY, style_code TEXT NOT NULL,
      old_cost_cny NUMERIC(12,2), cost_cny NUMERIC(12,2), actor JSONB NOT NULL,
      request_id UUID NOT NULL, created_at TIMESTAMPTZ NOT NULL DEFAULT clock_timestamp(),
      effective_at TIMESTAMPTZ NOT NULL DEFAULT clock_timestamp(),
      is_initial BOOLEAN NOT NULL DEFAULT false);
    -- Saved events are also the immutable cost timeline. Upgrade using the actual
    -- original save time, never the migration time or the first historical order.
    ALTER TABLE finance_style_cost_audit ADD COLUMN IF NOT EXISTS effective_at TIMESTAMPTZ;
    UPDATE finance_style_cost_audit SET effective_at=created_at WHERE effective_at IS NULL;
    ALTER TABLE finance_style_cost_audit ALTER COLUMN effective_at SET DEFAULT clock_timestamp();
    ALTER TABLE finance_style_cost_audit ALTER COLUMN effective_at SET NOT NULL;
    ALTER TABLE finance_style_cost_audit ADD COLUMN IF NOT EXISTS is_initial BOOLEAN NOT NULL DEFAULT false;
    -- Repair a cache-only legacy record at its known timestamp, without inventing
    -- costs for earlier orders. Normal saves always write both rows atomically.
    INSERT INTO finance_style_cost_audit(style_code,old_cost_cny,cost_cny,actor,request_id,created_at,effective_at)
      SELECT c.style_code,NULL,c.cost_cny,c.actor,md5('finance-style-baseline:'||c.style_code)::uuid,c.updated_at,c.updated_at
      FROM finance_style_costs c WHERE NOT EXISTS (
        SELECT 1 FROM finance_style_cost_audit a WHERE a.style_code=c.style_code);
    -- The first actual amount (including explicit zero) is the historical
    -- baseline. Clears before any amount do not count as a first cost.
    WITH initial AS (
      SELECT DISTINCT ON (a.style_code) a.id FROM finance_style_cost_audit a
      WHERE a.cost_cny IS NOT NULL AND NOT EXISTS (
        SELECT 1 FROM finance_style_cost_audit b WHERE b.style_code=a.style_code AND b.is_initial)
      ORDER BY a.style_code,a.effective_at,a.id)
    UPDATE finance_style_cost_audit a SET is_initial=true FROM initial i WHERE a.id=i.id;
    CREATE UNIQUE INDEX IF NOT EXISTS finance_style_cost_audit_initial
      ON finance_style_cost_audit(style_code) WHERE is_initial;
    CREATE INDEX IF NOT EXISTS finance_style_cost_audit_style
      ON finance_style_cost_audit(style_code,id DESC);
    CREATE INDEX IF NOT EXISTS finance_style_cost_audit_timeline
      ON finance_style_cost_audit(style_code,effective_at DESC,id DESC);
    ''')


def number(value, *, rate=False):
    if isinstance(value, bool) or value is None or str(value).strip() == '':
        raise ValueError('请输入汇率' if rate else '请输入成本；未设置不能当作零成本保存')
    try:
        d = Decimal(str(value))
        quantum = Decimal('.000001') if rate else CENT
        if not d.is_finite() or d < 0 or d > Decimal('9999999999.99') or d != d.quantize(quantum) or (rate and d == 0):
            raise ValueError()
    except (InvalidOperation, ValueError):
        raise ValueError('汇率必须为正数，最多六位小数' if rate else '成本须为非负金额，最多两位小数')
    return d


def day(value=None):
    try:
        result = date.fromisoformat(str(value)) if value else datetime.now(TZ).date()
        if not 1900 <= result.year <= 9998: raise ValueError()
        return result
    except (ValueError, TypeError):
        raise ValueError('日期格式无效')


def integer(value, minimum=1):
    if isinstance(value, bool) or not str(value).isdigit() or not minimum <= int(value) <= 2147483647:
        raise ValueError('编号或分页参数无效')
    return int(value)


def terms(value):
    if isinstance(value, str): value = value.replace('，', ',').split(',')
    if not isinstance(value, list) or len(value) > 100 or any(not isinstance(s, str) or len(s) > 200 for s in value):
        raise ValueError('款号筛选无效，最多 100 项')
    return list(dict.fromkeys(s.strip() for s in value if s.strip()))


def money(value):
    return None if value is None else format(value.quantize(CENT, rounding=ROUND_HALF_UP), '.2f')


def digest(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=False, default=str).encode()).hexdigest()


def preferences():
    row=db._fetch_one('SELECT exchange_rate,revision FROM finance_settings WHERE id=1')
    return {'exchangeRate':format(row['exchange_rate'],'.6f') if row else None,'version':str(row['revision']) if row else '0'}


def save_preferences(payload):
    if not isinstance(payload,dict):raise ValueError('汇率保存格式无效')
    rate=number(payload.get('exchangeRate'),rate=True)
    expected=str(integer(payload.get('version'),0))
    uid=g.current_user['id']
    db._fetch_one('SELECT pg_advisory_xact_lock(hashtextextended(%s,0))',('finance-shared-settings',))
    old=db._fetch_one('SELECT exchange_rate,revision FROM finance_settings WHERE id=1 FOR UPDATE')
    if old and old['exchange_rate']==rate:return preferences()
    if expected!=str(old['revision'] if old else 0):raise workbench.Conflict('公共汇率已更新，请核对最新保存值后重试')
    db._fetch_one('''INSERT INTO finance_settings(id,exchange_rate,updated_by) VALUES(1,%s,%s)
      ON CONFLICT(id) DO UPDATE SET exchange_rate=EXCLUDED.exchange_rate,
        revision=finance_settings.revision+1,updated_by=EXCLUDED.updated_by,updated_at=clock_timestamp() RETURNING revision''',(rate,uid))
    return preferences()


def options():
    return {
        'customers': db._fetch_all('SELECT id,name,company_name AS company FROM store_users ORDER BY name,id'),
        'salespeople': db._fetch_all("SELECT id,name,status FROM admin_users WHERE role='sales' ORDER BY name,id"),
        'styles': [r['code'] for r in db._fetch_all(f'SELECT DISTINCT {style_catalog.style_sql()} AS code FROM products p ORDER BY code')],
        'categories': db._fetch_all('SELECT category_key AS key FROM product_categories ORDER BY sort_order,id'),
    }


def costs(args):
    page=integer(args.get('page',1)); size=integer(args.get('pageSize',20))
    if size not in (12,20,24,48): raise ValueError('每页请选择 12、20、24 或 48 款')
    codes=terms(args.get('styleCodes',''))
    keyword=str(args.get('keyword','')).strip()
    if len(keyword)>200:raise ValueError('查询关键字最多 200 字符')
    where=[];params=[]
    if codes:where.append('lower(p.style_code)=ANY(%s)');params.append([c.lower() for c in codes])
    if keyword:
        where.append('EXISTS(SELECT 1 FROM catalog k WHERE k.style_code=p.style_code AND (k.sku ILIKE %s OR k.color_name ILIKE %s OR k.style_code ILIKE %s))')
        params.extend(['%'+keyword+'%']*3)
    if args.get('category'):
        where.append('EXISTS(SELECT 1 FROM catalog k WHERE k.style_code=p.style_code AND k.category_key=%s)');params.append(args['category'])
    if str(args.get('missingOnly','')).lower() in ('1','true'):where.append('c.cost_cny IS NULL')
    # Filter and page whole styles in SQL. No title, image, color/size matrix or N+1.
    source=f"""WITH catalog AS (SELECT {style_catalog.style_sql()} AS style_code,p.sku,p.color_name,c.category_key
        FROM products p JOIN product_categories c ON c.id=p.category_id),
      styles AS (SELECT DISTINCT style_code FROM catalog)
      SELECT p.style_code,c.cost_cny,COALESCE(c.revision,0) AS revision,c.updated_at
      FROM styles p LEFT JOIN finance_style_costs c ON c.style_code=p.style_code"""
    if where:source+=' WHERE '+' AND '.join(where)
    total=db._fetch_one('SELECT count(*) AS n FROM ('+source+') q',params)['n']
    rows=db._fetch_all(source+' ORDER BY p.style_code LIMIT %s OFFSET %s',(*params,size,(page-1)*size))
    return {'items':[{'styleCode':r['style_code'],'costCny':money(r['cost_cny']),
        'version':str(r['revision']),'updatedAt':db._iso(r['updated_at']) if r['updated_at'] else None} for r in rows],
        'total':total,'page':page,'pageSize':size,'costBasis':COST_BASIS}


def normalized(payload):
    if not isinstance(payload,dict):raise ValueError('请求格式无效')
    if any(key in payload for key in ('effectiveDate','effectiveAt')):
        raise ValueError('整款成本从服务端保存时间起生效，请刷新页面后保存')
    changes=payload.get('changes')
    if not isinstance(changes,list) or not 1<=len(changes)<=2000:raise ValueError('每次保存 1–2000 个款号成本')
    result=[];seen=set()
    for row in changes:
        if not isinstance(row,dict):raise ValueError('成本格式无效')
        code=row.get('styleCode')
        if not isinstance(code,str) or not code.strip() or len(code)>200 or '\x00' in code:raise ValueError('请选择有效款号')
        code=code.strip()
        if 'productId' in row or 'sizeCode' in row:raise ValueError('成本不再区分颜色尺码，请刷新页面后保存')
        if 'effectiveDate' in row or 'effectiveAt' in row:
            raise ValueError('成本生效时间由服务端确定，不接受自定义时间')
        if code in seen:raise ValueError('同一款号重复')
        seen.add(code)
        if 'costCny' not in row:raise ValueError('缺少成本字段；填写金额或显式传入 null 清空')
        cost=None if row['costCny'] is None else money(number(row['costCny']))
        result.append({'styleCode':code,'costCny':cost,'version':str(integer(row.get('version'),0))})
    return sorted(result,key=lambda r:r['styleCode'])


def preview(payload,lock=False):
    changes=normalized(payload);codes=[r['styleCode'] for r in changes]
    if lock:
        # Also serialize first writes, when no cost/version row exists yet. Stable
        # lock order avoids deadlocks between overlapping multi-style batches.
        for code in codes:
            db._fetch_one('SELECT pg_advisory_xact_lock(hashtextextended(%s,0))',('finance-style:'+code,))
    products=db._fetch_all(f'SELECT p.id,{style_catalog.style_sql()} AS style_code FROM products p WHERE {style_catalog.style_sql()}=ANY(%s) ORDER BY p.id'+(' FOR SHARE' if lock else ''),(codes,))
    existing={p['style_code'] for p in products}
    costs={r['style_code']:r for r in db._fetch_all('SELECT * FROM finance_style_costs WHERE style_code=ANY(%s) ORDER BY style_code'+(' FOR UPDATE' if lock else ''),(codes,))}
    initialized={r['style_code'] for r in db._fetch_all('SELECT style_code FROM finance_style_cost_audit WHERE style_code=ANY(%s) AND is_initial',(codes,))}
    initial_codes=[r['styleCode'] for r in changes if r['costCny'] is not None and r['styleCode'] not in initialized]
    # Only an initial amount fills missing historical orders. Later changes,
    # including re-entering a cleared cost, never reprice existing orders.
    orders=db._fetch_all(f'''SELECT DISTINCT o.id,{style_catalog.style_sql()} AS style_code
      FROM orders o JOIN order_items i ON i.order_id=o.id JOIN products p ON p.id=i.product_id
      WHERE {style_catalog.style_sql()}=ANY(%s) AND o.status=ANY(%s)''',(initial_codes,STATES)) if initial_codes else []
    output=[];affected=set()
    for row in changes:
        code=row['styleCode'];old=costs.get(code)
        if code not in existing:raise workbench.Conflict('款号已变更或商品已移除，请刷新后核对成本草稿')
        if row['version']!=str(old['revision'] if old else 0):raise workbench.Conflict('成本已被其他人更新，请刷新后核对；当前输入已保留')
        matching={r['id'] for r in orders if r['style_code']==code};affected.update(matching)
        output.append({**row,'oldCostCny':money(old['cost_cny']) if old else None,
            'initialCost':code in initial_codes,'affectedOrders':len(matching)})
    # Do not put a moving timestamp in the preview hash. The database assigns it
    # on save; only the first amount also applies before that timestamp.
    result={'changes':output,'affectedOrders':len(affected),'costBasis':COST_BASIS,'effectiveMode':'initial_then_on_save'}
    return {**result,'previewHash':digest(result)}


def cost_save_time():
    """Use the database clock once, after all style/version locks are held."""
    return db._fetch_one('SELECT clock_timestamp() AS saved_at')['saved_at']


def save_costs(payload):
    if not isinstance(payload,dict):raise ValueError('请求格式无效')
    try:rid=str(uuid.UUID(str(payload.get('requestId',''))))
    except (ValueError,AttributeError):raise ValueError('缺少有效提交编号')
    changes=normalized(payload)
    fingerprint=digest({'costBasis':COST_BASIS,'changes':changes,'previewHash':payload.get('previewHash')})
    db._fetch_one('SELECT pg_advisory_xact_lock(hashtextextended(%s,0))',('finance:'+rid,))
    previous=db._fetch_one('SELECT * FROM finance_cost_receipts WHERE request_id=%s',(rid,))
    if previous:
        if previous['actor_id']!=g.current_user['id'] or previous['fingerprint']!=fingerprint:raise workbench.Conflict('提交编号已用于其他修改')
        return {**previous['result'],'replayed':True}
    result=preview(payload,lock=True)
    if result['previewHash']!=payload.get('previewHash'):raise workbench.Conflict('预览已变化，请重新预览确认后保存')
    actor={key:g.current_user.get(key) for key in ('id','name','role')}
    effective_at=cost_save_time()
    for row in result['changes']:
        cost=None if row['costCny'] is None else Decimal(row['costCny'])
        old=None if row['oldCostCny'] is None else Decimal(row['oldCostCny'])
        db._fetch_one("""INSERT INTO finance_style_costs(style_code,cost_cny,revision,actor,updated_at) VALUES(%s,%s,1,%s,%s)
            ON CONFLICT(style_code) DO UPDATE SET cost_cny=EXCLUDED.cost_cny,
            revision=finance_style_costs.revision+1,actor=EXCLUDED.actor,updated_at=EXCLUDED.updated_at RETURNING revision""",(row['styleCode'],cost,Jsonb(actor),effective_at))
        db._fetch_one("""INSERT INTO finance_style_cost_audit(style_code,old_cost_cny,cost_cny,actor,request_id,created_at,effective_at,is_initial)
            VALUES(%s,%s,%s,%s,%s,%s,%s,%s) RETURNING id""",(row['styleCode'],old,cost,Jsonb(actor),rid,effective_at,effective_at,row['initialCost']))
    response={'saved':len(changes),'replayed':False,'effectiveAt':db._iso(effective_at)}
    db._fetch_one('INSERT INTO finance_cost_receipts(request_id,actor_id,fingerprint,result) VALUES(%s,%s,%s,%s) RETURNING request_id',(rid,g.current_user['id'],fingerprint,Jsonb(response)))
    return response


def history(args):
    if 'styleCode' in args:
        return style_history(args)
    pid=integer(args.get('productId')); size=str(args.get('sizeCode',''))
    page=integer(args.get('page',1))
    timeline=db._fetch_all('SELECT effective_date,cost_cny,actor,updated_at FROM finance_cost_history WHERE product_id=%s AND size_code=%s ORDER BY effective_date DESC',(pid,size))
    audits=db._fetch_all('SELECT * FROM finance_cost_audit WHERE product_id=%s AND size_code=%s ORDER BY id DESC LIMIT 10 OFFSET %s',(pid,size,(page-1)*10))
    return {'timeline':[{'effectiveDate':str(r['effective_date']),'costCny':money(r['cost_cny']),'actor':r['actor'],'updatedAt':db._iso(r['updated_at'])} for r in timeline],
        'items':[{'id':r['id'],'effectiveDate':str(r['effective_date']),'before':r['before_data'],'after':r['after_data'],'actor':r['actor'],'createdAt':db._iso(r['created_at'])} for r in audits],
        'total':db._fetch_one('SELECT count(*) AS n FROM finance_cost_audit WHERE product_id=%s AND size_code=%s',(pid,size))['n'],'page':page}


def style_history(args):
    code=str(args.get('styleCode','')).strip()
    if not code or len(code)>200:raise ValueError('请选择有效款号')
    page=integer(args.get('page',1))
    if integer(args.get('pageSize',20))!=20:raise ValueError('成本历史每页显示 20 条')
    rows=db._fetch_all('SELECT * FROM finance_style_cost_audit WHERE style_code=%s ORDER BY id DESC LIMIT 20 OFFSET %s',(code,(page-1)*20))
    return {'styleCode':code,'items':[{'id':r['id'],'styleCode':r['style_code'],'oldCostCny':money(r['old_cost_cny']),
        'costCny':money(r['cost_cny']),'actor':r['actor'],'createdAt':db._iso(r['created_at']),
        'effectiveAt':db._iso(r['effective_at']),'initialCost':r['is_initial']} for r in rows],
        'page':page,'pageSize':20,'total':db._fetch_one('SELECT count(*) AS n FROM finance_style_cost_audit WHERE style_code=%s',(code,))['n']}


def report(payload, paginate=True):
    with localcontext() as context:
        context.prec=60  # Quantity, stored amount, FX and cross-order sums stay exact.
        return _report(payload,paginate)


def _report(payload, paginate=True):
    if not isinstance(payload,dict): raise ValueError('筛选条件无效')
    end=day(payload.get('dateTo')); start=day(payload.get('dateFrom') or end-timedelta(days=29))
    if start>end: raise ValueError('开始日期不能晚于结束日期')
    rate=number(payload.get('exchangeRate'),rate=True)
    states=payload.get('statuses') or STATES
    if not isinstance(states,list) or any(s not in STATES for s in states): raise ValueError('仅支持有效订单状态')
    codes=terms(payload.get('styleCodes',[]))
    people=payload.get('salespersonIds',[])
    if not isinstance(people,list) or len(people)>200: raise ValueError('业务员筛选无效')
    owners=[integer(i) for i in people if i!='unassigned']
    params=[start,end+timedelta(days=1),states]
    where=["o.created_at >= (%s::date::timestamp AT TIME ZONE 'Asia/Shanghai')", "o.created_at < (%s::date::timestamp AT TIME ZONE 'Asia/Shanghai')",'o.status=ANY(%s)']
    customer=str(payload.get('customerName','')).strip()
    if customer: where.append('o.contact_name ILIKE %s');params.append('%'+customer+'%')
    if payload.get('customerId'): where.append('o.store_user_id=%s');params.append(integer(payload['customerId']))
    if people:
        where.append('(o.owner_admin_id=ANY(%s)'+(' OR o.owner_admin_id IS NULL' if 'unassigned' in people else '')+')');params.append(owners)
    style=f"CASE WHEN p.id IS NOT NULL THEN {style_catalog.style_sql()} ELSE COALESCE(ch.style_code,'未归类') END"
    if codes: where.append(f'LOWER({style})=ANY(%s)');params.append([v.lower() for v in codes])
    rows=db._fetch_all(f'''SELECT o.id AS order_id,o.order_no,o.created_at,o.status,o.contact_name,
        o.store_user_id,su.name AS account_name,su.company_name,COALESCE(a.name,'未分配') AS salesperson,
        o.owner_admin_id,i.id AS item_id,i.product_id,i.sku,i.size_code,i.quantity,i.unit_price,
        {style} AS style_code,h.cost_cny,h.effective_at AS cost_effective_at,h.is_initial
      FROM orders o JOIN order_items i ON i.order_id=o.id
      LEFT JOIN store_users su ON su.id=o.store_user_id LEFT JOIN admin_users a ON a.id=o.owner_admin_id
      LEFT JOIN products p ON p.id=i.product_id
      LEFT JOIN LATERAL (SELECT style_code FROM finance_cost_history c WHERE c.product_id=i.product_id AND c.size_code=i.size_code ORDER BY effective_date DESC LIMIT 1) ch ON true
      LEFT JOIN LATERAL (SELECT c.cost_cny,c.effective_at,c.is_initial FROM finance_style_cost_audit c
        WHERE c.style_code=({style}) AND (c.effective_at<=o.created_at OR c.is_initial)
        ORDER BY c.effective_at DESC,c.id DESC LIMIT 1) h ON true
      WHERE {' AND '.join(where)} ORDER BY o.created_at DESC,o.id DESC,i.id''',params)
    groups={}; lines=[]
    for r in rows:
        usd=r['unit_price']*r['quantity']; cny=(usd*rate).quantize(CENT,rounding=ROUND_HALF_UP)
        cost=r['cost_cny']*r['quantity'] if r['cost_cny'] is not None else None
        line={'orderId':r['order_id'],'orderNo':r['order_no'],'createdAt':db._iso(r['created_at']),
            'customerName':r['contact_name'],'accountName':r['company_name'] or r['account_name'] or '',
            'salesperson':r['salesperson'],'styleCode':r['style_code'],'sku':r['sku'],'sizeCode':r['size_code'] or '',
            'productId':r['product_id'],'quantity':r['quantity'],'salesUsd':money(usd),'salesCny':money(cny),
            'unitCostCny':money(r['cost_cny']),'costCny':money(cost),'profitCny':money(cny-cost) if cost is not None else None,
            'costEffectiveAt':db._iso(r['cost_effective_at']) if r['cost_effective_at'] else None,
            'costMatch':('首次成本补算' if r['is_initial'] and r['cost_effective_at']>r['created_at'] else '下单时生效成本') if cost is not None else '未设置成本'}
        lines.append(line)
        key=(r['style_code'],r['product_id'],r['sku'],r['size_code'])
        groups.setdefault(key,[]).append(line)
    def aggregate(items):
        missing=sum(v['quantity'] for v in items if v['costCny'] is None)
        total=lambda name:sum((Decimal(v[name]) for v in items if v[name] is not None),Decimal(0))
        return {'orderCount':len({v['orderId'] for v in items}),'quantity':sum(v['quantity'] for v in items),
            'salesUsd':money(total('salesUsd')),'salesCny':money(total('salesCny')),'knownCostCny':money(total('costCny')),
            'totalCostCny':None if missing else money(total('costCny')),
            'profitCny':None if missing else money(total('profitCny')),'knownProfitCny':money(total('profitCny')),'missingQuantity':missing}
    items=[]
    for key, values in groups.items():
        costs=[Decimal(v['unitCostCny']) for v in values if v['unitCostCny'] is not None]
        items.append({'styleCode':key[0],'productId':key[1],'sku':key[2],'sizeCode':key[3] or '',**aggregate(values),
            'minCostCny':money(min(costs)) if costs else None,'maxCostCny':money(max(costs)) if costs else None,'lines':values})
    items.sort(key=lambda r:(r['styleCode'],r['sku'],r['sizeCode']))
    page=integer(payload.get('page',1)); page_size=integer(payload.get('pageSize',25))
    if page_size not in (25,50,100): raise ValueError('分页无效')
    filters={'dateFrom':str(start),'dateTo':str(end),'customerName':customer,'customerId':payload.get('customerId'),
        'salespersonIds':people,'styleCodes':codes,'statuses':states,'exchangeRate':str(rate)}
    return {'items':items[(page-1)*page_size:page*page_size] if paginate else items,
        'total':len(items),'page':page,'pageSize':page_size,'summary':aggregate(lines),
        'filters':filters,'costBasis':COST_BASIS,'generatedAt':datetime.now(TZ).isoformat(),**({'lines':lines} if not paginate else {})}


def export(payload):
    # Existing Flask XLSX stack, not a spreadsheet authoring dependency for clients.
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    data=report(payload,paginate=False); book=Workbook(); summary=book.active;summary.title='核算汇总'
    def append(sheet, values):
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value,str): cell.data_type='s'  # Untrusted text never becomes a formula.
            if isinstance(cell.value,Decimal): cell.number_format='#,##0.00;[Red]-#,##0.00'
            cell.alignment=Alignment(vertical='center',wrap_text=True)
    amount=lambda v:Decimal(v) if v is not None else '待补全'
    append(summary,['GINGTTO 商品成本核算']); append(summary,['生成时间',data['generatedAt']])
    append(summary,['口径','每款首次成本补算此前订单；后续成本按订单下单时间匹配，从保存时间起生效，不覆盖旧订单；清空期间为未设置。美元销售额按本次汇率换算；不含运费、税费及其他费用；未付款订单不代表实际回款。'])
    labels={'dateFrom':'下单日期起（北京时间）','dateTo':'下单日期止（北京时间）','customerName':'客户名称关键词',
        'customerId':'商城账号','salespersonIds':'业务员','styleCodes':'款号','statuses':'订单状态','exchangeRate':'本次汇率（1 USD = CNY）',
        'orderCount':'涉及订单数','quantity':'商品件数','salesUsd':'商品销售额 USD','salesCny':'折合销售额 CNY',
        'knownCostCny':'已知成本 CNY','totalCostCny':'完整成本 CNY','profitCny':'商品毛利 CNY',
        'knownProfitCny':'已设置成本部分商品毛利 CNY','missingQuantity':'缺失成本件数'}
    lookup=options(); customers={c['id']:c['name'] for c in lookup['customers']}; people={c['id']:c['name'] for c in lookup['salespeople']}
    states={'pending_payment':'待付款','allocated':'已配货','paid':'已付款','shipped':'已发货','completed':'已完成'}
    for k,v in data['filters'].items():
        if k=='customerId': v=customers.get(v,str(v)) if v else '全部'
        if k=='salespersonIds': v=[('未分配' if i=='unassigned' else people.get(i,people.get(int(i),str(i)))) for i in v]
        if k=='statuses': v=[states[s] for s in v]
        if isinstance(v,list): v='、'.join(map(str,v)) or '全部'
        append(summary,[labels[k],str(v or '全部')])
    for k,v in data['summary'].items(): append(summary,[labels[k],amount(v) if k.endswith(('Cny','Usd')) else v])
    headers=['款号','颜色 SKU','真实尺码','订单数','件数','销售额 USD','折合销售额 CNY','成本下限 CNY','成本上限 CNY','已知成本 CNY','完整成本 CNY','商品毛利 CNY','缺失成本件数']
    append(summary,headers); header_row=summary.max_row
    for r in data['items']:
        append(summary,[r['styleCode'],r['sku'],r['sizeCode'],r['orderCount'],r['quantity'],*[amount(r[k]) for k in ('salesUsd','salesCny','minCostCny','maxCostCny','knownCostCny','totalCostCny','profitCny')],r['missingQuantity']])
    detail=book.create_sheet('订单明细');missing=book.create_sheet('成本缺失')
    columns=[('orderNo','订单号'),('createdAt','下单时间（北京时间）'),('customerName','客户'),('accountName','商城账号'),('salesperson','业务员'),('styleCode','款号'),('sku','SKU'),('sizeCode','尺码'),('quantity','数量'),('salesUsd','销售额 USD'),('salesCny','折合销售额 CNY'),('unitCostCny','单件成本 CNY'),('costEffectiveAt','成本生效时间（北京时间）'),('costCny','成本 CNY'),('profitCny','商品毛利 CNY'),('costMatch','成本匹配方式')]
    for sheet in (detail,missing): append(sheet,[label for _,label in columns])
    for row in data['lines']:
        vals=[amount(row[k]) if k.endswith(('Cny','Usd')) else datetime.fromisoformat(row[k]).astimezone(TZ).strftime('%Y-%m-%d %H:%M:%S') if k in ('createdAt','costEffectiveAt') and row[k] else row[k] for k,_ in columns]
        append(detail,vals)
        if row['costCny'] is None: append(missing,vals)
    if missing.max_row==1: append(missing,['无缺失成本'])
    for sheet,head in ((summary,header_row),(detail,1),(missing,1)):
        sheet.freeze_panes=f'A{head+1}';sheet.auto_filter.ref=f'A{head}:{get_column_letter(sheet.max_column)}{sheet.max_row}'
        sheet.sheet_view.showGridLines=False; sheet.sheet_properties.pageSetUpPr.fitToPage=True
        sheet.page_setup.orientation='landscape';sheet.page_setup.paperSize=sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=0;sheet.print_title_rows=f'{head}:{head}'
        for cell in sheet[head]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='243850')
        for i in range(1,sheet.max_column+1): sheet.column_dimensions[get_column_letter(i)].width=36 if i==1 else 22 if i<=5 else 18
        for cells in sheet:
            # Excel does not automatically grow wrapped rows generated by openpyxl.
            def wrapped_lines(cell):
                width=sheet.column_dimensions[cell.column_letter].width-3
                return sum(1+sum(2 if unicodedata.east_asian_width(ch) in ('W','F') else 1 for ch in line)//width for line in str(cell.value or '').split('\n'))
            sheet.row_dimensions[cells[0].row].height=max(30,min(409,16*max(wrapped_lines(c) for c in cells)))
    out=BytesIO();book.save(out);out.seek(0)
    return send_file(out,as_attachment=True,download_name='finance-cost-report.xlsx',mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


def register(app, require_auth):
    def guard(write=False):
        def decorate(fn):
            @wraps(fn)
            @require_auth
            def wrapped(*args,**kwargs):
                if not module_permissions.allows_request(g.current_user, request.path, request.method):
                    return jsonify({'message':'Forbidden'}),403
                try:
                    response=(workbench.atomic_admin_call if write else workbench.snapshot_admin_call)(fn,args,kwargs)
                    if hasattr(response,'headers'):response.headers['Cache-Control']='no-store'
                    return response
                except workbench.Conflict as exc:return jsonify({'message':str(exc),'code':'finance_conflict'}),409
                except ValueError as exc:return jsonify({'message':str(exc)}),400
            return wrapped
        return decorate
    routes=[('preferences','GET',lambda:jsonify(preferences()),False),
        ('preferences','PUT',lambda:jsonify(save_preferences(request.get_json(silent=True))),True),
        ('options','GET',lambda:jsonify(options()),False),('costs','GET',lambda:jsonify(costs(request.args)),False),
        ('costs/history','GET',lambda:jsonify(history(request.args)),False),
        ('costs/preview','POST',lambda:jsonify(preview(request.get_json(silent=True))),False),
        ('costs/save','POST',lambda:jsonify(save_costs(request.get_json(silent=True))),True),
        ('reports/query','POST',lambda:jsonify(report(request.get_json(silent=True))),False),
        ('reports/export','POST',lambda:export(request.get_json(silent=True)),False)]
    for path,method,handler,write in routes:
        app.add_url_rule('/api/admin/finance/'+path,'finance_'+path.replace('/','_')+'_'+method,guard(write)(handler),methods=[method])
