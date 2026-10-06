"""Category-first catalog ordering; uses only test_workbench's isolated *_test DB."""
import importlib.util
import io
from pathlib import Path
import unittest
from concurrent.futures import ThreadPoolExecutor
from unittest.mock import patch

from openpyxl import load_workbook
from test_workbench import seed, application, db
import test_workbench as workbench_tests
import catalog_order
import workbench


class CatalogOrderTest(unittest.TestCase):
    call = workbench_tests.WorkbenchTest.call

    def setUp(self):
        with db._connect() as conn:
            token = db.REQUEST_CONNECTION.set(conn)
            try:
                self.tokens, self.products = seed()
            finally:
                db.REQUEST_CONNECTION.reset(token)
        self.app = application.app
        self.app.config['TESTING'] = True
        db._fetch_one('UPDATE products SET category_id=2 WHERE id=1 RETURNING id')

    def categories(self):
        return self.call('categories').json

    def reorder(self, ids, version=None, role='admin'):
        return self.call('categories/order', 'PUT', {
            'categoryIds': ids,
            'orderVersion': version or self.categories()['orderVersion'],
        }, role=role)

    def inventory_snapshot(self):
        return {
            'products': db._fetch_all('SELECT id,stock,price,updated_at FROM products ORDER BY id'),
            'sizes': db._fetch_all('SELECT * FROM product_size_prices ORDER BY id'),
            'orders': db._fetch_all('SELECT * FROM orders ORDER BY id'),
        }

    def product(self, code, category='denim'):
        return db.create_product({
            'categoryKey': category, 'productCode': code, 'sku': code,
            'slug': code.lower(), 'colorGroup': 'GT-2026', 'colorName': code,
            'title': '排序测试商品', 'sizes': ['M'],
            'sizePrices': [{'sizeCode': 'M', 'stock': 1, 'price': 9}],
            'image': '/uploads/fixture.jpg', 'sizeChartImage': '/uploads/fixture.jpg',
            'descriptionImage': '/uploads/fixture.jpg',
        })

    def assert_order(self, items, ids):
        self.assertEqual([p['id'] for p in items], ids)
        self.assertTrue(all(isinstance(p['categorySortOrder'], int) for p in items))

    def test_move_permissions_versions_audit_and_no_inventory_changes(self):
        before = self.inventory_snapshot()
        initial = self.categories()
        self.assertEqual(self.reorder([2, 1], role='warehouse').status_code, 403)
        response = self.reorder([2, 1], initial['orderVersion'], role='sales')
        self.assertEqual(response.status_code, 200, response.json)
        self.assertEqual([c['id'] for c in response.json['items']], [2, 1])
        self.assertEqual([c['sortOrder'] for c in response.json['items']], [0, 1])
        self.assertEqual(self.reorder([1, 2], initial['orderVersion']).status_code, 409)
        self.assertEqual(self.inventory_snapshot(), before)
        logs = self.call('audit-logs?module=categories').json['items']
        self.assertEqual(len(logs), 1)  # Unchanged sort values do not generate audit rows.
        self.assertTrue(all(log['actor']['role'] == 'sales' for log in logs))
        self.assertEqual(len({log['batchId'] for log in logs}), 1)
        self.assertEqual([c['key'] for c in self.call('catalog-options').json['items']], ['outerwear', 'denim'])

    def test_invalid_lists_and_transaction_rollback(self):
        before = self.categories()
        for ids in ([1, 1], [1], [1, 3], [], ['1', 2], [True, 2], [0, 2]):
            response = self.reorder(ids, before['orderVersion'])
            self.assertEqual(response.status_code, 400, response.json)
        fetch = db._fetch_one

        def fail_second_update(query, *args, **kwargs):
            if query.startswith('UPDATE product_categories') and args[0][1] == 1:
                raise RuntimeError('Injected category update failure')
            return fetch(query, *args, **kwargs)

        with patch.object(db, '_fetch_one', side_effect=fail_second_update):
            with self.assertRaises(RuntimeError):
                self.reorder([2, 1], before['orderVersion'])
        self.assertEqual(self.categories(), before)
        self.assertEqual(self.call('audit-logs?module=categories').json['total'], 0)

    def test_concurrent_reorders_and_changed_category_set(self):
        version = self.categories()['orderVersion']
        with ThreadPoolExecutor(max_workers=2) as pool:
            responses = list(pool.map(lambda _: self.reorder([2, 1], version), range(2)))
        self.assertEqual(sorted(r.status_code for r in responses), [200, 409])
        version = self.categories()['orderVersion']
        self.assertEqual(self.call('categories', 'POST', {
            'key': 'tops', 'sortOrder': 8, 'labels': {'zh': '上衣', 'en': 'Tops'},
        }).status_code, 201)
        self.assertEqual(self.reorder([1, 2], version).status_code, 409)

    def test_all_catalog_endpoints_and_explicit_sorts(self):
        self.assert_order(db.list_products(), [2, 3, 1])
        for endpoint, role in [('products', 'admin'), ('inventory', 'warehouse'), ('inventory', 'customer')]:
            for sort in ('id', 'updatedAt', 'stock', 'price', 'sku'):
                response = self.call(f'{endpoint}?page=1&sort={sort}&direction=desc', role=role)
                self.assertEqual(response.status_code, 200, response.json)
                items = response.json['items']
                self.assertEqual([p['categoryKey'] for p in items], ['denim', 'denim', 'outerwear'])
                self.assertEqual(items[-1]['id'], 1)
        configured = self.call('products?page=1&sort=configured&ids=1,2,3').json['items']
        self.assert_order(configured, [2, 3, 1])
        self.assert_order(self.call('products/1/family').json['items'], [2, 3, 1])
        self.assert_order(workbench.products_for_export({}), [2, 3, 1])
        export = self.call('inventory/export')
        sheet = load_workbook(io.BytesIO(export.data)).active
        product_column = [cell.value for cell in sheet[1]].index('商品ID')
        product_ids = list(dict.fromkeys(int(row[product_column]) for row in sheet.iter_rows(min_row=2, values_only=True)))
        self.assertEqual(product_ids, [2, 3, 1])
        self.assertEqual([c['key'] for c in self.call('dashboard?view=workbench').json['filters']['categories']], ['denim', 'outerwear'])
        self.assertEqual(self.reorder([2, 1]).status_code, 200)
        self.assert_order(self.call('products?page=1&sort=configured&ids=3,1,2').json['items'], [1, 3, 2])
        self.assert_order(db.list_products(), [1, 2, 3])

    def test_category_order_precedes_pagination_and_new_products(self):
        with db._connect() as conn:
            token = db.REQUEST_CONNECTION.set(conn)
            try:
                for index in range(26):
                    self.product(f'NEW-{index}')
            finally:
                db.REQUEST_CONNECTION.reset(token)
        response = self.call('products?page=1&pageSize=25&sort=id&direction=desc').json
        self.assertEqual(response['total'], 29)
        self.assertTrue(all(p['categoryKey'] == 'denim' for p in response['items']))
        self.assertEqual(self.call('products?page=2&pageSize=25&sort=id&direction=desc').json['items'][-1]['id'], 1)
        self.assertEqual(self.reorder([2, 1]).status_code, 200)
        self.assertEqual(self.call('products?page=1&pageSize=25').json['items'][0]['id'], 1)
        new_outerwear = self.product('NEW-OUTER', 'outerwear')
        items = self.call('products?page=1&pageSize=25').json['items']
        self.assert_order(items[:2], [new_outerwear['id'], 1])
        db._fetch_one('UPDATE products SET category_id=1 WHERE id=%s RETURNING id', (new_outerwear['id'],))
        self.assertEqual(self.call('products?page=1&pageSize=25').json['items'][0]['id'], 1)

    def test_admin_tables_category_mode_cached_sorts_and_inventory_pagination(self):
        with db._connect() as conn:
            token = db.REQUEST_CONNECTION.set(conn)
            try:
                for index in range(26):
                    self.product(f'PAGE-{index}')
                # The newest product belongs to the last category. A global
                # time sort would put it on page 1 and split the first category.
                db._fetch_one("UPDATE products SET updated_at='2030-01-01 00:00:00+00' WHERE id=1 RETURNING id")
                db.update_product_inventory(1, {'M': -8, 'L': 10, 'Tall XL': 0})
            finally:
                db.REQUEST_CONNECTION.reset(token)
        before = self.inventory_snapshot()
        for endpoint, role in [('products', 'admin'), ('inventory', 'warehouse'), ('inventory', 'customer')]:
            for sort in ('category', 'updatedAt', 'stock'):
                for direction in ('asc', 'desc'):
                    pages = [self.call(f'{endpoint}?page={page}&pageSize=25&sort={sort}&direction={direction}', role=role).json
                             for page in (1, 2)]
                    items = [item for response in pages for item in response['items']]
                    self.assertEqual([item['categoryKey'] for item in items], ['denim'] * 28 + ['outerwear'])
                    self.assertEqual(len({item['id'] for item in items}), 29)
                    self.assertTrue(all(item['categoryKey'] == 'denim' for item in pages[0]['items']))
            category = self.call(f'{endpoint}?page=1&pageSize=25&sort=category&direction=asc', role=role).json
            default = self.call(f'{endpoint}?page=1&pageSize=25', role=role).json
            self.assertEqual([p['id'] for p in category['items']], [p['id'] for p in default['items']])
            filtered = self.call(f'{endpoint}?page=1&category=outerwear&sort=category', role=role).json
            self.assertEqual([p['id'] for p in filtered['items']], [1])
        self.assertEqual(self.inventory_snapshot(), before)
        self.assertEqual(self.reorder([2, 1]).status_code, 200)
        for endpoint in ('products', 'inventory'):
            items = self.call(f'{endpoint}?page=1&sort=category').json['items']
            self.assertEqual(items[0]['id'], 1)
            self.assertTrue(all(p['categoryKey'] == 'denim' for p in items[1:]))
        self.assertEqual(self.inventory_snapshot(), before)

    def test_home_activity_and_category_ties(self):
        config = db.get_homepage_config()
        config['heroBanners'] = {key: '/uploads/fixture.jpg' for key in db.HOME_SECTION_KEYS}
        config['sectionProductIds']['bestSeller'] = [1, 3, 2]
        config['collectionProductIds']['bestSeller'] = [1, 3, 2]
        config['displayCategoryKeys'] = ['outerwear', 'denim']
        response = self.call('home-config', 'PUT', {**config, 'version': workbench.version('homepage_configs', 1)})
        self.assertEqual(response.status_code, 200, response.json)
        for endpoint in ('home-config', 'activity-config'):
            config = self.call(endpoint).json['config']
            self.assertEqual(config['sectionProductIds']['bestSeller'], [3, 2, 1])
            self.assertEqual(config['collectionProductIds']['bestSeller'], [3, 2, 1])
            self.assertEqual(config['displayCategoryKeys'], ['denim', 'outerwear'])
        self.reorder([2, 1])
        config = db.get_homepage_config()
        self.assertEqual(config['sectionProductIds']['bestSeller'], [1, 3, 2])
        self.assertEqual(config['displayCategoryKeys'], ['outerwear', 'denim'])
        db._fetch_one('UPDATE product_categories SET sort_order=9 RETURNING id')
        self.assert_order(db.list_products(), [2, 3, 1])

    def test_storefront_parity_public_stock_and_colors(self):
        source = Path(__file__).resolve().parents[2] / 'shopify' / 'storefront-backend'
        if not source.exists():
            self.skipTest('Sibling storefront checkout required for cross-service regression')
        self.assertEqual((source / 'catalog_order.py').read_bytes(), Path(catalog_order.__file__).read_bytes())
        spec = importlib.util.spec_from_file_location('catalog_storefront_db', source / 'db.py')
        store = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(store)
        db.update_product_inventory(1, {'M': -10, 'L': 10, 'Tall XL': 0})
        self.assert_order(store.list_products(), [2, 3, 1])
        self.assertEqual(store.get_product_by_slug('gt-2026-1')['stock'], 10)
        self.assertEqual([c['id'] for c in store.get_product_by_slug('gt-2026-1')['colorOptions']], [2, 3, 1])
        self.assertEqual([c['id'] for c in db.get_product_by_id(1)['colorOptions']], [2, 3, 1])
        config = db.get_homepage_config()
        config['sectionProductIds']['bestSeller'] = [1, 3, 2]
        config['collectionProductIds']['bestSeller'] = [1, 3, 2]
        config['displayCategoryKeys'] = ['outerwear', 'denim']
        db.save_homepage_config(config)
        self.assertEqual(store.get_homepage_config(), db.get_homepage_config())
        self.reorder([2, 1])
        self.assert_order(store.list_products(), [1, 2, 3])
        self.assertEqual(store.get_homepage_config(), db.get_homepage_config())
        self.assertEqual(store.get_product_by_slug('gt-2026-1')['sizePrices'][0]['stock'], 0)
        self.assertEqual(db.get_product_by_id(1)['sizePrices'][0]['stock'], -10)

    def test_dashboard_product_details_and_style_filters(self):
        response = self.call('dashboard?view=style-detail&styleCode=GT-2026')
        self.assertEqual(response.status_code, 200, response.json)
        self.assertEqual([p['productId'] for p in response.json['items']], [2, 3, 1])
        db._fetch_one("UPDATE products SET style_code=CASE WHEN id=1 THEN 'AAA' ELSE 'ZZZ' END RETURNING id")
        self.assertEqual(self.call('dashboard?view=workbench').json['filters']['styles'], ['ZZZ', 'AAA'])
        self.reorder([2, 1])
        self.assertEqual(self.call('dashboard?view=workbench').json['filters']['styles'], ['AAA', 'ZZZ'])


if __name__ == '__main__':
    unittest.main()
