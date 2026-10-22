"""Finance regression writes are restricted by test_workbench to local *_test databases."""
import copy
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta, timezone
from decimal import Decimal
from io import BytesIO
import unittest
import uuid
from unittest.mock import patch
from openpyxl import load_workbook
from werkzeug.security import generate_password_hash
import test_orders as helpers
from test_workbench import db, application
import finance

DATABASE_COST_CLOCK=finance.cost_save_time


class FinanceTest(unittest.TestCase):
    call=helpers.AdminOrdersTest.call
    payload=helpers.AdminOrdersTest.payload
    stock=helpers.AdminOrdersTest.stock

    def clear(self):
        with db.get_connection() as conn:
            conn.execute('TRUNCATE finance_cost_history,finance_cost_versions,finance_cost_audit,finance_cost_receipts,finance_settings,finance_style_costs,finance_style_cost_audit RESTART IDENTITY')

    def setUp(self):
        self.clear();helpers.AdminOrdersTest.setUp(self)
        # Simulate a baseline saved before the September fixture orders. Browser
        # and production requests never accept a client-provided effective time.
        clock=patch.object(finance,'cost_save_time',return_value=datetime(2026,9,1,tzinfo=timezone.utc))
        clock.start();self.addCleanup(clock.stop)
        user=db.create_admin_user({'name':'Finance fixture','email':'finance@gingtto.test','passwordHash':generate_password_hash('Fixture-only',method='pbkdf2:sha256:1000'),'status':'active','role':'finance'})
        self.finance_id=user['id'];self.tokens['finance']=db.create_admin_session(user['id'])

    def tearDown(self): self.clear()

    def post(self,path,body,role='finance'):
        return self.call('finance/'+path,'POST',body,role)

    def change(self,code='GT-2026',cost='10.00'):
        row=db._fetch_one('SELECT revision FROM finance_style_costs WHERE style_code=%s',(code,))
        return {'changes':[{'styleCode':code,'costCny':cost,'version':str(row['revision'] if row else 0)}]}

    def prepared(self,body):
        response=self.post('costs/preview',body)
        self.assertEqual(response.status_code,200,response.json)
        return {**body,'previewHash':response.json['previewHash'],'requestId':str(uuid.uuid4())}

    def save_cost(self,at=None,**kwargs):
        body=self.prepared(self.change(**kwargs))
        if at:
            with patch.object(finance,'cost_save_time',return_value=datetime.fromisoformat(at)):
                response=self.post('costs/save',body)
        else:response=self.post('costs/save',body)
        self.assertEqual(response.status_code,200,response.json)
        return response.json

    def test_sql_summary_and_page_match_complete_report(self):
        self.save_cost(cost='0.00')
        with db.get_connection() as conn:
            for index in range(32):
                conn.execute("INSERT INTO order_items(order_id,product_id,product_name,sku,size_code,quantity,unit_price,total_price) SELECT order_id,product_id,product_name,sku,%s,quantity,unit_price,total_price FROM order_items ORDER BY id LIMIT 1",('Size-'+str(index),))
        payload={'dateFrom':'2026-09-01','dateTo':'2026-09-30','exchangeRate':'7.123456'}
        complete=finance.report(payload,paginate=False)
        first=self.query(page=1);second=self.query(page=2)
        self.assertEqual(first['summary'],complete['summary'])
        self.assertEqual(second['summary'],complete['summary'])
        self.assertEqual(first['items']+second['items'],complete['items'])
        self.assertEqual(len(first['items']),25)
        self.assertEqual(self.query(page=99)['items'],[])

    def query(self,**kwargs):
        body={'dateFrom':'2026-09-01','dateTo':'2026-09-30','exchangeRate':'7.123456',**kwargs}
        response=self.post('reports/query',body)
        self.assertEqual(response.status_code,200,response.json)
        return response.json

    def test_roles_permissions_and_private_options(self):
        for role in ('sales','warehouse','customer'):
            for path in ('options','costs','costs/history?productId=1&sizeCode=M','costs/history?styleCode=GT-2026'):
                self.assertEqual(self.call('finance/'+path,role=role).status_code,403)
            for path in ('costs/preview','costs/save','reports/query','reports/export'):
                self.assertEqual(self.post(path,{},role).status_code,403)
        self.assertEqual(self.call('finance/options',role='finance').status_code,200)
        self.assertEqual(self.call('products',role='finance').status_code,200)  # Contract picker dependency only.
        for path in ('orders','inventory','admin-users','store-users'):
            self.assertEqual(self.call(path,role='finance').status_code,403)
        self.assertEqual(self.call('orders','POST',self.payload(),'finance').status_code,403)
        db._fetch_one("UPDATE admin_users SET permissions='[]'::jsonb WHERE id=%s RETURNING id",(self.finance_id,))
        self.assertEqual(self.call('finance/options',role='finance').status_code,403)
        self.assertEqual(self.call('finance/options').status_code,200)
        db._fetch_one("UPDATE admin_users SET permissions='[\"orders\"]'::jsonb WHERE id=1 RETURNING id")
        self.assertEqual(self.call('finance/options').status_code,403)

    def test_new_cost_only_applies_after_save_and_history_is_not_repriced(self):
        initial=self.query();self.assertEqual(initial['summary']['quantity'],9)
        self.assertEqual(initial['summary']['missingQuantity'],9);self.assertIsNone(initial['summary']['profitCny'])
        first=self.save_cost(cost='12',at='2026-09-15T12:00:00+08:00')
        original=self.query();self.assertEqual(original['summary']['totalCostCny'],'108.00')
        self.save_cost(cost='18',at='2026-09-18T09:00:00+08:00')
        result=self.query();self.assertEqual(result['summary']['totalCostCny'],'126.00')
        # Sep 16 and Sep 18 00:30 Beijing use 12; Sep 18 17:00 uses 18.
        for row in result['items']:
            for line in row['lines']:
                self.assertEqual(line['unitCostCny'],'18.00' if line['orderId']==2 else '12.00')
            self.assertEqual((row['minCostCny'],row['maxCostCny']),('12.00','18.00'))
        self.assertEqual(self.query(dateFrom='2026-09-16',dateTo='2026-09-16')['summary']['totalCostCny'],'36.00')
        self.save_cost(cost='0',at='2026-10-15T11:00:00+08:00')
        self.assertEqual(self.query()['summary'],result['summary'])
        self.assertEqual(self.call('finance/costs',role='finance').json['items'][0]['costCny'],'0.00')
        history=self.call('finance/costs/history?styleCode=GT-2026',role='finance').json
        self.assertEqual(history['total'],3)
        self.assertEqual(history['items'][0]['oldCostCny'],'18.00')
        self.assertIsNone(history['items'][-1]['oldCostCny'])
        self.assertEqual(history['items'][-1]['effectiveAt'],first['effectiveAt'])
        self.assertEqual(result['costBasis'],'style_at_order_time')

    def test_first_cost_backfills_older_orders_once(self):
        saved=self.save_cost(cost='18',at='2026-09-18T09:00:00+08:00')
        result=self.query();self.assertEqual(result['summary']['missingQuantity'],0)
        self.assertEqual(result['summary']['totalCostCny'],'162.00')
        for row in result['items']:
            for line in row['lines']:
                self.assertEqual(line['unitCostCny'],'18.00')
                self.assertEqual(line['costEffectiveAt'],saved['effectiveAt'])
                self.assertEqual(line['costMatch'],'下单时生效成本' if line['orderId']==2 else '首次成本补算')
        self.save_cost(cost='25',at='2026-10-15T12:00:00+08:00')
        self.assertEqual(result['summary'],self.query()['summary'])
        self.assertEqual(db._fetch_one('SELECT count(*) AS n FROM finance_style_cost_audit WHERE is_initial')['n'],1)

    def test_first_zero_cost_backfills_but_cleared_period_is_not_filled_by_restore(self):
        self.save_cost(cost=None,at='2026-09-01T00:00:00+08:00')  # Never configured yet.
        proposal=self.post('costs/preview',self.change(cost='0')).json
        self.assertTrue(proposal['changes'][0]['initialCost']);self.assertEqual(proposal['affectedOrders'],3)
        self.save_cost(cost='0',at='2026-09-17T23:00:00+08:00')
        self.assertEqual(self.query()['summary']['totalCostCny'],'0.00')
        self.assertEqual(self.query()['summary']['missingQuantity'],0)
        self.save_cost(cost=None,at='2026-09-18T09:00:00+08:00')
        proposal=self.post('costs/preview',self.change(cost='25')).json
        self.assertFalse(proposal['changes'][0]['initialCost']);self.assertEqual(proposal['affectedOrders'],0)
        self.save_cost(cost='25',at='2026-09-18T18:00:00+08:00')
        self.assertEqual(self.query()['summary']['missingQuantity'],3)
        self.assertEqual(self.query()['summary']['knownCostCny'],'0.00')
        before=self.query()['summary']
        with db.get_connection() as conn:finance.migrate(conn.cursor());finance.migrate(conn.cursor())
        self.assertEqual(self.query()['summary'],before)

    def test_real_new_orders_and_idempotent_retry_keep_original_cost_and_stock(self):
        original_stock=self.stock()['stock']
        with patch.object(finance,'cost_save_time',side_effect=DATABASE_COST_CLOCK):
            before=db._fetch_one('SELECT clock_timestamp() AS t')['t']
            first_cost=self.save_cost(cost='12')
            after=db._fetch_one('SELECT clock_timestamp() AS t')['t']
            self.assertLessEqual(before,datetime.fromisoformat(first_cost['effectiveAt']))
            self.assertLessEqual(datetime.fromisoformat(first_cost['effectiveAt']),after)
            body=self.payload();body['contactName']='Cost timeline orders'
            first=self.call('orders','POST',body)
            self.assertEqual(first.status_code,200,first.json)
            old_created=first.json['order']['createdAt']
            self.save_cost(cost='18')
            replay=self.call('orders','POST',body)
            self.assertEqual(replay.status_code,200,replay.json);self.assertTrue(replay.json['replayed'])
            self.assertEqual(replay.json['order']['createdAt'],old_created)
            next_body={**body,'requestId':str(uuid.uuid4())}
            second=self.call('orders','POST',next_body)
            self.assertEqual(second.status_code,200,second.json)
        dates=[datetime.fromisoformat(o.json['order']['createdAt']).astimezone(finance.TZ).date() for o in (first,second)]
        result=self.query(dateFrom=str(min(dates)),dateTo=str(max(dates)),customerName='Cost timeline orders')
        self.assertEqual(result['summary']['quantity'],10);self.assertEqual(result['summary']['totalCostCny'],'150.00')
        self.assertEqual({line['orderId']:line['unitCostCny'] for row in result['items'] for line in row['lines']},
            {first.json['order']['id']:'12.00',second.json['order']['id']:'18.00'})
        self.assertEqual(self.stock()['stock'],original_stock-10)

    def test_same_day_boundary_clear_gap_zero_and_restore_use_exact_timestamps(self):
        at='2026-09-18T10:00:00+08:00';boundary=datetime.fromisoformat(at)
        with db.get_connection() as conn:
            for oid,t in [(1,boundary-timedelta(microseconds=1)),(2,boundary),(4,boundary+timedelta(microseconds=1))]:
                conn.execute('UPDATE orders SET created_at=%s WHERE id=%s',(t,oid))
        self.save_cost(cost='12',at='2026-09-18T09:00:00+08:00')
        self.save_cost(cost='18',at=at)
        result=self.query();self.assertEqual(result['summary']['totalCostCny'],'144.00')
        self.assertEqual({line['orderId']:line['unitCostCny'] for row in result['items'] for line in row['lines']},{1:'12.00',2:'18.00',4:'18.00'})
        self.save_cost(cost=None,at='2026-09-18T10:00:00.000001+08:00')
        self.assertEqual(self.query()['summary']['missingQuantity'],3)  # No fallback past a clear event.
        self.assertEqual(self.query()['summary']['knownCostCny'],'90.00')
        self.save_cost(cost='0',at='2026-09-18T10:00:00.000002+08:00')
        # An existing order in the clear interval stays missing even after restore.
        self.assertEqual(self.query()['summary']['missingQuantity'],3)
        db._fetch_one('UPDATE orders SET created_at=%s WHERE id=4 RETURNING id',(boundary+timedelta(microseconds=2),))
        self.assertEqual(self.query()['summary']['totalCostCny'],'90.00')
        self.assertEqual(self.query()['summary']['missingQuantity'],0)

    def test_shared_exchange_rate_save_validation_retry_and_migration(self):
        initial={'exchangeRate':None,'version':'0'}
        self.assertEqual(self.call('finance/preferences',role='finance').json,initial)
        stock=self.stock();orders=db._fetch_one('SELECT count(*) AS n FROM orders')['n']
        for value in ('0','-1','NaN','Infinity','7.1234567',True,None,'', '10000000000'):
            response=self.call('finance/preferences','PUT',{'exchangeRate':value,'version':'0'},'finance')
            self.assertEqual(response.status_code,400,response.json)
        self.assertEqual(self.call('finance/preferences','PUT',[], 'finance').status_code,400)
        first={'exchangeRate':'7.123456','version':'0'}
        response=self.call('finance/preferences','PUT',first,'finance')
        self.assertEqual(response.status_code,200,response.json)
        self.assertEqual(response.json,{'exchangeRate':'7.123456','version':'1'})
        self.assertEqual(self.call('finance/preferences').json,response.json)  # Other account, same setting.
        self.assertEqual(self.call('finance/preferences','PUT',first,'finance').json,response.json)  # Lost response retry.
        changed=self.call('finance/preferences','PUT',{'exchangeRate':'7.000001','version':'1'})
        self.assertEqual(changed.status_code,200,changed.json)
        self.assertEqual(self.call('finance/preferences',role='finance').json,changed.json)
        self.assertEqual(self.call('finance/preferences','PUT',{'exchangeRate':'6.90','version':'1'},'finance').status_code,409)
        with db.get_connection() as conn:finance.migrate(conn.cursor());finance.migrate(conn.cursor())
        self.assertEqual(self.call('finance/preferences',role='finance').json,{'exchangeRate':'7.000001','version':'2'})
        self.assertEqual(db._fetch_one('SELECT count(*) AS n FROM finance_settings')['n'],1)
        self.assertEqual(self.stock(),stock);self.assertEqual(db._fetch_one('SELECT count(*) AS n FROM orders')['n'],orders)

    def test_shared_exchange_rate_concurrency_and_permissions(self):
        for role in ('sales','warehouse','customer'):
            self.assertEqual(self.call('finance/preferences',role=role).status_code,403)
            self.assertEqual(self.call('finance/preferences','PUT',{'exchangeRate':'7','version':'0'},role).status_code,403)
        def write(row):return self.call('finance/preferences','PUT',{'exchangeRate':row[1],'version':'0'},row[0])
        with ThreadPoolExecutor(max_workers=2) as pool:
            results=list(pool.map(write,[('admin','7.1'),('finance','7.2')]))
        self.assertEqual(sorted(r.status_code for r in results),[200,409])
        winner=next(r.json for r in results if r.status_code==200)
        self.assertEqual(self.call('finance/preferences').json,winner)
        self.assertEqual(self.call('finance/preferences',role='finance').json,winner)
        db._fetch_one("UPDATE admin_users SET permissions='[]'::jsonb WHERE id=%s RETURNING id",(self.finance_id,))
        self.assertEqual(self.call('finance/preferences',role='finance').status_code,403)
        self.assertEqual(self.call('finance/preferences','PUT',{'exchangeRate':'8','version':'1'},'finance').status_code,403)
        self.assertEqual(self.call('finance/preferences').json,winner)

    def test_one_row_per_style_no_titles_sizes_or_colors_in_cost_results(self):
        self.save_cost(cost='10');self.save_cost(cost='15');self.save_cost(cost='0')
        history=self.call('finance/costs/history?styleCode=GT-2026',role='finance').json
        self.assertEqual(history['total'],3)
        self.assertEqual((history['items'][0]['oldCostCny'],history['items'][0]['costCny']),('15.00','0.00'))
        self.assertEqual((history['items'][1]['oldCostCny'],history['items'][1]['costCny']),('10.00','15.00'))
        row=self.call('finance/costs',role='finance').json['items'][0]
        self.assertEqual(set(row),{'styleCode','costCny','version','updatedAt','active'})
        self.assertEqual(row['costCny'],'0.00')
        self.assertEqual(self.call('finance/costs?keyword=长标题测试',role='finance').json['total'],0)
        self.assertEqual(self.call('finance/costs?keyword=GT-2026-1',role='finance').json['total'],1)
        self.assertEqual(self.call('finance/costs?keyword=水洗灰',role='finance').json['total'],1)
        self.assertEqual(self.call('finance/costs?category=denim',role='finance').json['total'],1)
        self.assertEqual(self.call('finance/costs?missingOnly=true',role='finance').json['total'],0)
        self.assertEqual(self.call('finance/costs?styleCodes=GT-2026，OTHER',role='finance').json['total'],1)

    def test_cost_and_history_survive_removed_sizes_deleted_color_and_style_changes(self):
        self.save_cost(cost='19')
        with db.get_connection() as conn:
            conn.execute("DELETE FROM product_size_prices WHERE product_id=1 AND size_code='M'")
            conn.execute("UPDATE products SET style_code='OTHER-STYLE' WHERE id=2")
            conn.execute('DELETE FROM products WHERE id=3')
        data=self.query()
        self.assertEqual(data['summary']['knownCostCny'],'114.00')  # 6 pieces remain GT-2026, no size dependency
        self.assertEqual(data['summary']['missingQuantity'],3)
        self.assertEqual(self.call('finance/costs/history?styleCode=OTHER-STYLE',role='finance').json['total'],0)
        self.assertEqual(self.call('finance/costs/history?styleCode=GT-2026',role='finance').json['total'],1)
        with db.get_connection() as conn:conn.execute("UPDATE products SET style_code='OTHER-STYLE' WHERE id=1")
        self.assertEqual(self.call('finance/costs/history?styleCode=GT-2026',role='finance').json['total'],1)
        self.assertEqual(self.query()['summary']['missingQuantity'],9)

    def test_style_history_paging_validation_empty_and_read_only(self):
        for n in range(21):self.save_cost(cost=str(n))
        stock=self.stock();costs=db._fetch_all('SELECT * FROM finance_style_costs')
        first=self.call('finance/costs/history?styleCode=GT-2026',role='finance').json
        second=self.call('finance/costs/history?styleCode=GT-2026&page=2',role='finance').json
        self.assertEqual(first['total'],21);self.assertEqual(len(first['items']),20);self.assertEqual(len(second['items']),1)
        self.assertGreater(first['items'][-1]['id'],second['items'][0]['id'])
        self.assertEqual(self.call('finance/costs/history?styleCode=NO-SUCH-STYLE',role='finance').json['total'],0)
        for query in ('styleCode=','styleCode=GT-2026&page=0','styleCode=GT-2026&pageSize=100'):
            self.assertEqual(self.call('finance/costs/history?'+query,role='finance').status_code,400)
        self.assertEqual(self.stock(),stock);self.assertEqual(db._fetch_all('SELECT * FROM finance_style_costs'),costs)
        self.assertEqual(db._fetch_one('SELECT count(*) AS n FROM finance_style_cost_audit')['n'],21)

    def test_preview_is_prospective_and_rejects_variant_or_effective_time_clients(self):
        body=self.change(cost='15');r=self.post('costs/preview',body).json
        self.assertEqual(r['affectedOrders'],3);self.assertTrue(r['changes'][0]['initialCost'])
        self.assertEqual(r['costBasis'],'style_at_order_time');self.assertEqual(r['effectiveMode'],'initial_then_on_save')
        self.assertNotIn('effectiveDate',r)
        self.assertEqual(self.query()['summary']['missingQuantity'],9)  # Preview never persists.
        self.assertEqual(self.post('costs/preview',{**body,'effectiveDate':'2026-09-01'}).status_code,400)
        self.assertEqual(self.post('costs/preview',{**body,'effectiveAt':'2020-01-01T00:00:00Z'}).status_code,400)
        body['changes'][0]['effectiveAt']='2020-01-01T00:00:00Z'
        self.assertEqual(self.post('costs/preview',body).status_code,400)
        self.assertEqual(self.post('costs/preview',{'changes':[{'productId':1,'sizeCode':'M','version':'0','costCny':'5'}]}).status_code,400)
        self.assertEqual(self.post('costs/preview',self.change(code='MISSING')).status_code,409)
        self.save_cost(cost='12')
        r=self.post('costs/preview',self.change(cost='18')).json
        self.assertEqual(r['affectedOrders'],0);self.assertFalse(r['changes'][0]['initialCost'])

    def test_decimal_sales_are_order_prices_and_shipping_excluded(self):
        self.save_cost(cost='10')
        result=self.query();s=result['summary']
        self.assertEqual(s['salesUsd'],'269.10');self.assertEqual(s['totalCostCny'],'90.00')
        self.assertEqual(s['salesCny'],'1916.91');self.assertEqual(s['profitCny'],'1826.91')
        db._fetch_one('UPDATE product_size_prices SET price=900 WHERE product_id=1 RETURNING id')
        self.assertEqual(self.query()['summary'],s)

    def test_combined_customer_owner_multistyle_filters_only_matching_lines(self):
        db._fetch_one('UPDATE orders SET owner_admin_id=2 WHERE id=1 RETURNING id')
        db._fetch_one("UPDATE products SET style_code='STYLE-A' WHERE id=1 RETURNING id")
        db._fetch_one("UPDATE products SET style_code='STYLE-B' WHERE id=2 RETURNING id")
        self.save_cost(code='STYLE-A',cost='10');self.save_cost(code='STYLE-B',cost='5')
        self.assertEqual(self.query()['summary']['totalCostCny'],'75.00')
        r=self.query(customerName='Alex',customerId=1,salespersonIds=[2],styleCodes='STYLE-A，NOT-HERE')
        self.assertEqual(r['summary']['quantity'],2);self.assertEqual(r['summary']['salesUsd'],'59.80');self.assertEqual(r['summary']['totalCostCny'],'20.00')
        self.assertEqual(self.query(styleCodes=['STYLE-A','STYLE-B'])['summary']['quantity'],9)
        self.assertEqual(self.query(salespersonIds=['unassigned'])['summary']['quantity'],6)
        self.assertEqual(self.query(customerName='No match')['summary']['quantity'],0)
        self.assertEqual(self.query(statuses=['completed'])['summary']['quantity'],3)

    def test_save_is_atomic_idempotent_and_versioned_without_business_changes(self):
        stock=self.stock();count=db._fetch_one('SELECT count(*) AS n FROM orders')['n']
        body=self.prepared(self.change())
        with ThreadPoolExecutor(max_workers=2) as pool:responses=list(pool.map(lambda _:self.post('costs/save',body),range(2)))
        self.assertTrue(all(r.status_code==200 for r in responses),[r.json for r in responses])
        self.assertEqual(sum(r.json['replayed'] for r in responses),1)
        self.assertEqual(responses[0].json['effectiveAt'],responses[1].json['effectiveAt'])
        self.assertEqual(self.post('costs/save',{**body,'requestId':str(uuid.uuid4())}).status_code,409)
        changed=copy.deepcopy(body);changed['changes'][0]['costCny']='20'
        self.assertEqual(self.post('costs/save',changed).status_code,409)
        self.assertEqual(self.post('costs/save',body,role='admin').status_code,409)  # Receipts are actor-bound.
        self.assertEqual(db._fetch_one('SELECT count(*) AS n FROM finance_style_cost_audit')['n'],1)
        self.assertEqual(self.stock(),stock);self.assertEqual(db._fetch_one('SELECT count(*) AS n FROM orders')['n'],count)
        db._fetch_one("UPDATE products SET style_code='OTHER' WHERE id=3 RETURNING id")
        batch=self.change(code='OTHER');batch['changes']+=self.change()['changes'];batch['changes'][1]['version']='0'
        self.assertEqual(self.post('costs/save',{**batch,'requestId':str(uuid.uuid4())}).status_code,409)
        self.assertIsNone(db._fetch_one("SELECT * FROM finance_style_costs WHERE style_code='OTHER'"))

    def test_save_audit_failure_rolls_back_every_finance_write(self):
        body=self.prepared(self.change());original=db._fetch_one
        def call(sql,*args,**kwargs):
            if 'INSERT INTO finance_style_cost_audit' in sql:raise ValueError('fixture audit failure')
            return original(sql,*args,**kwargs)
        with patch.object(db,'_fetch_one',side_effect=call):self.assertEqual(self.post('costs/save',body).status_code,400)
        for table in ('finance_style_costs','finance_cost_receipts','finance_style_cost_audit'):
            self.assertEqual(db._fetch_one(f'SELECT count(*) AS n FROM {table}')['n'],0)

    def test_new_colors_sizes_empty_historical_size_and_repeated_migrations_share_cost(self):
        self.save_cost(cost='12.30')
        # A newly added color (same explicit style, different SKU and sizes) inherits immediately.
        p=db.create_product({'categoryKey':'denim','productCode':'EXTRA-COLOR','sku':'EXTRA-COLOR','slug':'extra-color',
            'styleCode':'GT-2026','colorGroup':'GT-2026','colorName':'Extra','title':'Unrelated title','image':'/uploads/fixture.jpg',
            'sizes':['XS'],'sizePrices':[{'sizeCode':'XS','price':1,'stock':0}]})
        db._fetch_one("INSERT INTO order_items(order_id,product_id,product_name,sku,size_code,quantity,unit_price,total_price) VALUES(1,%s,'Extra','EXTRA-COLOR','XS',1,1,1) RETURNING id",(p['id'],))
        with db.get_connection() as conn:
            conn.execute("UPDATE order_items SET size_code=NULL WHERE product_id=1")
            conn.execute('DELETE FROM product_size_prices WHERE product_id=1')
            conn.execute('UPDATE products SET is_active=false WHERE id=1')
            finance.migrate(conn.cursor());finance.migrate(conn.cursor())
        r=self.query();self.assertEqual(r['summary']['totalCostCny'],'123.00');self.assertEqual(r['summary']['missingQuantity'],0)
        self.assertEqual(self.call('finance/costs',role='finance').json['total'],1)
        self.assertEqual(db._fetch_one('SELECT revision FROM finance_style_costs')['revision'],1)

    def test_invalid_money_styles_and_preview_are_rejected(self):
        for value in ('-1','1.001','NaN','Infinity',True,'', '10000000000'):
            self.assertEqual(self.post('costs/preview',self.change(cost=value)).status_code,400)
        omitted=self.change();del omitted['changes'][0]['costCny']
        self.assertEqual(self.post('costs/preview',omitted).status_code,400)
        for body in (None,[],{}, {'changes':[]}, {'changes':self.change()['changes']*2},self.change(code=''), self.change(code='X'*201)):
            self.assertEqual(self.post('costs/preview',body).status_code,400)
        body=self.prepared(self.change());body['previewHash']='wrong'
        self.assertEqual(self.post('costs/save',body).status_code,409)
        for rate in ('0','-1','NaN','7.1234567',None):self.assertEqual(self.post('reports/query',{'exchangeRate':rate}).status_code,400)
        self.assertEqual(self.post('reports/query',{'exchangeRate':'7','statuses':['cancelled']}).status_code,400)
        self.assertEqual(self.post('reports/query',{'exchangeRate':'7','dateFrom':'2026-10-20','dateTo':'2026-10-19'}).status_code,400)
        self.assertEqual(self.call('finance/costs?keyword='+('A'*201),role='finance').status_code,400)

    def test_exports_all_rows_decimal_cells_untrusted_text_and_missing(self):
        self.save_cost()
        db._fetch_one("UPDATE orders SET contact_name='=1+1',order_no=%s WHERE id=1 RETURNING id",('PO0918-'+('Long Buyer '*12)+'-000001',))
        payload={'dateFrom':'2026-09-01','dateTo':'2026-09-30','exchangeRate':'7.123456','page':99}
        response=self.post('reports/export',payload)
        self.assertEqual(response.status_code,200)
        book=load_workbook(BytesIO(response.data))
        self.assertEqual(book.sheetnames,['核算汇总','订单明细','成本缺失'])
        self.assertEqual(book['订单明细'].max_row,7);self.assertEqual(book['成本缺失']['A2'].value,'无缺失成本')
        self.assertEqual(book['订单明细']['M1'].value,'成本生效时间（北京时间）')
        self.assertIn('首次成本补算此前订单',book['核算汇总']['B3'].value)
        self.assertTrue(any(c.value=='=1+1' and c.data_type=='s' for row in book['订单明细'] for c in row))
        self.assertTrue(all(c.data_type!='f' for sheet in book for row in sheet for c in row))
        self.assertEqual(book['订单明细']['J2'].data_type,'n')
        self.assertTrue(book['订单明细']['A2'].alignment.wrap_text)

    def test_public_product_and_order_responses_do_not_include_cost(self):
        self.save_cost(cost='123.45')
        for route,role in [('products','sales'),('inventory/1','warehouse'),('orders','sales')]:
            response=self.call(route,role=role)
            self.assertEqual(response.status_code,200)
            self.assertNotIn('costCny',response.get_data(as_text=True))
            self.assertNotIn('finance_cost',response.get_data(as_text=True))

    def test_product_rebuild_and_deletion_keep_independent_cost_and_history(self):
        self.save_cost(cost='22.50')
        original=self.call('products/3').json['product']
        response=self.call('products/3','PUT',{**original,'title':'Edited fixture','version':original['version']})
        self.assertEqual(response.status_code,200,response.json)
        self.assertEqual(db._fetch_one('SELECT cost_cny FROM finance_style_costs')['cost_cny'],Decimal('22.50'))
        with db.get_connection() as conn:conn.execute('DELETE FROM products WHERE id=3')
        self.assertEqual(self.call('finance/costs/history?styleCode=GT-2026',role='finance').json['total'],1)
        self.assertEqual(self.query()['summary']['totalCostCny'],'202.50')

    def test_preview_change_requires_reconfirmation(self):
        body=self.prepared(self.change())
        self.save_cost(cost='8')
        self.assertEqual(self.post('costs/save',body).status_code,409)
        self.assertEqual(db._fetch_one('SELECT cost_cny FROM finance_style_costs')['cost_cny'],Decimal('8'))
        body=self.prepared(self.change())
        db._fetch_one("UPDATE orders SET status='cancelled' WHERE id=1 RETURNING id")
        # Old-order changes cannot invalidate a prospective-only cost preview.
        self.assertEqual(self.post('costs/save',body).status_code,200)

    def test_pagination_and_export_include_all_matching_groups(self):
        # Enough logical styles and variants to span both page sizes. Fixture database only.
        for n in range(30):
            code=f'PAGE{1000+n}'
            p=db.create_product({'categoryKey':'denim','productCode':code,'sku':code,'slug':code.lower(),
                'familyCode':code,'styleCode':code,'colorName':'Black','title':code,'image':'/uploads/fixture.jpg',
                'sizes':['M'],'sizePrices':[{'sizeCode':'M','price':1,'stock':0}]})
            db._fetch_one('INSERT INTO order_items(order_id,product_id,product_name,sku,size_code,quantity,unit_price,total_price) VALUES(1,%s,%s,%s,\'M\',1,1,1) RETURNING id',(p['id'],code,code))
        page=self.call('finance/costs?page=1&pageSize=12',role='finance').json
        self.assertEqual(page['total'],31);self.assertEqual(len(page['items']),12)
        last=self.call('finance/costs?page=3&pageSize=12',role='finance').json
        self.assertEqual(len(last['items']),7)
        # The cost editor now uses 20 styles, while old 12-style clients remain compatible.
        default=self.call('finance/costs',role='finance').json
        twenty=self.call('finance/costs?page=1&pageSize=20',role='finance').json
        remaining=self.call('finance/costs?page=2&pageSize=20',role='finance').json
        self.assertEqual(default['pageSize'],20);self.assertEqual(len(default['items']),20)
        self.assertEqual(default['items'],twenty['items']);self.assertEqual(twenty['total'],31)
        self.assertEqual(len(remaining['items']),11)
        first_codes={r['styleCode'] for r in twenty['items']};last_codes={r['styleCode'] for r in remaining['items']}
        self.assertFalse(first_codes & last_codes);self.assertEqual(len(first_codes | last_codes),31)
        first=self.query();second=self.query(page=2)
        self.assertEqual(first['total'],32);self.assertEqual(len(first['items']),25)
        self.assertEqual(len(second['items']),7);self.assertEqual(first['summary'],second['summary'])
        exported=self.post('reports/export',{**first['filters'],'page':2})
        self.assertEqual(exported.status_code,200)
        book=load_workbook(BytesIO(exported.data))
        self.assertEqual(book['订单明细'].max_row,37)

    def test_clear_only_affects_new_orders_and_retains_old_cost_history(self):
        self.save_cost(cost='10');stock=self.stock();order_count=db._fetch_one('SELECT count(*) AS n FROM orders')['n']
        body=self.prepared(self.change(cost=None))
        self.assertEqual(self.post('costs/preview',body).json['changes'][0]['oldCostCny'],'10.00')
        with patch.object(finance,'cost_save_time',return_value=datetime.fromisoformat('2026-09-18T09:00:00+08:00')):
            self.assertEqual(self.post('costs/save',body).status_code,200)
        result=self.query();self.assertEqual(result['summary']['missingQuantity'],3)
        self.assertEqual(result['summary']['knownCostCny'],'60.00');self.assertIsNone(result['summary']['totalCostCny']);self.assertIsNone(result['summary']['profitCny'])
        self.assertEqual(self.query(dateFrom='2026-09-16',dateTo='2026-09-16')['summary']['totalCostCny'],'30.00')
        self.assertEqual(self.call('finance/costs?missingOnly=true',role='finance').json['total'],1)
        history=self.call('finance/costs/history?styleCode=GT-2026',role='finance').json
        self.assertEqual(history['total'],2);self.assertIsNone(history['items'][0]['costCny']);self.assertEqual(history['items'][0]['oldCostCny'],'10.00')
        exported=self.post('reports/export',result['filters']);self.assertEqual(exported.status_code,200)
        self.assertEqual(load_workbook(BytesIO(exported.data))['成本缺失'].max_row,3)
        self.assertEqual(self.stock(),stock);self.assertEqual(db._fetch_one('SELECT count(*) AS n FROM orders')['n'],order_count)

    def test_clear_zero_cost_idempotency_conflict_and_restore(self):
        self.save_cost(cost='0')
        body=self.prepared(self.change(cost=None));self.assertEqual(self.post('costs/save',body).status_code,200)
        repeated=self.post('costs/save',body);self.assertEqual(repeated.status_code,200);self.assertTrue(repeated.json['replayed'])
        self.assertEqual(db._fetch_one('SELECT count(*) AS n FROM finance_style_costs')['n'],1)
        self.assertIsNone(db._fetch_one('SELECT cost_cny FROM finance_style_costs')['cost_cny'])
        self.assertEqual(db._fetch_one('SELECT count(*) AS n FROM finance_style_cost_audit')['n'],2)
        self.assertEqual(self.post('costs/save',{**body,'requestId':str(uuid.uuid4())}).status_code,409)
        self.save_cost(cost='20');history=self.call('finance/costs/history?styleCode=GT-2026',role='finance').json['items']
        self.assertEqual(history[0]['costCny'],'20.00');self.assertIsNone(history[0]['oldCostCny'])
        self.assertIsNone(history[1]['costCny']);self.assertEqual(history[1]['oldCostCny'],'0.00')
        self.assertEqual(history[2]['costCny'],'0.00')

    def test_clear_rollback_permissions_and_migration_keep_cost_and_audit(self):
        self.save_cost(cost='10');body=self.prepared(self.change(cost=None));original=db._fetch_one
        for role in ('sales','warehouse','customer'):
            self.assertEqual(self.post('costs/preview',body,role).status_code,403);self.assertEqual(self.post('costs/save',body,role).status_code,403)
        def call(sql,*args,**kwargs):
            if 'INSERT INTO finance_style_cost_audit' in sql:raise ValueError('fixture clear audit failure')
            return original(sql,*args,**kwargs)
        with patch.object(db,'_fetch_one',side_effect=call):self.assertEqual(self.post('costs/save',body).status_code,400)
        self.assertEqual(db._fetch_one('SELECT cost_cny FROM finance_style_costs')['cost_cny'],Decimal('10'))
        self.assertEqual(db._fetch_one('SELECT revision FROM finance_style_costs')['revision'],1)
        self.assertEqual(db._fetch_one('SELECT count(*) AS n FROM finance_cost_receipts')['n'],1)
        with db.get_connection() as conn:finance.migrate(conn.cursor());finance.migrate(conn.cursor())
        self.assertEqual(self.post('costs/save',body).status_code,200)
        with db.get_connection() as conn:finance.migrate(conn.cursor());finance.migrate(conn.cursor())
        self.assertIsNone(db._fetch_one('SELECT cost_cny FROM finance_style_costs')['cost_cny'])
        self.assertEqual(db._fetch_one('SELECT count(*) AS n FROM finance_style_cost_audit')['n'],2)

    def test_two_different_requests_compete_on_initial_style_version(self):
        first=self.prepared(self.change(cost='10'));second=self.prepared(self.change(cost='20'))
        with ThreadPoolExecutor(max_workers=2) as pool:results=list(pool.map(lambda body:self.post('costs/save',body),[first,second]))
        self.assertEqual(sorted(r.status_code for r in results),[200,409])
        self.assertEqual(db._fetch_one('SELECT revision FROM finance_style_costs')['revision'],1)
        self.assertEqual(db._fetch_one('SELECT count(*) AS n FROM finance_style_cost_audit')['n'],1)

    def test_legacy_history_remains_archival_and_never_used_as_style_cost(self):
        db._fetch_one("INSERT INTO finance_cost_history(product_id,size_code,effective_date,cost_cny,sku,style_code,actor) VALUES(1,'M','2020-01-01',2,'GT-2026-1','GT-2026','{}') RETURNING product_id")
        self.assertEqual(self.query()['summary']['missingQuantity'],9)
        self.save_cost(cost='10');self.save_cost(cost=None)
        with db.get_connection() as conn:finance.migrate(conn.cursor());finance.migrate(conn.cursor())
        self.assertEqual(self.query()['summary']['missingQuantity'],9)
        self.assertEqual(self.call('finance/costs/history?productId=1&sizeCode=M',role='finance').json['timeline'][0]['costCny'],'2.00')

    def test_timeline_migration_preserves_original_timestamps_and_cache_only_baseline(self):
        self.save_cost(cost='12',at='2026-09-15T12:00:00+08:00')
        self.save_cost(cost='18',at='2026-09-18T09:00:00+08:00')
        before=self.query()
        audits=db._fetch_all('SELECT * FROM finance_style_cost_audit ORDER BY id')
        with db.get_connection() as conn:
            # An actual old schema: add-column must backfill created_at, not NOW().
            conn.execute('ALTER TABLE finance_style_cost_audit DROP COLUMN effective_at')
            conn.execute('ALTER TABLE finance_style_cost_audit DROP COLUMN is_initial')
            conn.execute("UPDATE products SET style_code='CACHE-ONLY' WHERE id=3")
            conn.execute("INSERT INTO finance_style_costs(style_code,cost_cny,revision,updated_at) VALUES('CACHE-ONLY',5,1,'2026-09-18T10:00:00+08:00')")
            finance.migrate(conn.cursor());finance.migrate(conn.cursor())
        after=db._fetch_all("SELECT * FROM finance_style_cost_audit WHERE style_code='GT-2026' ORDER BY id")
        self.assertEqual(audits,after);self.assertEqual(before['summary'],self.query()['summary'])
        baseline=db._fetch_one("SELECT * FROM finance_style_cost_audit WHERE style_code='CACHE-ONLY'")
        self.assertEqual(baseline['cost_cny'],Decimal('5'))
        self.assertEqual(baseline['effective_at'],datetime.fromisoformat('2026-09-18T10:00:00+08:00'))
        self.assertTrue(baseline['is_initial']);self.assertTrue(after[0]['is_initial']);self.assertFalse(after[1]['is_initial'])
        self.assertEqual(db._fetch_one('SELECT count(*) AS n FROM finance_style_cost_audit')['n'],3)

    def test_export_matches_historical_cost_range_and_effective_time(self):
        self.save_cost(cost='12',at='2026-09-15T12:00:00+08:00')
        self.save_cost(cost='18',at='2026-09-18T09:00:00+08:00')
        result=self.query()
        book=load_workbook(BytesIO(self.post('reports/export',result['filters']).data))
        detail=list(book['订单明细'].iter_rows(min_row=2,values_only=True))
        self.assertEqual(sum(Decimal(str(r[13])) for r in detail),Decimal(result['summary']['totalCostCny']))
        self.assertEqual({(r[11],r[12]) for r in detail},{(12,'2026-09-15 12:00:00'),(18,'2026-09-18 09:00:00')})
        self.assertEqual(book['成本缺失']['A2'].value,'无缺失成本')

    def test_decimal_rounding_is_per_order_line_not_on_the_aggregate(self):
        self.save_cost(cost='0.01')
        with db.get_connection() as conn:
            conn.execute('UPDATE order_items SET unit_price=0.01,total_price=quantity*0.01')
        summary=self.query()['summary']
        self.assertEqual(summary['salesUsd'],'0.09')
        self.assertEqual(summary['salesCny'],'0.63')  # Rounding the total instead would give 0.64.
        self.assertEqual(summary['totalCostCny'],'0.09');self.assertEqual(summary['profitCny'],'0.54')

    def test_large_quantities_and_max_cost_use_exact_decimal_totals(self):
        self.save_cost(cost='9999999999.99')
        with db.get_connection() as conn:
            conn.execute('UPDATE order_items SET quantity=2000000001,unit_price=0.01,total_price=20000000.01 WHERE id=1')
        result=self.query();quantity=2000000008
        self.assertEqual(result['summary']['quantity'],quantity)
        self.assertEqual(result['summary']['totalCostCny'],str(Decimal('9999999999.99')*quantity))
        self.assertEqual(result['summary']['missingQuantity'],0)


if __name__=='__main__':unittest.main()
