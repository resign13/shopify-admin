"""Child process lifetime tests, plus opt-in isolated-database integration."""
import io
import json
import os
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor
import subprocess
import sys
import tempfile
import threading
import time
import unittest
import zipfile
from unittest.mock import patch
from flask import Flask, g
from openpyxl import load_workbook
from PIL import Image
import export_resources
import order_export_process as process
import order_export_worker as worker


class ExportProcessTest(unittest.TestCase):
    def test_read_statistics_are_counted_without_source_data(self):
        stats = worker.ImageReadStats()
        def read(url, source):
            if source and source.get('legacy_channel'):
                raise TimeoutError('private detail')
            return b'fixture'
        measured = stats.wrap(read)
        self.assertEqual(measured('private-url', None), b'fixture')
        self.assertEqual(measured('private-url', {'storage_key': 'private-key'}), b'fixture')
        with self.assertRaises(TimeoutError):
            measured('private-url', {'legacy_channel': 'admin'})
        self.assertEqual(stats.values['public_https'], {'attempts': 1, 'success': 1, 'errors': {}})
        self.assertEqual(stats.values['storage'], {'attempts': 1, 'success': 1, 'errors': {}})
        self.assertEqual(stats.values['legacy'], {'attempts': 1, 'success': 0, 'errors': {'TimeoutError': 1}})
        self.assertNotIn('private', json.dumps(stats.values))

    def test_failure_metadata_preserves_safe_diagnostics_only(self):
        code = """import json,os,sys
from pathlib import Path
payload=json.loads(sys.stdin.buffer.read().decode('utf-8'))
Path(os.environ['TMPDIR'],'status.json').write_text(json.dumps(payload))
"""
        stats = {key: {'attempts': 0, 'success': 0, 'errors': {}} for key in ('public_https','legacy','storage')}
        stats['public_https'] = {'attempts': 1, 'success': 0, 'errors': {'TimeoutError': 1}, 'url': 'private'}
        caches = {key: 1 for key in ('raw_hits','raw_misses','converted_hits','converted_misses','lock_waits','errors')}
        for name, exception in (('image_unavailable', process.ImageUnavailable), ('export_failed', process.ExportFailed)):
            with self.subTest(error=name), self.assertRaises(exception) as result:
                process._run_worker({'error': name, 'phase': 'images', 'completed': 2, 'total': 3,
                    'elapsed_ms': 900, 'exception_type': 'TimeoutError', 'url': 'private', 'message': 'private',
                    'image_read_stats': stats, 'image_cache_stats': caches},
                    command=[sys.executable, '-c', code])
            error = result.exception
            self.assertGreater(error.worker_pid, 0)
            self.assertEqual(error.error_type, 'TimeoutError')
            self.assertEqual(error.progress, {'phase': 'images', 'completed': 2, 'total': 3, 'elapsed_ms': 900})
            self.assertEqual(error.image_read_stats['public_https']['errors'], {'TimeoutError': 1})
            self.assertEqual(error.image_cache_stats, caches)
            self.assertNotIn('private', json.dumps(error.metadata))
            self.assertNotIn('message', error.metadata)

    def test_originating_failure_type_contains_no_message(self):
        try:
            try:
                raise TimeoutError('private endpoint')
            except TimeoutError:
                raise RuntimeError('private wrapping') from None
        except RuntimeError as error:
            self.assertEqual(worker.exception_type(error), 'TimeoutError')

    def test_generation_timeout_has_a_ninety_second_ceiling(self):
        app = Flask(__name__)
        with app.test_request_context('/'), patch.object(process, '_sources', return_value=([], [])):
            for value in ('0', '-1', '90.01', '3600'):
                with self.subTest(value=value), patch.dict(os.environ, {'ORDER_EXPORT_TIMEOUT_SECONDS': value}):
                    with self.assertRaises(RuntimeError):
                        process.build('combined', [])
            for value in ('0.05', '90'):
                with self.subTest(value=value), patch.dict(os.environ, {'ORDER_EXPORT_TIMEOUT_SECONDS': value, 'APP_ENV': 'development'}):
                    with patch.object(process, '_run_worker', return_value=(b'fixture', {})) as run:
                        stream = process.build('combined', [])
                    self.assertEqual(stream.read(), b'fixture')
                    self.assertEqual(run.call_args.kwargs['timeout'], float(value))

    def test_cache_statistics_only_accept_known_numeric_fields(self):
        keys = ('raw_hits','raw_misses','converted_hits','converted_misses','lock_waits','errors')
        stats = {key: 2 for key in keys}
        self.assertEqual(process.safe_cache_stats({**stats, 'url': 'private'}), stats)
        self.assertEqual(process.safe_cache_stats({**stats, 'errors': True}), {})
        self.assertEqual(process.safe_cache_stats({**stats, 'errors': -1}), {})
        self.assertEqual(process.safe_cache_stats({'raw_hits': 1}), {})

    def test_reader_only_uses_snapshot_and_peak_rss_is_numeric(self):
        read = worker.snapshot_reader({('public', False): None, ('private', True): {'storage_key': 'x'}})
        self.assertIsNone(read('public', False))
        self.assertEqual(read('private', True)['storage_key'], 'x')
        with self.assertRaises(PermissionError):
            read('private', False)
        with self.assertRaises(PermissionError):
            read('missing', False)
        peak = worker.peak_rss_bytes()
        self.assertIsInstance(peak, int)
        self.assertGreater(peak, 0)

    def test_timeout_reaps_child_even_when_stdin_is_blocked(self):
        code = """import os,sys,time
from pathlib import Path
Path(sys.argv[1]).write_text(os.environ['TMPDIR'])
time.sleep(3)
Path(sys.argv[2]).write_text('late')
"""
        with tempfile.TemporaryDirectory() as directory:
            marker = Path(directory) / 'started'
            late = Path(directory) / 'late'
            captured = []
            original = subprocess.Popen
            def spawn(*args, **kwargs):
                child = original(*args, **kwargs)
                captured.append(child)
                return child
            started = time.monotonic()
            with patch.object(process.subprocess, 'Popen', side_effect=spawn):
                with self.assertRaises(process.ExportTimeout):
                    process._run_worker({'data': 'x' * 2_000_000}, timeout=0.5,
                        command=[sys.executable, '-c', code, str(marker), str(late)])
            self.assertLess(time.monotonic() - started, 2.5)
            self.assertIsNotNone(captured[0].poll())
            self.assertTrue(marker.exists())
            self.assertFalse(Path(marker.read_text()).exists())
            self.assertFalse(late.exists())

    def test_timeout_reports_only_safe_progress_and_reaps(self):
        code = """import json,os,time
from pathlib import Path
stats={k:{'attempts':0,'success':0,'errors':{}} for k in ('public_https','legacy','storage')}
stats['public_https']={'attempts':3,'success':1,'errors':{'TimeoutError':1},'url':'private'}
Path(os.environ['TMPDIR'],'progress.json').write_text(json.dumps({'phase':'images','completed':3,'total':7,'elapsed_ms':200,'url':'private','image_read_stats':stats,'image_cache_stats':{k:1 for k in ('raw_hits','raw_misses','converted_hits','converted_misses','lock_waits','errors')}}))
time.sleep(3)
"""
        with self.assertRaises(process.ExportTimeout) as result:
            process._run_worker({}, timeout=0.5, command=[sys.executable, '-c', code])
        self.assertEqual(result.exception.progress, {'phase': 'images', 'completed': 3, 'total': 7, 'elapsed_ms': 200})
        self.assertEqual(result.exception.image_read_stats['public_https'], {'attempts': 3, 'success': 1, 'errors': {'TimeoutError': 1}})
        self.assertNotIn('private', json.dumps(result.exception.image_read_stats))
        self.assertEqual(result.exception.image_cache_stats['raw_hits'], 1)

    def test_result_and_environment(self):
        code = """import json,os,sys
from pathlib import Path
body=json.loads(sys.stdin.buffer.read().decode('utf-8'))
root=Path(os.environ['TMPDIR'])
assert os.environ['DB_SCHEMA_AUTO_MIGRATE']=='0'
(root/'export.xlsx').write_bytes(body['text'].encode('utf-8'))
(root/'status.json').write_text(json.dumps({'image_ms':12.0}))
"""
        result, status = process._run_worker({'text': '中文'}, timeout=5, command=[sys.executable, '-c', code])
        self.assertEqual(result.decode('utf-8'), '中文')
        self.assertEqual(status['image_ms'], 12)

    def test_sources_project_fields_and_denials(self):
        app = Flask(__name__)
        with app.test_request_context('/'):
            g.export_sources = {('public', False): None, ('private', True): {
                'storage_channel': 'admin', 'storage_key': 'key', 'byte_size': 2,
                'sha256': 'existing', 'ignored': object()}}
            export_resources.cache()[('denied', True)] = None
            sources, denied = process._sources()
            self.assertEqual(sources[0], ['public', False, None])
            self.assertNotIn('ignored', sources[1][2])
            self.assertEqual(denied, [['denied', True]])
            export_resources.cleanup()


@unittest.skipUnless(os.environ.get('EXPORT_PROCESS_INTEGRATION') == '1', 'requires isolated test DB')
class ExportProcessIntegrationTest(unittest.TestCase):
    def test_three_workbooks_keep_authorized_images(self):
        if os.environ.get('PGHOST') != '127.0.0.1' or not os.environ.get('PGDATABASE', '').endswith('_test'):
            self.fail('An isolated loopback *_test database is required')
        import app as application
        root = application.media.uploads['admin']
        root.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=root) as directory:
            image_path = Path(directory) / 'photo.png'
            Image.new('RGB', (30, 60), 'blue').save(image_path)
            source = {'legacy_path': str(image_path), 'legacy_channel': 'admin'}
            order = {'orderNo': 'fixture', 'shippingFee': 10, 'note': 'fixture note',
                     'labelImageUrls': ['/uploads/photo.png'], 'items': [
                         {'sku': 'one', 'image': '/uploads/photo.png', 'sizeCode': 'M', 'quantity': 1, 'unitPrice': 20}]}
            for kind in ('combined', 'split', 'invoice'):
                with self.subTest(kind=kind), application.app.test_request_context('/'):
                    g.export_sources = {('/uploads/photo.png', False): source, ('/uploads/photo.png', True): source}
                    book = load_workbook(process.build(kind, [order]))
                    images = [image for sheet in book for image in sheet._images]
                    template_images = []
                    if kind == 'invoice':
                        template = load_workbook(application.PROFORMA_TEMPLATE_PATH)
                        template_images = [(image.width, image.height, image.anchor._from.col,
                                            image.anchor._from.row, image._data())
                                           for sheet in template for image in sheet._images]
                    self.assertEqual(len(images), len(template_images) + 2)
                    serialized = [(image.width, image.height, image.anchor._from.col,
                                   image.anchor._from.row, image._data()) for image in images]
                    for original in template_images:
                        self.assertIn(original, serialized)
                        serialized.remove(original)
                    width, height, col, row, raw = serialized[0]
                    self.assertEqual((width, height, col, row), (30, 60, 1, 12 if kind == 'invoice' else 5))
                    for _, _, _, _, raw in serialized:
                        with Image.open(io.BytesIO(raw)) as authorized:
                            red, green, blue = authorized.convert('RGB').getpixel((15, 30))
                            self.assertLess(red, 15)
                            self.assertLess(green, 15)
                            self.assertGreater(blue, 240)
                    cells = '\n'.join(str(cell.value or '') for sheet in book for row in sheet for cell in row)
                    self.assertIn('fixture note', cells)
                    self.assertNotIn('Image unavailable', cells)
                    self.assertNotIn('Attachment images unavailable', cells)

    def test_denied_and_unreadable_images_fail_without_workbook(self):
        if os.environ.get('PGHOST') != '127.0.0.1' or not os.environ.get('PGDATABASE', '').endswith('_test'):
            self.fail('An isolated loopback *_test database is required')
        import app as application
        root = application.media.uploads['admin']
        root.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=root) as directory:
            invalid = Path(directory) / 'invalid.png'
            invalid.write_bytes(b'not-an-image')
            valid = Path(directory) / 'valid.png'
            Image.new('RGB', (30, 60), 'blue').save(valid)
            source = {'legacy_path': str(valid), 'legacy_channel': 'admin'}
            bad_source = {'legacy_path': str(invalid), 'legacy_channel': 'admin'}
            for kind in ('combined', 'split', 'invoice'):
                for failure in ('denied', 'unreadable'):
                    for attachment in (False, True):
                        with self.subTest(kind=kind, failure=failure, attachment=attachment):
                            order = {'orderNo': 'fixture', 'shippingFee': 10,
                                     'labelImageUrls': ['/uploads/attachment.png'] if attachment else [],
                                     'items': [{'sku': 'one', 'image': '/uploads/photo.png',
                                                'sizeCode': 'M', 'quantity': 1, 'unitPrice': 20}]}
                            sources = [['/uploads/photo.png', False, source]]
                            failed_key = ['/uploads/attachment.png', True] if attachment else ['/uploads/photo.png', False]
                            if attachment:
                                sources.append([*failed_key, source])
                            sources = [item for item in sources if item[:2] != failed_key]
                            if failure == 'unreadable':
                                sources.append([*failed_key, bad_source])
                            payload = {'kind': kind, 'orders': [order], 'include_images': True,
                                       'sources': sources, 'denied': [failed_key] if failure == 'denied' else [],
                                       'test_media_roots': {key: str(path) for key, path in application.media.roots.items()},
                                       'test_media_uploads': {key: str(path) for key, path in application.media.uploads.items()}}
                            original = subprocess.Popen
                            observations = []
                            def spawn(*args, **kwargs):
                                child = original(*args, **kwargs)
                                wait = child.wait
                                folder = Path(kwargs['env']['TMPDIR'])
                                def observe(*wait_args, **wait_kwargs):
                                    result = wait(*wait_args, **wait_kwargs)
                                    status = json.loads((folder / 'status.json').read_text(encoding='utf-8'))
                                    observations.append((folder, status.get('error'), (folder / 'export.xlsx').exists()))
                                    return result
                                child.wait = observe
                                return child
                            with patch.object(process.subprocess, 'Popen', side_effect=spawn):
                                with self.assertRaises(process.ImageUnavailable) as caught:
                                    process._run_worker(payload)
                            self.assertGreater(caught.exception.worker_pid, 0)
                            self.assertIn(caught.exception.phase, ('images', 'workbook'))
                            self.assertTrue(caught.exception.error_type)
                            self.assertIn('image_read_stats', caught.exception.metadata)
                            self.assertIn('image_cache_stats', caught.exception.metadata)
                            self.assertEqual(len(observations), 1)
                            folder, error, workbook_exists = observations[0]
                            self.assertEqual(error, 'image_unavailable')
                            self.assertFalse(workbook_exists)
                            self.assertFalse(folder.exists())

    def test_two_real_workers_large_synthetic_workbooks_and_peak_memory(self):
        if os.environ.get('PGHOST') != '127.0.0.1' or not os.environ.get('PGDATABASE', '').endswith('_test'):
            self.fail('An isolated loopback *_test database is required')
        import app as application
        root = application.media.uploads['admin']
        root.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=root) as directory:
            image_path = Path(directory) / 'synthetic.png'
            Image.new('RGB', (1600, 2400), '#365178').save(image_path)
            order = {'orderNo': 'fixture', 'currency': 'USD', 'items': [
                {'productId': i, 'sku': f'synthetic-{i}', 'image': '/uploads/synthetic.png',
                 'sizeCode': 'M', 'quantity': i, 'unitPrice': 10}
                for i in range(1, 21)]}
            payload = {'kind': 'combined', 'orders': [order], 'include_images': True,
                       'sources': [['/uploads/synthetic.png', False,
                                    {'legacy_path': str(image_path), 'legacy_channel': 'admin'}]],
                       'denied': [],
                       'test_media_roots': {key: str(path) for key, path in application.media.roots.items()},
                       'test_media_uploads': {key: str(path) for key, path in application.media.uploads.items()}}

            started = time.perf_counter()
            single = process._run_worker(payload)
            single_seconds = time.perf_counter() - started
            original = subprocess.Popen
            children = []
            overlap = []
            lock = threading.Lock()
            def spawn(*args, **kwargs):
                child = original(*args, **kwargs)
                with lock:
                    children.append(child)
                    if len(children) == 2:
                        overlap.append(all(item.poll() is None for item in children))
                return child
            started = time.perf_counter()
            with patch.object(process.subprocess, 'Popen', side_effect=spawn), ThreadPoolExecutor(max_workers=2) as pool:
                parallel = list(pool.map(lambda _: process._run_worker(payload), range(2)))
            parallel_seconds = time.perf_counter() - started
            self.assertEqual(overlap, [True])
            self.assertEqual(len({result[1]['pid'] for result in parallel}), 2)

            def inspect(result):
                workbook = load_workbook(io.BytesIO(result[0]))
                sheet = workbook.active
                self.assertEqual(len(sheet._images), 20)
                self.assertEqual(sheet['H26'].value, 210)
                self.assertEqual(sheet['J26'].value, 2100)
                values = tuple(tuple(cell.value for cell in row) for row in sheet.iter_rows())
                images = tuple((image.width, image.height) for image in sheet._images)
                with zipfile.ZipFile(io.BytesIO(result[0])) as archive:
                    media = tuple(archive.read(name) for name in archive.namelist() if name.startswith('xl/media/'))
                self.assertEqual(len(media), 20)
                return values, images, media

            expected = inspect(single)
            for result in parallel:
                self.assertEqual(inspect(result), expected)
            peaks = [result[1]['peak_rss_bytes'] for result in parallel]
            self.assertTrue(all(isinstance(value, int) and value > 0 for value in peaks))
            for result in parallel:
                self.assertEqual(result[1]['image_read_stats']['legacy'], {'attempts': 1, 'success': 1, 'errors': {}})
            print(json.dumps({'event': 'synthetic_export_worker_benchmark',
                  'workers': 2, 'images_per_workbook': 20, 'single_seconds': round(single_seconds, 3),
                  'parallel_seconds': round(parallel_seconds, 3),
                  'single_peak_rss_bytes': single[1]['peak_rss_bytes'],
                  'parallel_worker_peak_rss_bytes': peaks,
                  'sum_worker_peaks_upper_bound_bytes': sum(peaks),
                  'workbooks_match': True}))


if __name__ == '__main__':
    unittest.main()
