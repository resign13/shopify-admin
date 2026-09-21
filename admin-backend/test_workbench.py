"""Integration tests. Requires an explicitly named local *_test PostgreSQL database."""
import os
if not os.environ.get('PGDATABASE','').endswith('_test') or os.environ.get('PGHOST') != '127.0.0.1':
    raise RuntimeError('Set PGHOST=127.0.0.1 and an isolated PGDATABASE ending in _test')
import io
import json
import unittest
import uuid
from concurrent.futures import ThreadPoolExecutor
from unittest.mock import patch
from werkzeug.security import generate_password_hash
from openpyxl import load_workbook
import app as application
import db
import workbench


def seed():
    with db.get_connection() as conn:
        conn.execute('TRUNCATE purchase_contracts,admin_order_requests,admin_audit_logs,inventory_receipts,inventory_import_receipts,admin_users,store_users,products,product_categories,orders,homepage_configs,banners RESTART IDENTITY CASCADE')
    tokens={}
    for role,name in [('admin','陈管理员'),('sales','林外贸'),('warehouse','王仓管'),('customer','测试客户')]:
        user=db.create_admin_user({'name':name,'email':f'{role}@gingtto.test','passwordHash':generate_password_hash('Workbench-Test-2026',method='pbkdf2:sha256:1000'),'status':'active','role':role})
        tokens[role]=db.create_admin_session(user['id'])
    db.create_category({'key':'denim','sortOrder':1,'labels':{'zh':'牛仔系列','en':'Denim'}})
    db.create_category({'key':'outerwear','sortOrder':2,'labels':{'zh':'外套系列','en':'Outerwear'}})
    products=[]
    for index,color in enumerate(['深靛蓝','水洗灰','经典黑'],start=1):
        payload={'categoryKey':'denim','productCode':f'GT-2026-{index}','sku':f'GT-2026-{index}','slug':f'gt-2026-{index}',
                 'familyCode':'GT-2026','colorGroup':'GT-2026','colorName':color,'colorHex':'#334155',
                 'title':'GINGTTO 宽松直筒牛仔裤 长标题测试 适合日常通勤与休闲搭配的高品质基础款',
                 'sizes':['M','L','Tall XL'],'sizePrices':[{'sizeCode':size,'price':29.9,'stock':10} for size in ['M','L','Tall XL']],
                 'image':'/uploads/fixture.jpg','gallery':['/uploads/fixture.jpg']*3,'sizeChartImage':'/uploads/fixture.jpg','descriptionImage':'/uploads/fixture.jpg','featured':index==1}
        product=db.create_product(payload);products.append(product)
        db.update_product_inventory(product['id'],{}, {'M':12,'L':6,'Tall XL':3})
        db._fetch_one("UPDATE product_size_prices SET pending_inbound=4,pending_inspection=6 WHERE product_id=%s AND size_code='M' RETURNING id",(product['id'],))
        db._fetch_one("UPDATE product_size_prices SET pending_inbound=2,pending_inspection=3 WHERE product_id=%s AND size_code='L' RETURNING id",(product['id'],))
    buyer=db.create_store_user({'name':'Alex Morgan','companyName':'Northline Apparel','email':'buyer@gingtto.test','passwordHash':generate_password_hash('Fixture-Only'),'status':'active'})
    for index,(status,created,country) in enumerate([('paid','2026-09-17 16:30:00+00','Germany'),('pending_payment','2026-09-18 09:00:00+00','United States'),('cancelled','2026-09-18 10:00:00+00','Germany'),('completed','2026-09-16 12:00:00+00','France')],start=1):
        order=db._fetch_one('INSERT INTO orders(order_no,store_user_id,status,contact_name,phone,country,shipping_address,total_amount,shipping_fee,created_at) VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s) RETURNING id',
                           (f'GT20260918{index:04}',buyer['id'],status,'Alex','123456',country,'18 Market Street',99.7,10,created))
        for product,quantity in [(products[0],2),(products[1],1)]:
            db._fetch_one('INSERT INTO order_items(order_id,product_id,product_name,sku,size_code,quantity,unit_price,total_price) VALUES(%s,%s,%s,%s,%s,%s,%s,%s) RETURNING id',(order['id'],product['id'],product['name']['zh'],product['sku'],'M',quantity,29.9,round(29.9*quantity,2)))
    return tokens,products


class WorkbenchTest(unittest.TestCase):
    def setUp(self):
        self.tokens,self.products=seed()
        self.app=application.app
        self.app.config['TESTING']=True

    def call(self,path,method='GET',body=None,role='admin',**kwargs):
        with self.app.test_client() as client:
            return client.open('/api/admin/'+path,method=method,json=body,headers={'Authorization':'Bearer '+self.tokens[role]},**kwargs)

    def detail(self,role='admin'):
        response=self.call('inventory/1',role=role)
        self.assertEqual(response.status_code,200,response.json)
        return response.json['product']

    def test_pagination_summary_and_permissions(self):
        data=self.call('inventory?page=1&pageSize=25').json
        self.assertEqual(data['total'],3)
        self.assertEqual(data['summary']['stock'],90)
        self.assertEqual(data['summary']['pendingInbound'],18)
        customer=self.call('inventory?page=1&pageSize=25',role='customer').json
        self.assertNotIn('pendingInbound',customer['summary'])
        self.assertNotIn('contractPending',customer['items'][0]['sizePrices'][0])
        self.assertNotIn('pendingInbound',customer['items'][0]['sizePrices'][0])
        for role in ['sales','warehouse','customer']:
            self.assertEqual(self.call('audit-logs?page=1',role=role).status_code,403)
        self.assertEqual(self.call('inventory/1',method='PUT',body={'sizeStocks':{'M':1}},role='customer').status_code,403)
        self.assertEqual(self.call('products?page=1',role='warehouse').status_code,403)
        self.assertEqual(self.call('catalog-options',role='warehouse').status_code,200)
        self.assertEqual(self.call('orders?page=1&pageSize=25',role='warehouse').status_code,200)
        self.assertEqual(self.call('products?page=1&pageSize=999').status_code,400)

    def test_dashboard_timezone_cancelled_and_matching_lines(self):
        response=self.call('dashboard?view=workbench&dateFrom=2026-09-18&dateTo=2026-09-18&style=GT-2026-1')
        self.assertEqual(response.status_code,200,response.json)
        data=response.json
        self.assertEqual(data['metrics']['orders'],2)
        self.assertEqual(data['metrics']['units'],4)
        self.assertAlmostEqual(data['metrics']['amount'],119.6)
        self.assertEqual(data['metrics']['customers'],1)
        self.assertEqual(data['previous']['orders'],0)
        self.assertEqual(data['trend'][0]['orders'],2)
        self.assertEqual(data['snapshot']['stock'],90)

    def test_manual_inventory_version_conflict_and_audit(self):
        before=self.detail()
        body={'version':before['version'],'sizeStocks':{'M':11},'pendingInboundBySize':{'L':3}}
        self.assertEqual(self.call('inventory/1','PUT',body,role='warehouse').status_code,200)
        self.assertEqual(self.call('inventory/1','PUT',body).status_code,409)
        after=self.detail()
        self.assertEqual(after['stock'],31)
        logs=self.call('audit-logs?page=1&module=inventory&objectId=1').json
        self.assertGreater(logs['total'],0)
        self.assertTrue(all(row['actor']['role']=='warehouse' for row in logs['items']))

    def test_inspection_stage_flow_and_overdraft(self):
        db._fetch_one("UPDATE product_size_prices SET contract_pending=100,pending_inspection=0,pending_inbound=0 WHERE product_id=1 AND size_code='M' RETURNING id")
        def stage(**fields):
            p=self.detail()
            return self.call('inventory/1','PUT',{'version':p['version'],'sizeStocks':{},**fields})
        self.assertEqual(stage(pendingInspectionBySize={'M':20}).status_code,200)
        self.assertEqual(stage(pendingInboundBySize={'M':10}).status_code,200)
        size=self.detail()['sizePrices'][0]
        self.assertEqual((size['contractPending'],size['pendingInspection'],size['pendingInbound']),(80,10,10))
        before=self.detail()
        self.assertEqual(stage(pendingInboundBySize={'M':21}).status_code,400)
        self.assertEqual(self.detail()['sizePrices'],before['sizePrices'])
        self.assertEqual(stage(pendingInspectionBySize={'M':91}).status_code,400)
        r=self.call('inventory/1/receive','POST',{'version':self.detail()['version'],'requestId':str(uuid.uuid4())})
        self.assertEqual(r.status_code,200,r.json)
        size=self.detail()['sizePrices'][0]
        self.assertEqual((size['stock'],size['contractPending'],size['pendingInspection'],size['pendingInbound']),(20,80,10,0))
        customer=self.detail('customer')['sizePrices'][0]
        self.assertNotIn('pendingInspection',customer)

    def test_inspection_migration_runs_once(self):
        with db._connect() as conn:
            cur=conn.cursor()
            cur.execute('ALTER TABLE product_size_prices DROP COLUMN pending_inspection')
            cur.execute("UPDATE product_size_prices SET contract_pending=100,pending_inbound=20 WHERE product_id=1 AND size_code='M'")
            db._apply_schema_migrations(cur)
            db._apply_schema_migrations(cur)
            cur.execute("SELECT stock,contract_pending,pending_inspection,pending_inbound FROM product_size_prices WHERE product_id=1 AND size_code='M'")
            self.assertEqual(tuple(cur.fetchone().values()),(10,80,0,20))
            conn.rollback()

    def test_inspection_excel_delta_conflict_and_preservation(self):
        raw=self.workbook(lambda s:setattr(s['O2'],'value',8))
        preview=self.upload('preview',raw).json
        self.assertEqual(preview['errors'],[])
        self.assertEqual(preview['rows'][0]['contractPending'],10)
        self.assertEqual(self.upload('confirm',raw,preview['fileHash']).status_code,200)
        size=self.detail()['sizePrices'][0]
        self.assertEqual((size['contractPending'],size['pendingInspection']),(10,8))
        p=self.product_payload(1)
        r=self.call('products/save-group','POST',{'products':[p],'versions':{'1':p['version']}})
        self.assertEqual(r.status_code,200,r.json)
        self.assertEqual(self.detail()['sizePrices'][0]['pendingInspection'],8)
        raw=self.workbook(lambda s:setattr(s['M2'],'value',5))
        preview=self.upload('preview',raw).json
        db.update_product_inventory(1,{},None,None,{'M':9})
        self.assertEqual(self.upload('confirm',raw,preview['fileHash']).status_code,400)

    def test_inbound_validation_and_atomic_rollback(self):
        before=self.detail()
        response=self.call('inventory/1','PUT',{'version':before['version'],'sizeStocks':{'M':99},'pendingInboundBySize':{'L':100}})
        self.assertEqual(response.status_code,400)
        self.assertEqual(self.detail()['stock'],before['stock'])
        self.assertEqual(self.call('audit-logs?page=1').json['total'],0)

    def test_receive_idempotency_and_concurrency(self):
        before=self.detail();body={'version':before['version'],'requestId':str(uuid.uuid4())}
        def receive(_): return self.call('inventory/1/receive','POST',body)
        with ThreadPoolExecutor(max_workers=2) as pool:
            responses=list(pool.map(receive,range(2)))
        self.assertTrue(all(r.status_code==200 for r in responses),[r.json for r in responses])
        self.assertEqual(sum(bool(r.json.get('replayed')) for r in responses),1)
        self.assertEqual(responses[0].json['receivedUnits'],6)
        after=self.detail()
        self.assertEqual(after['stock'],36)
        self.assertEqual(sum(s['contractPending'] for s in after['sizePrices']),21)
        self.assertEqual(sum(s['pendingInbound'] for s in after['sizePrices']),0)
        self.assertEqual(self.call('inventory/1/receive','POST',{**body,'requestId':str(uuid.uuid4())}).status_code,409)

    def test_batch_all_or_nothing_and_group_versions(self):
        a=workbench.product_detail(1);b=workbench.product_detail(2)
        response=self.call('products/batch','POST',{'ids':[1,2],'versions':{'1':a['version'],'2':'stale'},'featured':True})
        self.assertEqual(response.status_code,409)
        self.assertFalse(workbench.product_detail(2)['featured'])
        self.assertEqual(self.call('audit-logs?page=1').json['total'],0)
        response=self.call('products/batch','POST',{'ids':[1,2],'versions':{'1':a['version'],'2':b['version']},'featured':True})
        self.assertEqual(response.status_code,200,response.json)
        self.assertEqual(self.call('products/1/family').status_code,200)

    def workbook(self,modify=None):
        raw=self.call('inventory/export?view=workbench&keyword=GT-2026-1').data
        book=load_workbook(io.BytesIO(raw));sheet=book.active
        if modify: modify(sheet)
        output=io.BytesIO();book.save(output);return output.getvalue()

    def upload(self,path,raw,file_hash=None):
        data={'file':(io.BytesIO(raw),'inventory.xlsx')}
        if file_hash:data['fileHash']=file_hash
        with self.app.test_client() as client:
            return client.post('/api/admin/inventory/import/'+path,data=data,headers={'Authorization':'Bearer '+self.tokens['admin']})

    def test_excel_v2_receive_fields_and_replay(self):
        raw=self.workbook(lambda s:setattr(s['M2'],'value',5))
        preview=self.upload('preview',raw)
        self.assertEqual(preview.status_code,200,preview.json)
        self.assertEqual(preview.json['errors'],[])
        self.assertEqual(preview.json['summary']['changedFieldCount'],2)
        result=self.upload('confirm',raw,preview.json['fileHash'])
        self.assertEqual(result.status_code,200,result.json)
        self.assertEqual(result.json['updatedFields'],2)
        replay=self.upload('confirm',raw,preview.json['fileHash'])
        self.assertEqual(replay.json['updatedFields'],0)
        self.assertTrue(replay.json['alreadyProcessed'])

    def test_excel_conflict_formula_and_limits(self):
        raw=self.workbook(lambda s:setattr(s['M2'],'value',99))
        self.assertTrue(self.upload('preview',raw).json['errors'])
        raw=self.workbook(lambda s:setattr(s['H2'],'value','=1+1'))
        self.assertTrue(self.upload('preview',raw).json['errors'])
        raw=self.workbook(lambda s:setattr(s['H2'],'value',12))
        preview=self.upload('preview',raw).json
        db.update_product_inventory(1,{'M':13})
        response=self.upload('confirm',raw,preview['fileHash'])
        self.assertEqual(response.status_code,400,response.json)
        self.assertEqual(self.detail()['sizePrices'][0]['stock'],13)

    def product_payload(self, product_id):
        p = self.call(f'products/{product_id}').json['product']
        return {**p, 'title':p['name']['zh'], 'familyCode':p['colorGroup']}

    def test_six_chinese_color_products(self):
        import copy
        rows=[]
        for color in ['黑色','白色','蓝色','棕色','灰色','绿色']:
            row=copy.deepcopy(self.product_payload(1))
            row.pop('id',None)
            row['productCode']=row['sku']='CS5101-'+color
            row['slug']='cs5101-'+color
            row['colorName']=color
            rows.append(row)
        response=self.call('products/save-group','POST',{'products':rows})
        self.assertEqual(response.status_code,200,response.json)
        self.assertEqual(len(response.json['items']),6)
        self.assertEqual(len({r['slug'] for r in response.json['items']}),6)
        duplicate=self.call('products/save-group','POST',{'products':[rows[-1]]})
        self.assertEqual(duplicate.status_code,400)
        self.assertIn('商品链接标识',duplicate.json['message'])

    def test_group_save_preserves_pending_and_rolls_back(self):
        rows=[self.product_payload(i) for i in [1,2]]
        versions={str(p['id']):p['version'] for p in rows}
        rows[0]['title']='新的标题'
        response=self.call('products/save-group','POST',{'products':rows,'versions':versions})
        self.assertEqual(response.status_code,200,response.json)
        sizes=self.detail()['sizePrices']
        self.assertEqual(sum(s['pendingInbound'] for s in sizes),6)
        self.assertEqual(sum(s['contractPending'] for s in sizes),21)
        rows=[self.product_payload(i) for i in [1,2]]
        versions={str(p['id']):p['version'] for p in rows}
        rows[0]['title']='应回滚标题'
        rows[1]['sizePrices']=rows[1]['sizePrices'][1:]
        rows[1]['sizes']=[p['sizeCode'] for p in rows[1]['sizePrices']]
        response=self.call('products/save-group','POST',{'products':rows,'versions':versions})
        self.assertEqual(response.status_code,400,response.json)
        self.assertEqual(self.detail()['name']['zh'],'新的标题')

    def test_excel_v1_and_invalid_rows(self):
        def legacy(sheet):
            sheet.delete_cols(13,2)
            for row in range(2,sheet.max_row+1):sheet.cell(row,1,'inventory-v1')
            sheet['H2']=11
        raw=self.workbook(legacy)
        self.assertEqual(self.upload('preview',raw).status_code,400)
        for value in ['',-2147483649,1.5,True,2147483648]:
            with self.subTest(value=value):
                raw=self.workbook(lambda s:setattr(s['H2'],'value',value))
                self.assertTrue(self.upload('preview',raw).json['errors'])
        raw=self.workbook(lambda s:s.append([c.value for c in s[2]]))
        self.assertTrue(self.upload('preview',raw).json['errors'])
        raw=self.workbook(lambda s:setattr(s['D2'],'value','WRONG-SKU'))
        self.assertTrue(self.upload('preview',raw).json['errors'])
        raw=self.workbook(lambda s:setattr(s['G2'],'value','UNKNOWN'))
        self.assertTrue(self.upload('preview',raw).json['errors'])

    def test_server_pagination_and_full_filter_summary(self):
        with db.get_connection() as conn:
            conn.execute("""INSERT INTO products(category_id,sku,slug,product_code,stock,price,is_active,main_image_url)
                SELECT 1,'PAGE-'||n,'page-'||n,'PAGE-'||n,0,0,TRUE,'/uploads/fixture.jpg' FROM generate_series(1,30) AS n""")
        first=self.call('products?page=1&pageSize=25&keyword=PAGE-&sort=id&direction=asc').json
        second=self.call('products?page=2&pageSize=25&keyword=PAGE-&sort=id&direction=asc').json
        self.assertEqual(first['total'],30)
        self.assertEqual(len(first['items']),25)
        self.assertEqual(len(second['items']),5)
        self.assertFalse({p['id'] for p in first['items']} & {p['id'] for p in second['items']})
        self.assertEqual(second['summary']['stock'],0)

    def test_config_conflicts_and_customer_receipt_denied(self):
        response=self.call('home-config');self.assertEqual(response.status_code,200,response.json)
        config=response.json['config']
        config['heroBanners']={key:'/uploads/fixture.jpg' for key in ['bestSeller','newArrival','specialPrice']}
        self.assertEqual(self.call('home-config','PUT',config).status_code,200)
        self.assertEqual(self.call('home-config','PUT',config).status_code,409)
        for path,method in [('inventory/1/receive','POST'),('inventory/export','GET'),('inventory/import/preview','POST')]:
            self.assertEqual(self.call(path,method,role='customer').status_code,403)

    def test_password_audit_has_no_credentials(self):
        response=self.call('admin-users/2','PUT',{'name':'林外贸','email':'sales@gingtto.test','role':'sales','status':'active','password':'New-Fixture-Password'})
        self.assertEqual(response.status_code,200,response.json)
        logs=db._fetch_all('SELECT before_data,after_data FROM admin_audit_logs')
        text=json.dumps(logs,ensure_ascii=False)
        self.assertNotIn('password_hash',text)
        self.assertNotIn('New-Fixture-Password',text)
        self.assertIn('passwordChanged',text)

    def test_audit_failure_rolls_back_business_write(self):
        with db.get_connection() as conn:
            conn.execute("ALTER TABLE admin_audit_logs ADD CONSTRAINT fixture_failure CHECK(action='IMPOSSIBLE')")
        try:
            with self.assertRaises(Exception):
                self.call('inventory/1','PUT',{'sizeStocks':{'M':22}})
            self.assertEqual(self.detail()['stock'],30)
        finally:
            with db.get_connection() as conn:conn.execute('ALTER TABLE admin_audit_logs DROP CONSTRAINT fixture_failure')


if __name__=='__main__':
    unittest.main()
