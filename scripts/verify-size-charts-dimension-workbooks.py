"""Edit browser-downloaded test workbooks to exercise real Excel row operations."""
from pathlib import Path
from openpyxl import load_workbook

p=Path('output/playwright')
source=p/'size-charts-preview-cm-maintenance.xlsx'
for invalid in [False,True]:
    book=load_workbook(source);sheet=book.worksheets[0]
    assert sheet['C5'].value=='cm' and len(sheet._images)==1
    # Same operation as deleting the four top measurement rows in Excel.
    sheet.delete_rows(9,4)
    sheet['A9']='膝围成衣';sheet['B10']='Hem Opening'
    last=sheet.max_row
    data=[c.value for c in sheet[last]]
    assert data[0]=='外长'
    sheet.delete_rows(last);sheet.insert_rows(9)
    for col,value in enumerate(data,1):sheet.cell(9,col,value)
    sheet.append(['胸围','Chest','cm',94,96,0,None])
    if invalid:sheet.append(data)
    book.save(p/('size-charts-dimensions-invalid.xlsx' if invalid else 'size-charts-dimensions-import.xlsx'))
print('Dimension deletion, rename, reorder, addition and duplicate-row fixtures prepared')
