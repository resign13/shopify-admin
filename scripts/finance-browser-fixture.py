"""Finance browser fixture; never point it at a manual or production database."""
import os
import sys
import uuid
from pathlib import Path

if os.environ.get('PGHOST') != '127.0.0.1' or not os.environ.get('PGDATABASE', '').endswith('_test'):
    raise RuntimeError('Explicit isolated local *_test database required')

import psycopg
from psycopg.sql import SQL, Identifier
from werkzeug.security import generate_password_hash

with psycopg.connect(host='127.0.0.1', port=os.environ['PGPORT'], dbname='postgres',
                     user=os.environ['PGUSER'], password=os.environ['PGPASSWORD'], autocommit=True) as conn:
    name = os.environ['PGDATABASE']
    if not conn.execute('SELECT 1 FROM pg_database WHERE datname=%s', (name,)).fetchone():
        conn.execute(SQL('CREATE DATABASE {}').format(Identifier(name)))

root = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(root / 'admin-backend'))
os.environ['ORDER_VOICE_ENABLED'] = '0'
import test_workbench as fixtures

if __name__ == '__main__':
    fixtures.seed()
    with fixtures.db.get_connection() as conn:
        conn.execute('TRUNCATE finance_cost_history,finance_cost_versions,finance_cost_audit,finance_cost_receipts,finance_settings,finance_style_costs,finance_style_cost_audit RESTART IDENTITY')
        conn.execute("DELETE FROM product_size_prices WHERE product_id=3 AND size_code='L'")
        conn.execute("DELETE FROM product_sizes WHERE product_id=3 AND size_code='L'")
        conn.execute("UPDATE products SET style_code='GT-OTHER' WHERE id=2")
    fixtures.db.create_admin_user({'name':'财务测试员','email':'finance@gingtto.test',
        'passwordHash':generate_password_hash('Workbench-Test-2026', method='pbkdf2:sha256:1000'),
        'status':'active','role':'finance'})
    # Historical events in the isolated fixture let the UI prove that current
    # edits/clears do not reprice earlier orders, including same-day changes.
    with fixtures.db.get_connection() as conn:
        for old,cost,stamp in [(None,12,'2026-09-15T12:00:00+08:00'),(12,18,'2026-09-18T09:00:00+08:00')]:
            conn.execute("INSERT INTO finance_style_cost_audit(style_code,old_cost_cny,cost_cny,actor,request_id,created_at,effective_at,is_initial) VALUES('GT-2026',%s,%s,'{\"name\":\"历史成本夹具\"}',%s,%s,%s,%s)",(old,cost,uuid.uuid4(),stamp,stamp,old is None))
        conn.execute("INSERT INTO finance_style_costs(style_code,cost_cny,revision,updated_at) VALUES('GT-2026',18,2,'2026-09-18T09:00:00+08:00')")
    if '--excel' in sys.argv:
        from io import BytesIO
        from openpyxl import load_workbook
        empty_product=fixtures.db.create_product({'categoryKey':'denim','productCode':'GT-EMPTY','sku':'GT-EMPTY',
            'slug':'gt-empty-excel','styleCode':'GT-EMPTY','title':'Excel isolated fixture','image':'','sizes':['M'],
            'sizePrices':[{'sizeCode':'M','price':9,'stock':12}]})
        fixtures.db._fetch_one("UPDATE products SET style_code='GT-EMPTY' WHERE id=%s RETURNING id",(empty_product['id'],))
        user=fixtures.db._fetch_one("SELECT id FROM admin_users WHERE email='finance@gingtto.test'")
        token=fixtures.db.create_admin_session(user['id'])
        response=fixtures.application.app.test_client().post('/api/admin/finance/costs/export',json={},headers={'Authorization':'Bearer '+token})
        assert response.status_code==200
        book=load_workbook(BytesIO(response.data));sheet=book['商品成本']
        for cells in sheet.iter_rows(min_row=2):
            if cells[0].value=='GT-2026':cells[3].value=18
            if cells[0].value=='GT-OTHER':cells[3].value=0
        folder=root/'output/playwright';folder.mkdir(parents=True,exist_ok=True)
        book.save(folder/'cost-excel-valid.xlsx')
        for cells in sheet.iter_rows(min_row=2):
            if cells[0].value=='GT-OTHER':cells[3].value='=1+1'
        book.save(folder/'cost-excel-invalid.xlsx')
        book.close()
    fixtures.application.app.run(host='127.0.0.1', port=5362, threaded=True, use_reloader=False)
