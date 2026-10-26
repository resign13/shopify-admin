from openpyxl import load_workbook
from pathlib import Path
p=Path('output/playwright')
b=load_workbook(p/'size-charts-maintenance.xlsx');s=b.worksheets[0]
assert len(s._images)==1 and b['_尺码表模板'].sheet_state=='veryHidden'
assert s['D9'].value==28 and s['D14'].value==0
s['D9']=None;s['E9']=31.25;s['D11']=0
b.save(p/'size-charts-import.xlsx')
e=load_workbook(p/'size-charts-english.xlsx');assert len(e.worksheets)==1 and len(e.worksheets[0]._images)==1
vals=[c.value for r in e.worksheets[0] for c in r if c.value is not None]
assert 'Waist (Stretched)' not in vals and 'Hip' in vals and 'Inseam' in vals
assert e.worksheets[0]['C9'].value==28
print('Downloaded Excel data, image, signed metadata and English isolation passed')

