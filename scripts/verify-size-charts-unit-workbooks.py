"""Inspect downloads from the isolated browser unit conversion regression."""
from pathlib import Path
from openpyxl import load_workbook

for unit in ('cm','in'):
    for mode in ('maintenance','english'):
        book=load_workbook(Path('output/playwright') / f'size-charts-{unit}-{mode}.xlsx')
        sheet=book.worksheets[0]
        assert sheet['C5'].value==unit
        assert len(sheet._images)==1 and sheet._images[0].anchor._from.row==0
        unit_col=3 if mode=='maintenance' else 2
        assert all(sheet.cell(row,unit_col).value==unit for row in range(9,sheet.max_row+1))
        assert sheet.cell(sheet.max_row,unit_col+1).value==(100.51 if unit=='cm' else 39.57)
        assert sheet.sheet_view.showGridLines and sheet.sheet_view.pane is None
        assert sheet['A8'].value==('测量维度' if mode=='maintenance' else 'Measurement')
        assert len(book.worksheets)==(2 if mode=='maintenance' else 1)
print('Four cm/in maintenance/customer workbooks: uniform units, converted values, images and layout passed')
